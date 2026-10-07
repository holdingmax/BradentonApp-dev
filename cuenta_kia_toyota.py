"""
Cuenta corriente de Kia (Gettel) y Toyota (pedido del usuario, 2026-10-07).

Cómo funciona (confirmado con el manager y con los Excel de Cierre desde
mayo 2024):
- Los autos de las concesionarias cargan nafta con un vale y no pagan en el
  momento: el POS lo registra como Local Account (Store Info). Toyota empezó
  el 24/10/2025; antes era solo Kia.
- Lo que se les cobra por cada período es (vales − $0.02 por galón de
  rebate) × 1.03 (el 3% del recargo de la tarjeta): REBATE_PER_GALLON y
  CARD_CHARGE, las mismas fórmulas de las hojas "Pendiente" del Excel.
- Cada una o dos semanas pagan lo acumulado con una American Express, en
  cobros de hasta $999; el POS lo registra como el departamento LOCAL ACCT
  (la columna VS del Excel) y lo deposita J.H.

Cálculo puro (sin UI ni base) más la lectura de las hojas de Gettel-Toyota
de los Excel de Cierre (o de los Excel de control de Gettel), que es de
donde sale el historial: cargas por día y por empresa (con galones) y los
pagos (con la empresa entre paréntesis desde febrero 2026).
"""

import calendar
import os
import re
from datetime import date, datetime, timedelta

KIA = "Kia"
TOYOTA = "Toyota"
COMPANIES = (KIA, TOYOTA)
TOYOTA_START = "2025-10-24"  # primer vale de Toyota: antes, todo pago es de Kia
# Cada una paga con su propia tarjeta (el manager, 2026-10-07) y el recibo
# trae los últimos 4 dígitos: así se asignaron los pagos sin nombre de
# octubre 2025 a febrero 2026 (Toyota ya pagaba con la 0953 desde noviembre
# 2025). La Amex 1008 (16/02/2026) no se sabe de quién es.
KNOWN_CARDS = {"3924": KIA, "5145": KIA, "5449": KIA, "0953": TOYOTA, "1363": TOYOTA, "9451": TOYOTA}
# Un pago anulado (el recibo dice VOID y el POS no lo cobró) no cuenta.
VOIDED = "Anulado"
REBATE_PER_GALLON = 0.02
CARD_CHARGE = 0.03
TOLERANCE = 0.005
# Un pago que se atrasa más que esto se avisa (pagan cada una o dos semanas).
LATE_PAYMENT_DAYS = 21


def to_charge(amount, gallons):
    """Lo que se cobra por unas cargas: (vales − rebate por galón) × 1.03."""
    return round(((amount or 0.0) - REBATE_PER_GALLON * (gallons or 0.0)) * (1 + CARD_CHARGE), 2)


# ---------------------------------------------------------------------------
# Lectura de los Excel (hojas de Gettel-Toyota)
# ---------------------------------------------------------------------------

def _norm(value):
    return str(value).strip().upper() if value is not None else ""


def _num(value):
    return float(value) if isinstance(value, (int, float)) and not isinstance(value, bool) else None


def _fix_date(day, file_month):
    """Una fecha tipeada al revés (12/01 en vez de 01/12) se da vuelta si así cae cerca del mes del archivo."""
    if file_month is None or day.day > 12:
        return day
    lo = date(file_month[0], file_month[1], 1) - timedelta(days=40)
    hi = date(file_month[0], file_month[1], 1) + timedelta(days=70)
    if lo <= day <= hi:
        return day
    try:
        swapped = date(day.year, day.day, day.month)
    except ValueError:
        return day
    for year in (file_month[0], file_month[0] - 1, file_month[0] + 1):
        try:
            candidate = swapped.replace(year=year)
        except ValueError:
            continue
        if lo <= candidate <= hi:
            return candidate
    return day


def _file_month(path):
    """(año, mes) del nombre del archivo: "Cierre 09-25.xlsx", "Gettel-Toyota Enero.xlsx" con la carpeta, etc."""
    name = os.path.basename(path)
    match = re.search(r"(\d{2})-(\d{2})\b", name)
    if match and 1 <= int(match.group(1)) <= 12:
        return 2000 + int(match.group(2)), int(match.group(1))
    match = re.search(r"[\\/](\d{4})[\\/](\d{2}) ", path)
    if match:
        return int(match.group(1)), int(match.group(2))
    return None


def read_control_workbook(path):
    """
    Lee las hojas de Gettel-Toyota de un Excel de Cierre (o de control de
    Gettel). Devuelve {"days": [{date, la, kia, kia_gal, toyota, toyota_gal,
    main}], "payments": [{date, transc, amount, company}], "pos": [{date, la,
    vs}]}. "main" es la hoja del mes (no la de días pendientes del mes
    anterior). La empresa de un pago sale de la aclaración "(Kia)"/"(Toyota)"
    del primer renglón de cada pago; antes de TOYOTA_START, todo es Kia.
    """
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    file_month = _file_month(path)
    result = {"days": [], "payments": [], "pos": []}
    try:
        for sheet in workbook.worksheets:
            title = sheet.title.upper()
            if title.startswith("STORE INFO"):
                result["pos"].extend(_read_store_info(sheet, file_month))
                continue
            if not any(key in title for key in ("GETTEL", "PENDIENTE", "PAGO", "PEND ")):
                continue
            rows = list(sheet.iter_rows(max_col=16, values_only=True))
            header = None
            for i, row in enumerate(rows[:6]):
                if "LOCAL ACCOUNT" in [_norm(v) for v in row]:
                    header = ("days", i)
                    break
            if header is None:
                for i, row in enumerate(rows[:6]):
                    if any("TRANSC" in _norm(v) for v in row):
                        header = ("payments", i)
                        break
            if header is None:
                continue
            kind, i = header
            labels = [_norm(v) for v in rows[i]]
            if kind == "days":
                result["days"].extend(_read_days(rows[i + 1:], labels, "PEND" not in title, file_month))
            else:
                result["payments"].extend(_read_payments(rows[i + 1:], labels, file_month))
    finally:
        workbook.close()
    return result


def _read_days(rows, labels, main, file_month):
    la = labels.index("LOCAL ACCOUNT")
    kia = next((j for j, x in enumerate(labels) if x in ("1 X 1", "GETTEL")), None)
    toyota = next((j for j, x in enumerate(labels) if x == "TOYOTA"), None)
    date_col = labels.index("FECHA") if "FECHA" in labels else 0
    days = []
    for row in rows:
        value = row[date_col] if date_col < len(row) else None
        if not isinstance(value, datetime):
            continue

        def cell(j):
            return _num(row[j]) if j is not None and j < len(row) else None

        days.append({
            "date": _fix_date(value.date(), file_month).isoformat(), "la": cell(la),
            "kia": cell(kia), "kia_gal": cell(kia + 1) if kia is not None else None,
            "toyota": cell(toyota), "toyota_gal": cell(toyota + 1) if toyota is not None else None,
            "main": main,
        })
    return days


def _read_payments(rows, labels, file_month):
    transc = next(j for j, x in enumerate(labels) if "TRANSC" in x)
    amount = next((j for j, x in enumerate(labels) if x.startswith("TOTAL CUPON")), transc + 1)
    payments, last, company = [], None, None
    for row in rows:
        if isinstance(row[0], datetime):
            day = _fix_date(row[0].date(), file_month).isoformat()
            if day != last:
                company = None
            last = day
        number = row[transc] if transc < len(row) else None
        value = _num(row[amount]) if amount < len(row) else None
        if value is None or last is None or not re.fullmatch(r"\d+(\.0)?", str(number or "").strip()):
            continue
        note = " ".join(str(v) for v in row[amount + 1:] if isinstance(v, str)).upper()
        if "KIA" in note:
            company = KIA
        elif "TOYOTA" in note:
            company = TOYOTA
        who = company or (KIA if last < TOYOTA_START else None)
        payments.append({"date": last, "transc": str(number).replace(".0", ""), "amount": round(value, 2),
                         "company": who})
    return payments


def _read_store_info(sheet, file_month):
    rows = list(sheet.iter_rows(max_col=30, values_only=True))
    for i, row in enumerate(rows[:6]):
        labels = [_norm(v) for v in row]
        if "LOCAL ACCOUNT" in labels and "VS" in labels:
            la, vs = labels.index("LOCAL ACCOUNT"), labels.index("VS")
            return [{"date": r[0].date().isoformat(), "la": _num(r[la]), "vs": _num(r[vs])}
                    for r in rows[i + 1:] if isinstance(r[0], datetime)]
    return []


def best_days(days):
    """Un renglón por día: el que tiene las empresas separadas y, entre esos, el de la hoja del mes."""
    best = {}
    for d in days:
        score = (bool((d["kia"] or 0) or (d["toyota"] or 0)), d.get("main", True))
        current = best.get(d["date"])
        if current is None or score >= current[0]:
            best[d["date"]] = (score, d)
    return {k: v[1] for k, v in best.items()}


# ---------------------------------------------------------------------------
# Cuenta corriente
# ---------------------------------------------------------------------------

def payment_groups(payments):
    """
    Junta los cobros de un mismo pago (mismo día, N° de transacción
    seguidos): un pago de $4,153.67 son cinco cobros. Cada grupo tiene su
    empresa si todos sus cobros la tienen.
    """
    groups = []
    for p in sorted(payments, key=lambda p: (p["date"], int(p["transc"]) if p["transc"].isdigit() else 0)):
        number = int(p["transc"]) if p["transc"].isdigit() else None
        last = groups[-1] if groups else None
        if (last and last["date"] == p["date"] and number is not None and last["last_transc"] is not None
                and 0 < number - last["last_transc"] <= 3 and last["company"] == p["company"]):
            last["items"].append(p)
            last["last_transc"] = number
        else:
            groups.append({"date": p["date"], "items": [p], "last_transc": number, "company": p["company"]})
    for g in groups:
        g["amount"] = round(sum(p["amount"] for p in g["items"]), 2)
        g["ids"] = [p.get("id") for p in g["items"] if p.get("id") is not None]
        g["transcs"] = [p["transc"] for p in g["items"]]
        g["pos_only"] = any(p.get("pos_only") for p in g["items"])
    return groups


# Días en que el POS cobró LOCAL ACCT (VS) y no hay recibo anotado: el pago
# existe (lo cobró la caja y lo depositó J.H.) aunque falte el papel. Se
# suman como "cobrado sin recibo", asignables a una empresa (pedido del
# usuario, 2026-10-07: abril, mayo y septiembre 2026 tenían $14,184.73 así,
# que eran de Kia). Un recibo anotado con otra fecha que el cobro del POS
# (hasta POS_MATCH_DAYS de distancia) se compensa y no cuenta dos veces.
POS_ONLY = "POS"
POS_MATCH_DAYS = 7


def pos_only_payments(payments, pos, assignments=None):
    """Lo cobrado en el POS sin recibo, un pago por grupo de días cercanos; y lo anotado de más."""
    assignments = assignments or {}
    paid = {}
    for p in payments:
        if p["company"] != VOIDED:
            paid[p["date"]] = round(paid.get(p["date"], 0.0) + p["amount"], 2)
    diffs = []
    for day in sorted(set(paid) | {d for d, v in pos.items() if v.get("vs")}):
        if day not in pos or pos[day].get("vs") is None:
            continue
        diff = round((pos[day].get("vs") or 0.0) - paid.get(day, 0.0), 2)
        if abs(diff) > TOLERANCE:
            diffs.append((day, diff))
    clusters = []
    for day, diff in diffs:
        if clusters and (date.fromisoformat(day) - date.fromisoformat(clusters[-1][-1][0])).days <= POS_MATCH_DAYS:
            clusters[-1].append((day, diff))
        else:
            clusters.append([(day, diff)])
    extra, over = [], []
    for cluster in clusters:
        net = round(sum(d for _, d in cluster), 2)
        day = max(cluster, key=lambda x: x[1])[0]
        if net > TOLERANCE:
            company = assignments.get(day) or (KIA if day < TOYOTA_START else None)
            extra.append({"date": day, "transc": POS_ONLY, "amount": net, "company": company,
                          "pos_only": True})
        elif net < -TOLERANCE:
            over.append({"date": min(cluster, key=lambda x: x[1])[0], "amount": -net})
    return extra, over


def build_ledger(charges, payments, pos, today=None):
    """
    charges: {fecha: {kia, kia_gal, toyota, toyota_gal, la}}; payments:
    [{date, transc, amount, company}] (company None = sin asignar); pos:
    {fecha: {la, vs}} (Local Account y LOCAL ACCT cobrado del POS). Devuelve
    los saldos por empresa, el resumen por mes y lo que hay que revisar.
    """
    today = today or date.today()
    events = {}  # mes -> empresa -> {...}
    by_day = {}

    def bucket(month, company):
        return events.setdefault(month, {}).setdefault(company, {
            "vouchers": 0.0, "gallons": 0.0, "charged": 0.0, "paid": 0.0})

    for day, c in charges.items():
        month = day[:7]
        row = by_day.setdefault(day, {"date": day})
        for company, amount_key, gal_key in ((KIA, "kia", "kia_gal"), (TOYOTA, "toyota", "toyota_gal")):
            amount, gallons = c.get(amount_key) or 0.0, c.get(gal_key) or 0.0
            charged = to_charge(amount, gallons) if amount else 0.0
            b = bucket(month, company)
            b["vouchers"] += amount
            b["gallons"] += gallons
            b["charged"] += charged
            row[company] = {"vouchers": amount, "gallons": gallons, "charged": charged}
        row["la_excel"] = c.get("la")
    unassigned = []
    payments = [p for p in payments if p["company"] != VOIDED]
    for p in payments:
        row = by_day.setdefault(p["date"], {"date": p["date"]})
        row.setdefault("paid", {}).setdefault(p["company"] or "?", 0.0)
        row["paid"][p["company"] or "?"] = round(row["paid"][p["company"] or "?"] + p["amount"], 2)
        if p["company"]:
            bucket(p["date"][:7], p["company"])["paid"] += p["amount"]
        else:
            unassigned.append(p)
    for day, values in pos.items():
        row = by_day.setdefault(day, {"date": day})
        row["la_pos"], row["vs"] = values.get("la"), values.get("vs")

    months = sorted(events)
    balances = {company: 0.0 for company in COMPANIES}
    monthly = []
    for month in months:
        entry = {"month": month, "companies": {}}
        for company in COMPANIES:
            b = events[month].get(company) or {"vouchers": 0.0, "gallons": 0.0, "charged": 0.0, "paid": 0.0}
            opening = balances[company]
            balances[company] = round(opening + b["charged"] - b["paid"], 2)
            entry["companies"][company] = {
                "opening": round(opening, 2), "vouchers": round(b["vouchers"], 2), "gallons": round(b["gallons"], 2),
                "charged": round(b["charged"], 2), "paid": round(b["paid"], 2), "closing": balances[company]}
        entry["unassigned"] = round(sum(p["amount"] for p in unassigned if p["date"][:7] == month), 2)
        monthly.append(entry)

    # Días: Local Account contra los vales, y lo cobrado en el POS contra los pagos anotados.
    days = []
    running = {company: 0.0 for company in COMPANIES}
    for day in sorted(by_day):
        row = by_day[day]
        # Saldo de cada una al cierre del día (pedido del usuario: "con qué
        # saldo quedan día a día"). Los pagos sin empresa no lo mueven.
        for company in COMPANIES:
            running[company] = round(running[company] + (row.get(company) or {}).get("charged", 0.0)
                                     - (row.get("paid") or {}).get(company, 0.0), 2)
        row["balance"] = dict(running)
        vouchers = sum((row.get(c) or {}).get("vouchers", 0.0) for c in COMPANIES)
        la = row.get("la_pos") if row.get("la_pos") is not None else row.get("la_excel")
        row["vouchers_total"] = round(vouchers, 2)
        row["la"] = la
        row["la_diff"] = None if la is None or day not in charges else round(la - vouchers, 2)
        paid = round(sum((row.get("paid") or {}).values()), 2)
        row["paid_total"] = paid
        row["vs_diff"] = None if row.get("vs") is None else round(row["vs"] - paid, 2)
        days.append(row)

    last_payment = {}
    for p in payments:
        if p["company"] and p["date"] > last_payment.get(p["company"], ""):
            last_payment[p["company"]] = p["date"]
    issues = []
    status = {}
    for company in COMPANIES:
        last = last_payment.get(company)
        days_since = (today - date.fromisoformat(last)).days if last else None
        status[company] = {"balance": balances[company], "last_payment": last, "days_since": days_since,
                           "last_amount": round(sum(p["amount"] for p in payments
                                                    if p["company"] == company and p["date"] == last), 2) if last else None,
                           "late": days_since is not None and days_since > LATE_PAYMENT_DAYS}
        if status[company]["late"]:
            issues.append(f"{company} no paga hace {days_since} días (último pago: {last[8:10]}/{last[5:7]}/{last[:4]}).")
    if unassigned:
        issues.append(f"{len(payment_groups(unassigned))} pago(s) sin empresa (${sum(p['amount'] for p in unassigned):,.2f}): "
                      f"asignalos abajo para que el saldo de cada una sea el real.")
    return {
        "monthly": monthly,
        "days": days,
        "status": status,
        "total_balance": round(sum(balances.values()) - sum(p["amount"] for p in unassigned), 2),
        "unassigned_groups": payment_groups(unassigned),
        "issues": issues,
    }


def month_days(ledger, year, month):
    """Los días del mes (todos, aunque no tengan movimientos)."""
    by_date = {d["date"]: d for d in ledger["days"]}
    first = f"{year:04d}-{month:02d}-01"
    previous = [d for d in ledger["days"] if d["date"] < first]
    balance = dict(previous[-1]["balance"]) if previous else {company: 0.0 for company in COMPANIES}
    rows = []
    for day in range(1, calendar.monthrange(year, month)[1] + 1):
        key = f"{year:04d}-{month:02d}-{day:02d}"
        row = by_date.get(key) or {"date": key}
        balance = row.get("balance") or balance
        rows.append(dict(row, balance=balance))
    return rows
