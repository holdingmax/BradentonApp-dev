"""
Módulo "Gettel -- Pagos de Cupones" (2026-09-19, pedido explícito del
usuario): reconciliación mensual de lo que hay que pagarle a Gettel
(Total del mes ya cargado en el módulo de cupones diarios de Gettel/
Toyota + lo que quedó pendiente del mes anterior) contra lo efectivamente
cobrado (los pagos/cupones que Gettel realmente envió y se cargaron a
mano). Reproduce la hoja real "Pago Cupones" del Excel Cierre --
decodificada con openpyxl leyendo un ejemplo real (`hoja_de_pagos_
gettel.xlsx`), no adivinada.

Reemplaza la lectura automática de PDF de pagos (gettel_toyota_parser.
extract_pago_batch_from_pdf/process_gettel_pagos -- la herramienta vieja
de /gettel/pagos, que escribe directo sobre el Excel Cierre real) --
pedido explícito del usuario: "no esta leyendo bien los pagos que subi en
pdf". La carga de cada pago ahora es a mano, un renglón por cupón/
transacción -- mismo criterio "Guardar en base" ya elegido para
Combustible/Físico (nunca se escribe sobre ningún Excel real).

Fórmula real (decodificada de hoja_de_pagos_gettel.xlsx):
    Total del mes          = gettel_db.get_month_gettel_amount(year, month)
                              (F4 real: `='[1]Gettel-Toyota MM.YYYY'!E40` --
                              el mismo dato que ya alimenta la categoría
                              "Gettel" de Ventas por Departamento).
    Pendiente Mes Anterior  = F7 real -- en la hoja real SIEMPRE es un
                              valor pegado a mano mes a mes; acá se
                              automatiza encadenando UN solo mes atrás
                              (mismo patrón que caja._resolve_opening_
                              balance/fisico._resolve_initial): un
                              override manual de ESTE mes manda siempre;
                              si no hay, se encadena el "Pendiente que
                              pasa al mes siguiente" ya calculado del mes
                              anterior (nunca más de un mes atrás); si
                              tampoco hay nada ahí, 0.
    Total a Pagar           = Total del mes + Pendiente Mes Anterior
                              (F32 real: `=SUM(F4:F8)` -- suma los dos
                              bloques merged F4:F5 y F7:F8).
    Cobrado                 = SUMA de los cupones/pagos cargados este mes
                              (D32 real: `=SUM(D4:D30)`).
    Diferencia               = Cobrado - Total a Pagar (E32 real: `=+D32-F32`).
        Diferencia >= 0   -> "a acreditarse en 48 hs" (D33 real) -- sobró
                              lo cobrado, no queda nada pendiente para el
                              mes que viene.
        Diferencia < 0    -> falta cobrar -- el valor absoluto pasa a ser
                              el Pendiente del mes SIGUIENTE.
"""

import gettel_db


def _resolve_pendiente_anterior(year, month, _chain=True):
    """(valor, fuente) del Pendiente Mes Anterior -- mismo patrón que caja._resolve_opening_balance/fisico._resolve_initial."""
    settings = gettel_db.get_pago_month_settings(year, month)
    if settings and settings.get("pendiente_anterior_override") is not None:
        return settings["pendiente_anterior_override"], "manual"
    if not _chain:
        return 0.0, "default"

    prev_year, prev_month = (year - 1, 12) if month == 1 else (year, month - 1)
    prev_report = build_month_report(prev_year, prev_month, _chain=False)
    if prev_report["has_any_data"]:
        return prev_report["pendiente_siguiente"], "prev_month_computed"
    return 0.0, "default"


def build_month_report(year, month, _chain=True):
    pagos = gettel_db.get_month_pagos(year, month)
    settings = gettel_db.get_pago_month_settings(year, month)

    total_mes = round(gettel_db.get_month_gettel_amount(year, month), 2)
    pendiente_anterior, pendiente_source = _resolve_pendiente_anterior(year, month, _chain=_chain)
    pendiente_anterior = round(pendiente_anterior, 2)

    total_a_pagar = round(total_mes + pendiente_anterior, 2)
    cobrado = round(sum(p.get("total_cupon") or 0.0 for p in pagos), 2)
    diferencia = round(cobrado - total_a_pagar, 2)

    if diferencia >= 0:
        credito_48hs = diferencia
        pendiente_siguiente = 0.0
    else:
        credito_48hs = 0.0
        pendiente_siguiente = round(-diferencia, 2)

    has_any_data = bool(pagos) or (settings is not None) or bool(total_mes)

    return {
        "year": year,
        "month": month,
        "has_any_data": has_any_data,
        "pagos": pagos,
        "total_mes": total_mes,
        "pendiente_anterior": pendiente_anterior,
        "pendiente_source": pendiente_source,
        "total_a_pagar": total_a_pagar,
        "cobrado": cobrado,
        "diferencia": diferencia,
        "credito_48hs": credito_48hs,
        "pendiente_siguiente": pendiente_siguiente,
    }


def grouped_pagos(pagos):
    """
    Agrupa filas consecutivas que comparten Fecha+Pago N° -- para el
    rowspan de la tabla en pantalla y las celdas unificadas del export a
    Excel/PDF (mismo criterio visual que la hoja real, donde un pago
    puede traer varias transacciones/cupones abajo, todas con la misma
    Fecha y el mismo Pago N°).
    """
    groups = []
    current = None
    for pago in pagos:
        key = (pago.get("fecha"), pago.get("pago_n"))
        if current is not None and current["key"] == key:
            current["rows"].append(pago)
        else:
            current = {"key": key, "rows": [pago]}
            groups.append(current)
    return groups


def get_available_years():
    years = set(gettel_db.get_pago_years()) | set(gettel_db.get_pago_settings_years())
    return sorted(years)


def _fmt_money(value):
    return "—" if value is None else "{:,.2f}".format(value)


def build_pagos_export_workbook(report, year, month, dest_path):
    """
    Excel NUEVO (nunca toca ningún archivo real) con el cuadro de Pagos de
    Gettel -- mismas filas/formatos/colores/negritas/celdas unificadas/
    títulos/centrados que la hoja real "Pago Cupones" (`hoja_de_pagos_
    gettel.xlsx`, decodificada con openpyxl) -- pedido explícito del
    usuario. Mismos colores hex que ya usa build_caja_export_workbook
    (gris D9D9D9 para encabezados/totales, verde 92D050 para el bloque de
    Pendiente Mes Anterior -- confirmados 1:1 contra el archivo real).

    Total del mes/Pendiente Anterior/Total a Pagar/Cobrado/Diferencia se
    exportan siempre como VALOR, nunca como fórmula de Excel -- mismo
    criterio que build_caja_export_workbook, para que el número que ve
    acá sea siempre el mismo que ya calculó la página.
    """
    import openpyxl
    from openpyxl.styles import Alignment as XlAlignment, Border as XlBorder, Font as XlFont, PatternFill, Side as XlSide

    HEADER_GRAY = PatternFill("solid", fgColor="FFD9D9D9")
    DATA_GREEN = PatternFill("solid", fgColor="FF92D050")
    MONEY_FMT = '"$" #,##0.00'
    THIN_SIDE = XlSide(style="thin", color="FF000000")
    THIN_BORDER = XlBorder(left=THIN_SIDE, right=THIN_SIDE, top=THIN_SIDE, bottom=THIN_SIDE)
    CENTER = XlAlignment(horizontal="center", vertical="center")
    CENTER_WRAP = XlAlignment(horizontal="center", vertical="center", wrap_text=True)

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Pago Cupones"

    # Título -- mismo texto real (adaptado al mes/año pedido), fila 1 unificada.
    title_cell = sheet.cell(row=1, column=1, value=f"CONTROL - {month:02d}/{year} (CUPONES QUE ENVIA RICK A FINAL DE MES)")
    title_cell.font = XlFont(name="Segoe UI", bold=True, size=14)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=8)
    sheet.row_dimensions[1].height = 21

    # Encabezado de la tabla de cupones (fila 3, igual que la hoja real).
    headers = ["Fecha", "Pago N°", "Transc N°", "Total Cupon"]
    for col, text in enumerate(headers, start=1):
        cell = sheet.cell(row=3, column=col, value=text)
        cell.font = XlFont(bold=True)
        cell.fill = HEADER_GRAY
        cell.alignment = CENTER_WRAP
        cell.border = THIN_BORDER

    # Bloque de totales a la derecha, arriba (F4:F5 Total del mes, G4:G5
    # su período; F7:F8 Pendiente Mes Anterior con fondo verde, G7:G8 su
    # período, H7:J7 la etiqueta) -- mismo layout real, con "leve
    # separación" real (columna E vacía entre D y F).
    f3 = sheet.cell(row=3, column=6, value="Total")
    f3.font = XlFont(bold=True)
    f3.fill = HEADER_GRAY
    f3.alignment = CENTER
    f3.border = THIN_BORDER

    sheet.merge_cells(start_row=4, start_column=6, end_row=5, end_column=6)
    f4 = sheet.cell(row=4, column=6, value=report["total_mes"])
    f4.font = XlFont(bold=True)
    f4.number_format = MONEY_FMT
    f4.alignment = CENTER
    f4.border = THIN_BORDER
    sheet.merge_cells(start_row=4, start_column=7, end_row=5, end_column=7)
    g4 = sheet.cell(row=4, column=7, value=f"Total del mes de Gettel ({month:02d}/{year})")
    g4.font = XlFont(bold=True)
    g4.alignment = CENTER_WRAP
    g4.border = THIN_BORDER

    for col in (6, 7):
        sheet.merge_cells(start_row=7, start_column=col, end_row=8, end_column=col)
    f7 = sheet.cell(row=7, column=6, value=report["pendiente_anterior"])
    f7.font = XlFont(bold=True)
    f7.number_format = MONEY_FMT
    f7.fill = DATA_GREEN
    f7.alignment = CENTER
    f7.border = THIN_BORDER
    g7 = sheet.cell(row=7, column=7)
    g7.fill = DATA_GREEN
    g7.border = THIN_BORDER
    sheet.merge_cells(start_row=7, start_column=8, end_row=7, end_column=10)
    h7 = sheet.cell(row=7, column=8, value="(PENDIENTE DEL MES ANTERIOR)")
    h7.alignment = XlAlignment(horizontal="left", vertical="center")

    # Filas de cupones -- Fecha/Pago N° con celda unificada por grupo
    # (mismas filas consecutivas que comparten pago -- ver grouped_pagos).
    row_cursor = 4
    for group in grouped_pagos(report["pagos"]):
        start_row = row_cursor
        for pago in group["rows"]:
            sheet.cell(row=row_cursor, column=3, value=pago.get("transc_n") or "").alignment = CENTER
            sheet.cell(row=row_cursor, column=3).border = THIN_BORDER
            amount_cell = sheet.cell(row=row_cursor, column=4, value=pago.get("total_cupon"))
            amount_cell.number_format = MONEY_FMT
            amount_cell.alignment = CENTER
            amount_cell.border = THIN_BORDER
            for col in (1, 2):
                sheet.cell(row=row_cursor, column=col).border = THIN_BORDER
            row_cursor += 1
        end_row = row_cursor - 1
        fecha_value = group["rows"][0].get("fecha")
        pago_n_value = group["rows"][0].get("pago_n")
        if end_row > start_row:
            sheet.merge_cells(start_row=start_row, start_column=1, end_row=end_row, end_column=1)
            sheet.merge_cells(start_row=start_row, start_column=2, end_row=end_row, end_column=2)
        fecha_cell = sheet.cell(row=start_row, column=1, value=fecha_value)
        fecha_cell.alignment = CENTER
        fecha_cell.font = XlFont(bold=True)
        pago_cell = sheet.cell(row=start_row, column=2, value=pago_n_value)
        pago_cell.alignment = CENTER
        pago_cell.font = XlFont(bold=True)

    # Fila de Cobrado/Diferencia/Total a Pagar -- mismas columnas/fórmulas
    # reales (C32 COBRADO, D32 suma, E32 diferencia, F32 total a pagar,
    # D33 "a acreditarse en 48 hs") pero como VALOR ya calculado.
    totals_row = max(row_cursor, 32)
    c_label = sheet.cell(row=totals_row, column=3, value="COBRADO")
    c_label.font = XlFont(bold=True)
    c_label.alignment = XlAlignment(horizontal="right")
    d_total = sheet.cell(row=totals_row, column=4, value=report["cobrado"])
    d_total.font = XlFont(bold=True)
    d_total.number_format = MONEY_FMT
    d_total.fill = HEADER_GRAY
    d_total.alignment = CENTER
    d_total.border = THIN_BORDER
    e_dif = sheet.cell(row=totals_row, column=5, value=report["diferencia"])
    e_dif.font = XlFont(bold=True)
    e_dif.number_format = MONEY_FMT
    e_dif.alignment = CENTER
    f_pagar = sheet.cell(row=totals_row, column=6, value=report["total_a_pagar"])
    f_pagar.font = XlFont(bold=True)
    f_pagar.number_format = MONEY_FMT
    f_pagar.fill = HEADER_GRAY
    f_pagar.alignment = CENTER
    f_pagar.border = THIN_BORDER

    note_row = totals_row + 1
    note_label = sheet.cell(
        row=note_row, column=4,
        value="a acreditarse en 48 hs" if report["diferencia"] >= 0 else "pasa a Pendiente del mes siguiente",
    )
    note_label.font = XlFont(bold=True, italic=True)
    note_label.alignment = XlAlignment(horizontal="right")

    for col_letter, width in zip("ABCDEFGHIJ", (12, 9, 12, 13, 3, 13, 22, 14, 8, 8)):
        sheet.column_dimensions[col_letter].width = width

    workbook.save(dest_path)
    return dest_path


def build_pagos_export_pdf(report, year, month, dest_path):
    """
    Versión PDF (básica, con los mismos colores clave) del cuadro de
    Pagos de Gettel -- pedido explícito del usuario: "quiero que se vea
    el excel... y su PDF de al lado del cuadro, el de reporte todavia no
    lo hacemos" (es decir, un PDF propio de este cuadro, NO el PDF
    resumido del módulo "Reportes"). Reusa pdf_export.build_multi_
    section_pdf -- dos secciones: la tabla de cupones y el resumen de
    totales, mismo criterio Detalle/Total ya usado en los resúmenes de
    Reportes (ver caja.build_caja_pdf_resumen).
    """
    from pdf_export import build_multi_section_pdf

    headers = ["Fecha", "Pago N°", "Transc N°", "Total Cupón"]
    table_rows = [
        [pago.get("fecha") or "—", str(pago.get("pago_n")) if pago.get("pago_n") is not None else "—",
         pago.get("transc_n") or "—", _fmt_money(pago.get("total_cupon"))]
        for pago in report["pagos"]
    ]
    table_rows.append(["", "", "COBRADO", _fmt_money(report["cobrado"])])

    totals_rows = [
        ["Total del mes (Gettel)", _fmt_money(report["total_mes"])],
        ["Pendiente Mes Anterior", _fmt_money(report["pendiente_anterior"])],
        ["Total a Pagar", _fmt_money(report["total_a_pagar"])],
        ["Cobrado", _fmt_money(report["cobrado"])],
        ["Diferencia", _fmt_money(report["diferencia"])],
        [
            "A acreditarse en 48 hs" if report["diferencia"] >= 0 else "Pendiente Mes Siguiente",
            _fmt_money(report["credito_48hs"] if report["diferencia"] >= 0 else report["pendiente_siguiente"]),
        ],
    ]

    sections = [
        {
            "heading": "Cupones cargados",
            "headers": headers,
            "rows": table_rows,
            "col_widths_mm": [35, 25, 50, 35],
            "bold_last_row": True,
        },
        {
            "heading": "Resumen del mes",
            "headers": ["Detalle", "Total"],
            "rows": totals_rows,
            "col_widths_mm": [110, 80],
            "bold_last_row": True,
        },
    ]

    return build_multi_section_pdf(
        dest_path,
        f"Gettel — Pago Cupones — {month:02d}/{year}",
        sections,
        company_header=True,
    )
