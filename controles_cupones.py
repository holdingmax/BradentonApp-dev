"""
Controles: Cupones -- tercer módulo de la sección Controles.

Cruza el saldo final del "Mayor" (extracto contable) de la cuenta J.H.
Williams - Recaudación a Liquidar -- un .xls/.xlsx que el usuario exporta
del sistema contable para un período dado -- contra la suma de los
cupones que todavía no se aplicaron a ningún EFT en la hoja Cupones del
Excel "Aplicacion TC y EFT -": los que no tienen fórmula en la columna F
(el cruce contra Cta Cte J.H.Williams que arma cupones_append.py) ni
tienen relleno de color -- en la práctica las dos señales del mismo
estado (toda fila ya cruzada contra un EFT queda con las dos a la vez),
se piden ambas por las dudas de que alguna quede desincronizada.

Es de solo lectura: nunca escribe nada en ningún Excel ni genera un
archivo para descargar -- el resultado se muestra en pantalla, según lo
ya decidido para toda la sección Controles.

Sensible al momento en que se corre (confirmado con el usuario 2026-09-08):
la hoja Cupones es un archivo único y acumulativo -- no hay "un archivo de
Cupones por mes" -- así que el lado de cupones pendientes de este control
refleja el estado de la hoja AL MOMENTO DE LA CARGA, no el estado a fin
del mes que cubre el Mayor. Correrlo con un Excel de Cupones de un mes
posterior no sirve: para entonces los cupones que estaban pendientes a
fin del mes del Mayor ya se van a haber aplicado a algún EFT posterior,
así que la diferencia da mucho más grande de lo real. Hay que usarlo con
un Excel de Cupones de cerca de la fecha de cierre de ese mes.
"""

import os
import re
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook

from cupones_append import (
    CUPONES_COL_COUPON,
    CUPONES_COL_DATE,
    CUPONES_COL_EFT_FORMULA,
    CUPONES_COL_FEES,
    CUPONES_COL_GROSS,
    CUPONES_COL_NET,
    CUPONES_SCAN_START_ROW,
    CUPONES_SHEET,
    DDC_COUPON_PATTERN,
    find_last_cupones_row,
)

# A diferencia de los otros controles (±$0.01, redondeo), acá siempre hay
# alguna diferencia real -- movimientos por el mismo estilo de "COBRO
# GETTEL-KIA" que pasan por esta misma cuenta sin ser cupones, cupones
# recién resueltos entre el cierre del mes y el momento de la carga, etc.
# Margen pedido explícitamente por el usuario (2026-09-08, según lo que
# le indicó su compañera de trabajo): hasta $15.000 de diferencia es
# normal; por encima, vale la pena una alerta -- nunca un error duro
# como el resto de los controles, porque una diferencia acá no implica
# necesariamente que algo esté mal cargado.
TOLERANCE = 15000.0

# Columnas que el motor de cupones_append.py pinta cuando una fila queda
# resuelta -- B/C/D/E, ya sea por el color de "grupo partido" (solo en B,
# ver SPLIT_GROUP_FILL_COLORS) o el rojo rosado de "posible duplicado"
# (solo en C/D/E, ver DUPLICATE_COUPON_FILL_COLOR). Se revisan las cuatro
# para capturar cualquiera de las dos señales -- alcanza con que UNA tenga
# relleno para considerar la fila "pintada".
_PAINTED_CHECK_COLUMNS = (CUPONES_COL_COUPON, CUPONES_COL_GROSS, CUPONES_COL_FEES, CUPONES_COL_NET)

_PERIOD_LINE_RE = re.compile(r"desde:\s*([\d/-]+)\s*hasta:\s*([\d/-]+)", re.IGNORECASE)
_MAYOR_DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y")


def _parse_mayor_date(text):
    text = text.strip()
    for fmt in _MAYOR_DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _read_mayor_ledger(path):
    """
    Lee el Mayor (.xls o .xlsx, pandas elige el motor solo según la
    extensión) de la cuenta Recaudación a Liquidar.

    Busca la línea "Desde: ... Hasta: ..." (en cualquier fila -- el
    reporte siempre la trae en algún lugar de las primeras filas, junto
    con el título) y la fila de encabezado real por su propio contenido
    ("Fecha"/"Saldo" como texto de columna, nunca una posición fija) en
    vez de asumir un número de fila -- el reporte contable no promete la
    misma cantidad de filas de título en cada exportación. El saldo final
    es el último valor numérico real de la columna Saldo -- la fila de
    totales al pie no tiene Saldo propio, así que queda afuera sola.

    Devuelve (period_from, period_to, saldo_final).
    """
    df = pd.read_excel(path, header=None)

    period_from = period_to = None
    header_row_index = None
    saldo_col_index = None

    for idx in range(len(df)):
        cells = ["" if pd.isna(v) else str(v).strip() for v in df.iloc[idx]]
        if period_from is None:
            joined = " ".join(c for c in cells if c)
            match = _PERIOD_LINE_RE.search(joined)
            if match:
                period_from = _parse_mayor_date(match.group(1))
                period_to = _parse_mayor_date(match.group(2))
        if header_row_index is None:
            lowered = [c.lower() for c in cells]
            if "fecha" in lowered and "saldo" in lowered:
                header_row_index = idx
                saldo_col_index = lowered.index("saldo")

    if header_row_index is None:
        raise ValueError('No se encontró la fila de encabezado ("Fecha" / "Saldo") en el Mayor.')
    if period_from is None or period_to is None:
        raise ValueError('No se encontró la línea "Desde: ... Hasta: ..." en el Mayor.')

    saldo_final = None
    for idx in range(header_row_index + 1, len(df)):
        value = df.iat[idx, saldo_col_index]
        if pd.notna(value) and isinstance(value, (int, float)):
            saldo_final = float(value)

    if saldo_final is None:
        raise ValueError('No se encontró ningún valor de Saldo en las filas del Mayor.')

    return period_from, period_to, saldo_final


def _format_date(value):
    """
    Formatea una fecha para mostrar en el template -- tolera que la celda
    de Excel (o la fecha del Mayor) haya quedado como texto en vez de una
    fecha real (visto en la práctica en filas de Cupones cargadas a mano),
    en vez de asumir siempre un objeto date/datetime.
    """
    if value is None:
        return None
    if hasattr(value, "strftime"):
        return value.strftime("%d/%m/%Y")
    text = str(value).strip()
    return text or None


def _read_float(value):
    if value is None:
        return 0.0
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def _cell_is_painted(cell):
    fill = cell.fill
    return fill is not None and fill.patternType is not None


def _row_is_painted(sheet, row):
    return any(_cell_is_painted(sheet.cell(row=row, column=col)) for col in _PAINTED_CHECK_COLUMNS)


def _find_pending_coupons(sheet):
    """
    Cupones sin resolver: sin fórmula en F (columna EFT) Y sin relleno de
    color en ninguna de B/C/D/E.

    La columna B solo cuenta como cupón si matchea el patrón real
    ("DDC-####", mismo que usa cupones_append.py) -- algunos meses tienen,
    justo debajo del último cupón real, un cuadrito de reconciliación a
    mano del usuario (ej. "SI"/"EFT"/"VENTAS TC"/"VENTAS GETTEL"/"SF" en
    filas sueltas de la columna B) que de otra forma se cuela como si
    fueran cupones más -- confirmado contra un Excel real de julio-2026.

    También se descarta una fila que tiene el N° de cupón tipeado en B
    pero Gross/Fees/Net completamente vacíos (sin relleno de color como
    sus vecinas tampoco) -- un hueco de carga, no un cupón real -- visto
    en el mismo Excel real (fila con "DDC-1078" sin ningún monto).
    """
    last_row = find_last_cupones_row(sheet)
    pending = []
    for row in range(CUPONES_SCAN_START_ROW, last_row + 1):
        coupon = sheet.cell(row=row, column=CUPONES_COL_COUPON).value
        if coupon is None or str(coupon).strip() == "":
            continue
        if not DDC_COUPON_PATTERN.fullmatch(str(coupon).strip()):
            continue
        gross_raw = sheet.cell(row=row, column=CUPONES_COL_GROSS).value
        fees_raw = sheet.cell(row=row, column=CUPONES_COL_FEES).value
        net_raw = sheet.cell(row=row, column=CUPONES_COL_NET).value
        if gross_raw is None and fees_raw is None and net_raw is None:
            continue
        eft_value = sheet.cell(row=row, column=CUPONES_COL_EFT_FORMULA).value
        has_eft_link = eft_value not in (None, "")
        if has_eft_link or _row_is_painted(sheet, row):
            continue
        pending.append(
            {
                "row": row,
                "date": sheet.cell(row=row, column=CUPONES_COL_DATE).value,
                "date_display": _format_date(sheet.cell(row=row, column=CUPONES_COL_DATE).value),
                "coupon": str(coupon).strip(),
                "gross": _read_float(sheet.cell(row=row, column=CUPONES_COL_GROSS).value),
                "fees": _read_float(sheet.cell(row=row, column=CUPONES_COL_FEES).value),
                "net": _read_float(sheet.cell(row=row, column=CUPONES_COL_NET).value),
            }
        )
    return pending


def check_cupones_pending(mayor_path, eft_excel_path):
    """
    Cruza el saldo final del Mayor de Recaudación a Liquidar contra la
    suma de los cupones sin aplicar a ningún EFT en la hoja Cupones.

    Devuelve un dict con el período del Mayor, el saldo final, el detalle
    y total de cupones pendientes, y el chequeo (diferencia + ok). Nunca
    escribe ni descarga nada.
    """
    mayor_path = os.path.abspath(str(mayor_path).strip())
    if not os.path.isfile(mayor_path):
        raise FileNotFoundError(f"Mayor de Recaudación a Liquidar no encontrado: {mayor_path}")

    eft_excel_path = os.path.abspath(str(eft_excel_path).strip())
    if not os.path.isfile(eft_excel_path):
        raise FileNotFoundError(f"Excel de Aplicacion TC y EFT no encontrado: {eft_excel_path}")
    extension = os.path.splitext(eft_excel_path)[1].lower()
    if extension not in {".xlsx", ".xlsm"}:
        raise ValueError("El Excel de Aplicacion TC y EFT debe ser .xlsx o .xlsm.")

    period_from, period_to, saldo_final = _read_mayor_ledger(mayor_path)

    workbook = load_workbook(eft_excel_path, data_only=False)
    try:
        if CUPONES_SHEET not in workbook.sheetnames:
            raise ValueError(f'Hoja "{CUPONES_SHEET}" no encontrada en el Excel de Aplicacion TC y EFT.')
        sheet = workbook[CUPONES_SHEET]
        pending_coupons = _find_pending_coupons(sheet)
    finally:
        workbook.close()

    pending_total_gross = round(sum(c["gross"] for c in pending_coupons), 2)
    pending_total_fees = round(sum(c["fees"] for c in pending_coupons), 2)
    pending_total_net = round(sum(c["net"] for c in pending_coupons), 2)

    diff = round(saldo_final - pending_total_gross, 2)

    return {
        "period_from": period_from,
        "period_to": period_to,
        "period_from_display": _format_date(period_from),
        "period_to_display": _format_date(period_to),
        "saldo_final": round(saldo_final, 2),
        "pending_coupons": pending_coupons,
        "pending_count": len(pending_coupons),
        "pending_total_gross": pending_total_gross,
        "pending_total_fees": pending_total_fees,
        "pending_total_net": pending_total_net,
        "diff": diff,
        "ok": abs(diff) <= TOLERANCE,
    }
