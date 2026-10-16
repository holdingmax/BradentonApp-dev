"""
Compara la app sobre SQLite y sobre PostgreSQL con los mismos datos reales
(pase a Postgres, punto 6 de PLAN_POSTGRES.md). Antes de cada commit de un
módulo, sus páginas, descargas y escrituras tienen que dar idénticas.

    python comparar_bases.py jh_facturas [--datos CARPETA]

- Toma dos copias de la carpeta de datos (por defecto, la de la app; los
  originales no se tocan): una para correr la app sobre SQLite y otra para
  la corrida sobre Postgres.
- Copia el módulo a Postgres con migrar_sqlite_a_postgres.py --vaciar, así
  que solo corre contra un Postgres de esta PC (DATABASE_URL apuntando a
  localhost, nunca a testing ni producción).
- En cada corrida abre las páginas del módulo con el test client de Flask
  (logueado como admin, sin contraseña), hace sus pruebas de escritura y
  vuelve a abrir las páginas. Después compara todo: el HTML, el texto de
  los PDF y las celdas de los Excel.
"""

import argparse
import difflib
import io
import json
import os
import re
import shutil
import subprocess
import sys
import tempfile
import time
from datetime import date

APP_DIR = os.path.dirname(os.path.abspath(__file__))
LOCAL_HOSTS = {"localhost", "127.0.0.1", "::1"}
CONFIG_FILES = ("users.json", "chase_master_rules.json", "chase_rules.json", "proveedores_dynamic_suppliers.json", "proveedores_pago_rules.json")
MONTHS = [(2026, 7), (2026, 8), (2026, 9), (2026, 10)]


def _months(path):
    return [f"{path}?year={y}&month={m}" for y, m in MONTHS]


# Páginas que leen cada módulo. Se agregan a medida que se pasa cada uno.
PAGES = {
    "jh_facturas": _months("/carga-datos/eft/varias") + ["/carga-datos/eft/proximos", "/controles/rapidos/eft", "/carga-datos/eft"],
    "ice_machine": _months("/carga-datos/depositos"),
    "reporte_mensual": [url for path in ("/carga-datos/reporte-mensual", "/carga-datos/reporte-diario", "/reporte/historial",
                                         "/reporte/store-info/historial", "/controles/cierre", "/controles/cierre/exportar/pdf",
                                         "/controles/cmv") for url in _months(path)],
    "cheques": ["/carga-datos/chase/cheques"],
    "control_cmv": _months("/controles/cmv"),
    "depositos": [url for path in ("/carga-datos/depositos", "/controles/depositos", "/carga-datos/chase") for url in _months(path)],
    "horas_trabajo": ["/carga-datos/horas-trabajo"] + _months("/carga-datos/horas-trabajo/historial"),
    "jh_mensual": [url for path in ("/carga-datos/eft/jh", "/carga-datos/eft/historial", "/carga-datos/eft/cupones/historial",
                                    "/carga-datos/reportes/eft/pdf", "/controles/tarjetas", "/controles/cierre",
                                    "/controles/rapidos/tarjetas") for url in _months(path)]
                  + ["/carga-datos/eft/proximos", "/controles/rapidos/eft"],
    "documents": ["/buscar/documentos?q=close"],  # /documentos/... redirige mientras el guardado esté apagado
    "fisico": [url for path in ("/fisico", "/fisico/exportar", "/fisico/exportar/pdf") for url in _months(path)]
              + ["/carga-datos/eft/proximos", "/controles/rapidos/eft"],
    "caja": [f"/carga-datos/caja?year=2026&month={m}" for m in range(1, 7)]
            + [url for path in ("/carga-datos/caja", "/carga-datos/caja/exportar", "/carga-datos/caja/exportar/pdf",
                                "/carga-datos/reportes/caja/pdf", "/controles/caja", "/controles/rapidos/caja") for url in _months(path)]
            + ["/carga-datos/reportes"],
    "gettel_toyota": [url for path in ("/carga-datos/gettel", "/carga-datos/gettel/historial", "/carga-datos/reportes/gettel/excel",
                                       "/carga-datos/reportes/gettel/pdf", "/controles/kia-toyota", "/controles/rapidos/kia")
                      for url in _months(path)],
    "lottery": ["/carga-datos/lottery", "/carga-datos/lottery/dia/2026-09-15", "/carga-datos/lottery/dia/2026-10-04",
                "/carga-datos/lottery/dia/2026-07-31"]
               + [url for path in ("/carga-datos/lottery/historial", "/carga-datos/lottery/exportar", "/carga-datos/lottery/exportar/pdf",
                                   "/carga-datos/lottery/mensual", "/carga-datos/lottery/resumen-mensual", "/carga-datos/reportes/lottery/pdf",
                                   "/controles/lottery", "/controles/caja", "/carga-datos/caja") for url in _months(path)],
    # /reporte/dia/<d>/pdf redirige: abre la copia del PDF, que ya no se guarda.
    "reportes_diarios": [f"/reporte/dia/{d}" for d in ("2026-07-27", "2026-08-19", "2026-09-15", "2026-10-04")]
                        + ["/carga-datos/eft/proximos", "/carga-datos/proyecciones", "/buscar/documentos?q=close"]
                        + [url for path in ("/carga-datos/reporte-diario", "/reporte/historial", "/reporte/store-info/historial",
                                            "/reporte/store-info/exportar", "/reporte/store-info/exportar/pdf",
                                            "/carga-datos/reportes/reporte-diario/pdf", "/controles/tarjetas", "/controles/cierre",
                                            "/controles/cierre/exportar", "/controles/cierre/exportar/pdf", "/controles/cmv",
                                            "/controles/kia-toyota", "/carga-datos/caja", "/controles/rapidos/tarjetas",
                                            "/controles/rapidos/caja") for url in _months(path)],
    "cuenta_kia_toyota": ["/carga-datos/gettel"]
                         + [url for path in ("/controles/kia-toyota", "/controles/rapidos/kia", "/carga-datos/gettel") for url in _months(path)],
    "eft": ["/carga-datos/eft/proximos", "/controles/rapidos/eft", "/carga-datos/reportes"]
           + [url for path in ("/carga-datos/eft", "/carga-datos/eft/historial", "/carga-datos/eft/cupones/historial", "/carga-datos/eft/jh",
                               "/carga-datos/eft/varias", "/carga-datos/reportes/eft/pdf", "/controles/tarjetas", "/controles/cierre",
                               "/controles/rapidos/tarjetas") for url in _months(path)],
    "chase": ["/carga-datos/chase", "/carga-datos/chase/cheques", "/proveedores", "/carga-datos/proveedores/guardado",
              "/carga-datos/proveedores/guardado/ht_hackney", "/carga-datos/proveedores/guardado/fpl", "/carga-datos/eft/proximos",
              "/controles/rapidos/eft"]
             + [url for path in ("/carga-datos/chase/historial", "/carga-datos/chase/exportar", "/carga-datos/reportes/chase/pdf",
                                 "/controles/depositos", "/carga-datos/caja", "/controles/lottery", "/controles/tarjetas",
                                 "/controles/cierre", "/controles/rapidos/caja") for url in _months(path)],
    "cmv": ["/carga-datos/cmv", "/carga-datos/cmv/costo/historial", "/carga-datos/cmv/costo/comparar", "/controles/ventas-costos",
            "/controles/productos/lista"]
           + [url for path in ("/carga-datos/cmv/ventas/historial", "/controles/cmv") for url in _months(path)],
    "proveedores": ["/carga-datos/proveedores/guardado", "/controles/productos", "/controles/productos/lista", "/controles/ventas-costos",
                    "/carga-datos/chase/cheques", "/controles/productos/ht_hackney", "/controles/productos/cec",
                    "/controles/productos/ht_hackney/2026-10-09/916761", "/controles/productos/ht_hackney/2026-10-09/916761/cambios.pdf",
                    "/controles/productos/ht_hackney/2026-10-09/916761/cambios.xlsx", "/controles/productos/cec/2026-10-06/1900656"]
                   + [f"/carga-datos/proveedores/guardado/{key}" for key in ("ht_hackney", "airgas", "fpl", "manatee_county", "gce", "cec",
                                                                             "coca", "bimbo", "red_bull", "kooler_ice")]
                   + [url for path in ("/carga-datos/proveedores/historial", "/carga-datos/reportes/proveedores/pdf") for url in _months(path)],
    # Usuarios, reglas, botonera y Sale Tax (config_store.py; en Postgres, el esquema config).
    "config": ["/admin/users", "/perfil/password", "/carga-datos/chase", "/carga-datos/proveedores/reglas", "/proveedores", "/proveedores/nuevo"]
              + [url for path in ("/controles/rapidos/caja", "/controles/cierre", "/carga-datos/chase/historial") for url in _months(path)],
}


def _escenario_jh_facturas():
    """Altas, repetidos, conflicto y baja de facturas varias (datos de prueba,
    solo en las copias)."""
    import jh_facturas_db

    invoice = {"invoice_number": "SI-999901", "invoice_date": date(2026, 10, 3), "due_date": date(2026, 10, 6),
               "amount": 131.004, "kind": "Mako", "description": "Prueba del comparador", "detail": "con 'comillas' y 100%"}
    results = [jh_facturas_db.save_invoice(invoice, "Prueba (Mako).pdf"), jh_facturas_db.save_invoice(invoice)]
    try:
        jh_facturas_db.save_invoice({**invoice, "amount": 140.0})
    except ValueError as exc:
        results.append(str(exc))
    month = jh_facturas_db.get_month_invoices(2026, 10)
    results.append([{k: v for k, v in inv.items() if k != "uploaded_at"} for inv in month])
    jh_facturas_db.delete_invoice("SI-999901")
    results.append(len(jh_facturas_db.list_all_invoices()))
    return results


def _escenario_ice_machine():
    """Alta, reemplazo (mismo número) y baja de un Payment Summary de prueba."""
    import ice_machine_db

    summary = {"summary_no": "999901", "reference": "IND-PRUEBA", "from_date": "2026-10-03", "to_date": "2026-10-09",
               "transactions": 12, "gross": 150.0, "process_fees": 4.5, "service_fees": 1.0, "net": 144.5,
               "notes": "con 'comillas' y 100%"}
    results = [ice_machine_db.save_summary(summary, "Payment Summary prueba.pdf"),
               ice_machine_db.save_summary({**summary, "net": 140.25, "notes": None})]
    results.append([{k: v for k, v in s.items() if k != "uploaded_at"} for s in ice_machine_db.list_month(2026, 10)])
    results += [ice_machine_db.delete_summary("999901"), ice_machine_db.delete_summary("999901")]
    results.append(len(ice_machine_db.list_month(2026, 10)))
    return results


def _escenario_reporte_mensual():
    """El reporte de septiembre guardado como diciembre, corregido a mano, listado y borrado."""
    import reporte_mensual_db

    report = reporte_mensual_db.get_report(2026, 9)
    copy = {**report, "month": 12, "from_date": "2026-12-01", "to_date": "2026-12-31"}
    reporte_mensual_db.save_report(copy, "Resumen Ventas prueba.pdf")
    store_info = {**report["store_info"], "cash": (report["store_info"]["cash"] or 0) + 1.5, "network_revenue": None}
    departments = report["departments"][:5] + [{"department": "PRUEBA 100%", "count": 3, "amount": 0.0}]
    reporte_mensual_db.update_report(2026, 12, store_info, departments, 123.45, 7, ["aviso de prueba"], from_sheet=True)
    saved = reporte_mensual_db.get_report(2026, 12)
    results = [{k: v for k, v in saved.items() if k not in ("updated_at", "edited_at", "sheet_added_at")}, reporte_mensual_db.list_months()]
    reporte_mensual_db.delete_report(2026, 12)
    results.append(reporte_mensual_db.get_report(2026, 12))
    return results


def _escenario_cheques():
    """Cheque cargado a mano, corregido, con datos automáticos, buscado y borrado."""
    import cheques_db

    fields = {"manual_date": "2026-10-05", "manual_amount": 250.5, "manual_supplier": "Prueba 100%", "manual_invoice": "F-1",
              "manual_debit_date": None, "voided": False}
    check_id = cheques_db.save_manual(None, 99901, fields)
    results = [check_id, cheques_db.save_manual(check_id, 99901, {**fields, "manual_amount": 260.0, "voided": True})]
    try:
        cheques_db.save_manual(None, 99901, fields)
    except ValueError as exc:
        results.append(str(exc))
    cheques_db.set_auto_fields(check_id, {"auto_supplier_key": "coca", "auto_supplier_label": "Coca-Cola", "auto_invoice_no": "123",
                                          "auto_invoice_date": "2026-10-01", "auto_amount": 260.0})
    results += [cheques_db.find_by_number(99901), cheques_db.get_check(check_id), len(cheques_db.list_checks()),
                cheques_db.delete_check(check_id), cheques_db.delete_check(check_id), cheques_db.find_by_number(99901)]
    return [_without(r, "uploaded_at") for r in results]


def _escenario_control_cmv():
    """Un reporte de Elistar de septiembre copiado a diciembre, reemplazado y borrado."""
    import control_cmv_db

    september = control_cmv_db.get_month(2026, 9)
    results = [sorted(september)]
    for kind, report in september.items():
        copy = {k: v for k, v in report.items() if k not in ("filename", "uploaded_at")}
        copy.update(month=12)
        control_cmv_db.save_report(copy, "Depts prueba.xls")
        control_cmv_db.save_report({**copy, "nota": "reemplazo con 100%"}, "Depts prueba 2.xls")
    results.append({k: _without(v, "uploaded_at") for k, v in control_cmv_db.get_month(2026, 12).items()})
    for kind in september:
        control_cmv_db.delete_report(2026, 12, kind)
    results.append(control_cmv_db.get_month(2026, 12))
    return results


def _escenario_depositos():
    """Recibos de prueba (uno sin N°, para el orden de los vacíos), duplicado, cambio de tipo, edición y baja."""
    import depositos_db

    a = depositos_db.add_deposit(2026, 10, "2026-10-08", 1234.5, "Recibo con 100%", "555001", "normal", "", "Dep (prueba).pdf", 0)
    b = depositos_db.add_deposit(2026, 10, "2026-10-08", 99.99, None, None, "normal", "", "Dep (prueba).pdf", 1)
    results = [depositos_db.find_duplicate("555001", "2026-10-08", 1234.501), depositos_db.find_duplicate("555001", "2026-10-08", 1234.6),
               depositos_db.find_by_source_page("Dep (prueba).pdf", 1), depositos_db.find_by_source_page("Dep (prueba).pdf", 0, "555001")]
    depositos_db.set_kind(a, "ice", "Ice Machine")
    depositos_db.update_deposit(b, "2026-10-09", 100.0, "editado", "food_truck")
    results += [[_without(d, "uploaded_at") for d in depositos_db.list_month(2026, 10)],
                [d["id"] for d in depositos_db.list_kind("ice")], _without(depositos_db.get_deposit(b), "uploaded_at")]
    for deposit_id in (a, b):
        depositos_db.delete_deposit(deposit_id)
    results.append(len(depositos_db.list_month(2026, 10)))
    return results


def _escenario_horas_trabajo():
    """Una semana de prueba cargada, vuelta a subir, con un empleado agregado, editado y borrado."""
    import horas_trabajo_db

    employees = [{"employee_name": "Prueba Uno", "hours_label": "26:48", "hours": 26.8}, {"employee_name": "Prueba Dos", "hours": 40}]
    week = horas_trabajo_db.upsert_week("2026-10-12", "2026-10-05", "2026-10-11", employees)
    again = horas_trabajo_db.upsert_week("2026-10-12", "2026-10-05", "2026-10-11", employees[:1], source="ocr")
    extra = horas_trabajo_db.add_employee(week, "Prueba Tres 100%", hours=5.5, rate=16.0, deduct=10.0, hours_label="05:30")
    horas_trabajo_db.update_employee(extra, hours=6.0, deduct=0.0)
    results = [week == again, horas_trabajo_db.get_employee_week(extra), horas_trabajo_db.known_employee_names(),
               [[_without(dict(w), "created_at", "employees"), [_without(dict(e), "id") for e in w.get("employees", [])]]
                for w in horas_trabajo_db.get_month_weeks(2026, 10)]]
    horas_trabajo_db.delete_employee(extra)
    horas_trabajo_db.delete_week(week)
    results += [horas_trabajo_db.get_week(week), len(horas_trabajo_db.get_month_weeks(2026, 10))]
    return results


def _escenario_jh_mensual():
    """Los reportes de J.H. de septiembre sumados a diciembre: primero un renglón, después todos, y la baja."""
    import jh_mensual_db

    september = jh_mensual_db.get_reports(2026, 9)
    results = [sorted(september)]
    for kind, report in sorted(september.items()):
        month_report = {**report, "month": 12, "from_date": "2026-12-01", "to_date": "2026-12-31"}
        results.append(jh_mensual_db.merge_report(kind, {**month_report, "rows": report["rows"][:1]}, "prueba.pdf"))
        results.append(jh_mensual_db.merge_report(kind, month_report, "prueba 2.pdf"))
        results.append(len(jh_mensual_db.get_rows(kind)))
    results.append({k: _without(v, "updated_at") for k, v in jh_mensual_db.get_reports(2026, 12).items()})
    for kind in september:
        jh_mensual_db.delete_report(kind, 2026, 12)
    results += [jh_mensual_db.get_reports(2026, 12), len(jh_mensual_db.get_all_invoice_numbers())]
    return results


def _escenario_documents():
    """Con el guardado prendido solo en la copia: listar, buscar, guardar uno de prueba, leerlo y borrarlo."""
    import tempfile as _tempfile

    import documents_db

    documents_db.GUARDAR_DOCUMENTOS = True
    try:
        clean = lambda docs: [_without(d, "uploaded_at", "stored_path") for d in docs]
        results = [clean(documents_db.list_all_documents("lottery")), clean(documents_db.search_documents("CLOSE", limit=5))]
        with _tempfile.NamedTemporaryFile("wb", suffix=".pdf", delete=False) as handle:
            handle.write(b"%PDF-1.4 prueba")
        doc_id = documents_db.store_document("lottery", handle.name, "Prueba 100%.pdf", 2026, 12, label="prueba")
        results += [doc_id, clean(documents_db.list_documents("lottery", 2026, 12)), _without(documents_db.get_document(doc_id), "uploaded_at", "stored_path")]
        documents_db.delete_document(doc_id)
        results.append(documents_db.get_document(doc_id))
        os.remove(handle.name)
    finally:
        documents_db.GUARDAR_DOCUMENTOS = False
    return results


def _escenario_fisico():
    """Factura de combustible de prueba con renglones, buscada sin distinguir mayúsculas, editada y borrada; ajustes del mes."""
    import fisico_db

    lines = [{"product_code": "087", "description": "Regular 100%", "qty_billing": 6000, "qty_freight": 6000,
              "unit_price_wo_tax": 2.1, "unit_price_with_tax": 2.75, "line_total": 16500.0},
             {"product_code": "093", "description": "Premium", "qty_billing": 2600, "qty_freight": 2600,
              "unit_price_wo_tax": 2.844, "unit_price_with_tax": 3.494, "line_total": 9084.4}]
    invoice_id = fisico_db.add_invoice("2026-12-03", "2026-12-08", "SI-PRUEBA1", 8600, 25584.4, source="pdf", bol_number="B1",
                                       lines=lines, source_filename="Fuel prueba.pdf")
    clean = lambda d: _without(d, "created_at", "updated_at") if d else d
    results = [invoice_id, clean(fisico_db.find_invoice_by_number("si-prueba1")), [_without(l, "id") for l in fisico_db.get_invoice_lines(invoice_id)]]
    fisico_db.update_invoice(invoice_id, "2026-12-04", "2026-12-09", "SI-PRUEBA1", 8601, 25600.0)
    fisico_db.set_month_initial_override(2026, 12, 2746.15, 7000.5)
    fisico_db.set_month_real_ending(2026, 12, 3000, "2026-12-31")
    results += [[clean(i) for i in fisico_db.get_month_invoices(2026, 12)], fisico_db.get_invoice_years(),
                _without(fisico_db.get_month_settings(2026, 12) or {}, "updated_at"), fisico_db.get_settings_years()]
    fisico_db.delete_invoice(invoice_id)
    results += [fisico_db.find_invoice_by_number("SI-PRUEBA1"), len(fisico_db.list_all_invoices())]
    return results


def _escenario_caja():
    """Gastos de prueba (con N° de ticket como número), repetidos, pendientes y saldos del mes."""
    import caja_db

    a = caja_db.add_expense_item("2026-12-05", 50.0, "Gasto con 100%", source="ocr", trans_no=77)
    b = caja_db.add_expense_item("2026-12-05", 50.0, None)
    results = [caja_db.find_same_expense("2026-12-05", 50.0, trans_no=77), caja_db.find_same_expense("2026-12-05", 50.004),
               caja_db.find_same_expense("2026-12-05", 50.0, trans_no="78")]
    caja_db.update_expense_detail(b, "detalle nuevo")
    pending = caja_db.add_pending_expense(2026, 12, "2026-12-06", 12.5, "ticket dudoso", "Gasto prueba.pdf")
    results += [caja_db.add_pending_expense(2026, 12, "2026-12-06", 12.5, "otra vez", "Gasto prueba.pdf"),
                caja_db.get_month_pending_expenses(2026, 12), caja_db.delete_pending_expense(pending), caja_db.delete_pending_expense(pending)]
    caja_db.set_month_opening_balance(2026, 12, 1234.56)
    caja_db.set_month_closing_override(2026, 12, 999.99)
    results += [caja_db.get_month_expense_items(2026, 12), caja_db.get_month_expenses(2026, 12), caja_db.get_expense_years(),
                _without(caja_db.get_month_settings(2026, 12) or {}, "updated_at")]
    for item in (a, b):
        caja_db.delete_expense_item(item)
    results.append(caja_db.get_month_expense_items(2026, 12))
    return [[_without(x, "created_at") for x in r] if isinstance(r, list) else _without(r, "created_at") for r in results]


def _escenario_gettel_toyota():
    """Días de Gettel y Toyota de prueba, pagos (con N° de transacción en otra caja), ajuste del mes y bajas."""
    from datetime import date as _date

    import gettel_db

    results = [gettel_db.upsert_vendor_totals("gettel", {_date(2026, 12, 1): {"amount": 812.34, "gallons": 250.5},
                                                         _date(2026, 12, 2): {"amount": 0.0, "gallons": 0.0}}),
               gettel_db.upsert_vendor_totals("toyota", {_date(2026, 12, 1): {"amount": 99.99, "gallons": 30.0}}, source="pdf")]
    pago = gettel_db.add_pago(_date(2026, 12, 5), 1, "TX-9001", 1234.5, source="pdf", empresa="Kia")
    gettel_db.update_pago(pago, _date(2026, 12, 6), 2, "TX-9001", 1300.0, empresa="Toyota")
    gettel_db.set_month_pendiente_anterior_override(2026, 12, 456.78)
    clean = lambda rows: [_without(r, "created_at", "updated_at") for r in rows]
    results += [_without(gettel_db.get_day("2026-12-01"), "updated_at"), clean(gettel_db.get_month_days(2026, 12)),
                gettel_db.get_month_gettel_amount(2026, 12), gettel_db.get_day_gettel_amount("2026-12-01"),
                _without(gettel_db.find_pago_by_transaction("tx-9001"), "created_at", "updated_at"), clean(gettel_db.get_month_pagos(2026, 12)),
                gettel_db.get_pago_years(), gettel_db.get_toyota_days_years(), gettel_db.get_earliest_month(),
                _without(gettel_db.get_pago_month_settings(2026, 12) or {}, "updated_at"), gettel_db.get_pago_settings_years()]
    gettel_db.delete_pago(pago)
    results.append(gettel_db.find_pago_by_transaction("TX-9001"))
    return results


def _escenario_lottery():
    """Un día de prueba por sus tres caminos, fechas de pago de los bloques (con el reajuste) y el reporte mensual."""
    import lottery_db

    real = lottery_db.get_day("2026-09-15")
    sales = {k: real.get(k) for k in lottery_db._SALES_REPORT_FIELDS}
    lottery_db.upsert_department_fields("2026-12-07", 3, 45.5, 2, 20.0, "ocr", "Close Store prueba.pdf")
    lottery_db.upsert_sales_report_fields("2026-12-07", sales, "portal", "Daily Sales prueba.pdf")
    lottery_db.upsert_manual_day("2026-12-08", {**sales, "online_count": 1, "online_net_sales": 10})
    lottery_db.blank_day("2026-12-09")
    results = [_without(lottery_db.get_day("2026-12-07"), "updated_at"), sorted(lottery_db.get_days_between("2026-12-01", "2026-12-31")),
               lottery_db.set_block_chase_date(2026, 49, "2026-12-11"), lottery_db.set_block_chase_date(2026, 48, "2026-12-03"),
               lottery_db.build_month_blocks(2026, 12), lottery_db.monthly_debit_total(2026, 12), lottery_db.get_available_years(),
               lottery_db.compute_month_closing(2026, 12)]
    report = lottery_db.get_monthly_report(2026, 9)
    if report:
        lottery_db.save_monthly_report({**report, "month": 12}, "Monthly prueba.pdf")
        lottery_db.save_monthly_report({**report, "month": 12, "warnings": ["reemplazo"]}, "Monthly prueba 2.pdf")
        results.append(_without(lottery_db.get_monthly_report(2026, 12), "updated_at"))
        lottery_db.delete_monthly_report(2026, 12)
    results += [lottery_db.get_monthly_report(2026, 12), lottery_db.set_block_chase_date(2026, 48, None)]
    return results


def _escenario_reportes_diarios():
    """Un día de prueba con departamentos (OCR y a mano), Store Info, totales impresos, márgenes y bajas."""
    import reportes_db

    real = reportes_db.get_day("2026-09-15")
    info = {k: v for k, v in (real["store_info"] or {}).items() if k not in ("date", "credit_terms_json", "store_info_source", "pdf_filename")}
    info["credit_terms"] = real["store_info"]["credit_terms"] if real["store_info"] else []
    records = [{"department": d["department"], "count": d["count"], "amount": d["amount"]} for d in real["departments"]]
    results = [reportes_db.replace_department_sales("2026-12-07", records, "Close Store 06-12.pdf")]
    reportes_db.upsert_department_row("2026-12-07", "PRUEBA 100%", 2, 0.0)
    reportes_db.upsert_department_row("2026-12-07", records[0]["department"], 999, 12.5)
    reportes_db.delete_department_row("2026-12-07", records[1]["department"])
    reportes_db.upsert_store_info("2026-12-07", info, "ocr", "Close Store 06-12.pdf")
    reportes_db.upsert_store_info("2026-12-07", {**info, "cash": None}, "ocr", keep_existing_for_none=True)
    reportes_db.upsert_printed_totals("2026-12-07", 1234.5, 321)
    margins = reportes_db.get_department_margins({"Tabacco": 0.1, "Soda": 0.3})
    reportes_db.save_department_margins({**margins, "Soda": 0.35})
    results += [reportes_db.get_day("2026-12-07"), reportes_db.get_month_overview(2026, 12), reportes_db.get_month_department_totals(2026, 12),
                reportes_db.get_month_departments_by_date(2026, 12), reportes_db.get_month_store_info(2026, 12),
                reportes_db.get_month_department_amounts(2026, 12, records[0]["department"]), reportes_db.get_month_local_accounts(2026, 12),
                reportes_db.get_store_info_years(), len(reportes_db.get_card_sales_by_date()), len(reportes_db.get_local_account_sales_by_date()),
                reportes_db.list_known_departments(), reportes_db.get_department_margins({"Tabacco": 0.1, "Soda": 0.3}),
                reportes_db.delete_departments_for_date("2026-12-07"), reportes_db.get_day("2026-12-07")["departments"]]
    return results


def _escenario_cuenta_kia_toyota():
    """Un Excel de prueba que reemplaza días ya cargados (vuelven al final, como en SQLite), pagos nuevos y repetidos,
    días del POS, asignaciones a mano; y la cuenta corriente que arma la página con eso."""
    import cuenta_kia_toyota_db

    charges = cuenta_kia_toyota_db.get_charges()
    payments = cuenta_kia_toyota_db.get_payments()
    days = [{**{k: c[k] for k in ("la", "kia", "kia_gal", "toyota", "toyota_gal")}, "date": d, "main": True, "kia": (c["kia"] or 0) + 1.25}
            for d, c in list(charges.items())[:3]]
    days.append({"date": "2026-12-05", "la": 812.5, "kia": 500.0, "kia_gal": 150.0, "toyota": 312.5, "toyota_gal": 90.0, "main": True})
    data = {"days": days, "payments": [{k: p[k] for k in ("date", "transc", "amount", "company")} for p in payments[:2]]
            + [{"date": "2026-12-06", "transc": "9001", "amount": 999.0, "company": None}],
            "pos": [{"date": "2026-12-05", "la": 812.5, "vs": 812.5}, {"date": "2026-12-06", "la": None, "vs": None}]}
    results = [cuenta_kia_toyota_db.save_import(data, "Cierre prueba.xlsx")]
    new = [p["id"] for p in cuenta_kia_toyota_db.get_payments() if p["transc"] == "9001"]
    cuenta_kia_toyota_db.assign_company(new, "Toyota")
    cuenta_kia_toyota_db.assign_pos_only("2026-12-05", "Kia")
    cuenta_kia_toyota_db.assign_pos_only("2026-12-05", "Toyota")
    results += [list(cuenta_kia_toyota_db.get_charges())[-5:], cuenta_kia_toyota_db.get_payments()[-3:],
                list(cuenta_kia_toyota_db.get_pos_days().items())[-2:], cuenta_kia_toyota_db.get_pos_only_assignments()]
    return results


def _escenario_eft():
    """Un EFT de prueba copiado de uno real, DDC corregido a mano, cupones nuevos y repetidos, detalle de cupones
    repetido y nuevo, y los cálculos que se arman con eso."""
    import eft_db

    real = eft_db.get_month_deposits(2026, 9)[0]
    coupons = [{k: c[k] for k in ("date", "invoice", "coupon", "gross_amount", "fees_amount", "paid_amount")} for c in real["coupons"]]
    invoices = [{k: i[k] for k in ("invoice", "paid_amount")} for i in real["paid_invoices"]]
    eft_date = real["deposit"]["eft_date"]
    deposit = eft_db.save_eft({"draft_no": "RCV-99901", "eft_date": eft_date}, invoices, coupons, "EFT prueba.pdf")
    net = sum(c["paid_amount"] or 0 for c in coupons)
    results = [deposit, bool(eft_db.find_existing_deposit("RCV-99901", eft_date, net))]
    first_line = next(d for d in eft_db.get_month_deposits(2026, 9) if d["deposit"]["id"] == deposit)["coupons"][0]["id"]
    eft_db.update_coupon_row(first_line, coupon="DDC-99901")
    record = {"coupon": "DDC-99901", "date": coupons[0]["date"], "gross": coupons[0]["gross_amount"], "fees": coupons[0]["fees_amount"],
              "net": coupons[0]["paid_amount"], "reported_group_text": None, "reported_group_total": None}
    results += [eft_db.upsert_cupones([record], "Cupones prueba.xlsx"), eft_db.upsert_cupones([{**record, "net": (record["net"] or 0) + 1}]),
                eft_db.insert_new_cupones([{**record, "coupon": "DDC-99902"}, record], "Daily Summary prueba.pdf"),
                eft_db.autolink_missing_ddc_by_amount(), eft_db.backfill_grouped_cupones_from_eft()]
    groups = eft_db.get_detail_groups()
    existing = groups[0]
    new_batches = [{**b, "batch": "PRUEBA" + str(i)} for i, b in enumerate(existing["batches"])]
    results.append(eft_db.save_coupon_detail(
        [{"batches": existing["batches"], "totals": {k: existing[k] for k in ("gross", "fees", "net")}, "coupons": existing["coupons"]},
         {"batches": new_batches, "totals": {k: existing[k] for k in ("gross", "fees", "net")}, "coupons": None}], "Detalle prueba.pdf"))
    results += [eft_db.identify_detail_groups(), eft_db.get_cupones_with_status(limit=15), len(eft_db.get_cupones_flat()),
                sorted(eft_db.get_coupon_gross_by_date().items())[-5:], eft_db.coupon_month_summary(2026, 9),
                eft_db.cupones_month_view(2026, 9), eft_db.get_unmatched_eft_coupons(), eft_db.get_cupones_years(), eft_db.get_deposit_years()]
    eft_db.delete_deposit(deposit)
    results += [eft_db.delete_cupones_month(2026, 12), len(eft_db.list_all_deposits())]
    return results


CHASE_REAL = "G:/.shortcut-targets-by-id/1NdiYjB3tA6w90-ekRpM2ksmVuK3i57_1/Bradenton/Book keeping/2026/10 Octubre - 26/CHASE 10-2026.xlsx"


def _escenario_chase(client):
    """El extracto real de octubre vuelto a subir por la página, correcciones a mano, depósitos y recategorizar."""
    import chase_db
    import webapp
    from chase_rules import categorize_chase_description

    with open(CHASE_REAL, "rb") as handle:  # el Drive solo se lee
        response = client.post("/carga-datos/chase", data={"chase_file": (io.BytesIO(handle.read()), "CHASE 10-2026.xlsx")},
                               content_type="multipart/form-data")
    with client.session_transaction() as session:
        flashes = session.get("_flashes", [])
    month = chase_db.get_month_transactions(2026, 10)
    first = month[0]
    chase_db.set_manual_detalle(first["posting_date"], first["description"], first["amount"], "PRUEBA 100%")
    chase_db.set_manual_supplier(first["posting_date"], first["description"], first["amount"], "fpl")
    deposits = chase_db.deposits_on("2026-10-08")
    results = [response.status_code, response.headers.get("Location"), flashes, len(month), deposits,
               chase_db.recategorize_all(categorize_chase_description, webapp._resolve_supplier_key_for_chase),
               chase_db.get_uncategorized_count(2026, 10), chase_db.get_month_transactions(2026, 10)[:5],
               chase_db.get_supplier_transactions("fpl")[-3:], chase_db.get_check_transactions()[-3:],
               chase_db.has_detalle_on_date("2026-10-08", "DEPOSITO"), chase_db.list_known_details()[:10],
               chase_db.get_first_posting_date(), chase_db.get_last_posting_date(), chase_db.get_available_years()]
    return results


def _escenario_cmv():
    """El CMV entero vuelto a cargar con precios cambiados y un producto nuevo, foto del día, ventas de un
    departamento cargadas dos veces y lo que se calcula con eso."""
    import cmv_db

    costs = cmv_db.get_all_costs()
    keys = ("upc", "upc_mod", "name", "cost", "price", "dept_id", "dept_name")
    records = [{k: c[k] for k in keys} for c in costs]
    for r in records[:3]:
        r["price"] = (r["price"] or 0) + 0.5
    records.append({**records[0], "upc": "999999999901", "name": "PRUEBA 100%", "cost": 1.234, "price": 2.5})
    results = [cmv_db.replace_costs_for_departments(records), cmv_db.save_snapshot("2026-12-01"), cmv_db.list_snapshots()[:2],
               len(cmv_db.get_snapshot_items("2026-12-01")), cmv_db.get_snapshot_items("2026-12-01")[:3]]
    dept = records[0]["dept_name"]
    rows = [{"upc": r["upc"], "name": r["name"], "count": i + 1, "amount": round(1.1 * (i + 1), 2)} for i, r in enumerate(records[:6])]
    cmv_db.replace_month_department_sales(2026, 12, dept, rows + rows[:1])
    cmv_db.replace_month_department_sales(2026, 12, dept.lower(), rows[:4])
    results += [cmv_db.get_costs_by_department(dept.upper())[:5], cmv_db.get_month_department_totals(2026, 12),
                cmv_db.get_monthly_sales_rows([(2026, 12), (2026, 9)])[:8], len(cmv_db.get_all_monthly_sales()),
                cmv_db.get_recent_price_changes(5), cmv_db.list_departments()[:5], cmv_db.get_month_years_with_data()]
    return results


def _escenario_proveedores():
    """Facturas nuevas y repetidas (la repetida da IntegrityError y la conexión sigue andando), renglones,
    ocultar, configuración, credit memos, pagos históricos y a mano, y las bajas."""
    import proveedores_db

    results = [proveedores_db.save_invoice("cec", "CEC", "2026-12-03", "PRUEBA-1", 1546.8, "CEC prueba.pdf"),
               proveedores_db.save_invoice("cec", "CEC", "2026-12-03", "PRUEBA-1", 9.99),
               proveedores_db.save_invoice("cec", "CEC", "2026-12-04", "PRUEBA-2", 10.0),
               proveedores_db.invoice_exists("cec", "PRUEBA-1")]
    lines = [{"line_no": i, "upc": f"0000{i}", "description": f"Producto {i} 100%", "pack": 10, "qty": 2, "unit_cost": 1.234 * i,
              "line_total": 2.468 * i} for i in range(1, 4)]
    try:
        results.append(proveedores_db.replace_invoice_lines("cec", "PRUEBA-1", "2026-12-03", lines))
    except Exception as exc:  # el formato de renglones lo valida el módulo: alcanza con que falle igual en los dos
        results.append(f"{type(exc).__name__}: {exc}")
    proveedores_db.set_supplier_hidden("prueba_oculto", True)
    proveedores_db.set_supplier_hidden("prueba_oculto", True)
    proveedores_db.set_supplier_settings("cec", True, True)
    memo = proveedores_db.save_credit_memo("cec", "CEC", "2026-12-05", "CM-1", 15.22, "devolución 100%")
    proveedores_db.replace_history_payments("cec", [{"date": "2026-01-02", "amount": 100.0, "detail": "OP"},
                                                    {"date": "2026-01-02", "amount": 50.0, "detail": "OP 2"}], "Excel prueba")
    manual = proveedores_db.save_manual_payment("bimbo", "Bimbo", "2026-12-06", 75.5, "pago en caja", caja_expense_item_id=999901)
    # Sin ids de factura: en Postgres el INSERT rechazado por duplicado consume un número de la secuencia
    # (los ids pueden saltear números; el orden no cambia).
    def no_ids(value):
        if isinstance(value, dict):
            return {k: no_ids(v) for k, v in value.items() if k != "id"}
        return [no_ids(v) for v in value] if isinstance(value, list) else value
    results += [no_ids(proveedores_db.get_month_invoices(2026, 12)), proveedores_db.list_suppliers_with_counts(), no_ids(proveedores_db.get_supplier_invoices("cec")[:4]),
                len(proveedores_db.get_all_invoice_lines()), proveedores_db.list_hidden_suppliers(), proveedores_db.get_supplier_settings("cec"),
                proveedores_db.get_supplier_credit_memos("cec"), proveedores_db.get_supplier_history_payments("cec"),
                proveedores_db.get_supplier_manual_payments("bimbo")[:3], proveedores_db.get_available_years()]
    proveedores_db.delete_credit_memo(memo)
    proveedores_db.delete_manual_payment(manual)
    proveedores_db.delete_manual_payment_by_caja_item(999901)
    proveedores_db.set_supplier_hidden("prueba_oculto", False)
    for inv in proveedores_db.get_supplier_invoices("cec"):
        if inv["invoice_no"].startswith("PRUEBA-"):
            proveedores_db.delete_invoice(inv["id"])
    results += [proveedores_db.get_month_invoices(2026, 12), proveedores_db.list_hidden_suppliers()]
    return results


def _escenario_config():
    """Usuarios (alta, contraseña, baja), reglas de Chase y de pagos (alta, edición, baja, categorizar), límite de
    Caja, Sale Tax del mes y un proveedor sin código; y que la clave de sesión no cambie entre lecturas."""
    import auth
    import chase_rules
    import control_cierre
    import controles_rapidos
    import proveedores_dynamic_extractors as dyn
    import proveedores_pago_rules

    results = []
    auth.create_user("prueba", "clave-de-prueba", is_admin=False)
    results += [auth.list_users(), auth.must_change_password("prueba"), bool(auth.verify_user("prueba", "clave-de-prueba")),
                bool(auth.verify_user("prueba", "otra"))]
    auth.set_password("prueba", "clave-nueva")
    results += [auth.must_change_password("prueba"), bool(auth.verify_user("prueba", "clave-nueva"))]
    try:
        auth.create_user("prueba", "x")
    except ValueError as exc:
        results.append(str(exc))
    auth.delete_user("prueba")
    results.append(auth.list_users())
    chase_rules.add_dynamic_rule("PRUEBA 100% COMERCIO", "MOVILIDAD Y VIATICOS")
    rules = chase_rules.load_dynamic_rules()
    chase_rules.edit_dynamic_rule_by_index(len(rules) - 1, "PRUEBA 100% COMERCIO", "GASTOS VARIOS")
    results += [chase_rules.load_dynamic_rules()[-2:], chase_rules.categorize_chase_description("PAGO PRUEBA 100% COMERCIO 1234"),
                len(chase_rules.load_master_rules()), chase_rules.list_display_rules()[:3]]
    chase_rules.delete_dynamic_rule_by_index(len(rules) - 1)
    results.append(len(chase_rules.load_dynamic_rules()))
    proveedores_pago_rules.add_dynamic_rule("PRUEBA PAGO", "PRUEBA")
    results.append(proveedores_pago_rules.match_supplier_sheet("PAGO A PRUEBA PAGO 99"))
    pago = proveedores_pago_rules.load_dynamic_rules()
    proveedores_pago_rules.delete_dynamic_rule_by_index(len(pago) - 1)
    results.append(len(proveedores_pago_rules.load_dynamic_rules()))
    controles_rapidos.set_caja_limit(9876.5)
    control_cierre.set_sale_tax_record(2026, 9, True, "A-123")
    results += [controles_rapidos.caja_limit(), {k: v for k, v in (control_cierre.sale_tax_record(2026, 9) or {}).items() if k != "marked_at"}]
    control_cierre.set_sale_tax_record(2026, 9, False)
    results.append(control_cierre.sale_tax_record(2026, 9))
    dyn._save_dynamic_suppliers_file({"prueba": {"label": "Prueba 100%", "sheet_name": "PRUEBA", "fields": {"total": "TOTAL"}}})
    results.append(dyn.load_dynamic_suppliers())
    dyn._save_dynamic_suppliers_file({})
    import webapp
    results.append(webapp._load_or_create_secret_key() == webapp._load_or_create_secret_key())
    return results


def _without(value, *keys):
    return {k: v for k, v in value.items() if k not in keys} if isinstance(value, dict) else value


SCENARIOS = {"jh_facturas": _escenario_jh_facturas, "ice_machine": _escenario_ice_machine,
             "reporte_mensual": _escenario_reporte_mensual, "cheques": _escenario_cheques,
             "control_cmv": _escenario_control_cmv, "depositos": _escenario_depositos,
             "horas_trabajo": _escenario_horas_trabajo, "jh_mensual": _escenario_jh_mensual,
             "documents": _escenario_documents, "fisico": _escenario_fisico,
             "caja": _escenario_caja, "gettel_toyota": _escenario_gettel_toyota,
             "lottery": _escenario_lottery,
             "reportes_diarios": _escenario_reportes_diarios, "cuenta_kia_toyota": _escenario_cuenta_kia_toyota,
             "eft": _escenario_eft, "chase": _escenario_chase,
             "cmv": _escenario_cmv, "proveedores": _escenario_proveedores, "config": _escenario_config}


# --- Una corrida (proceso aparte por motor) --------------------------------

def _normalize_html(text):
    text = re.sub(r"\b\d{1,2}:\d{2}(:\d{2})?\b", "HH:MM", text)  # horas de "generado a las ..."
    return text


def _body(response):
    kind = response.headers.get("Content-Type", "")
    data = response.get_data()
    if "pdf" in kind:
        from pypdf import PdfReader

        return "\n".join(page.extract_text() or "" for page in PdfReader(io.BytesIO(data)).pages)
    if "spreadsheet" in kind or "excel" in kind:
        import openpyxl

        book = openpyxl.load_workbook(io.BytesIO(data))
        return "\n".join(f"{sheet.title}!{cell.coordinate}={cell.value!r}" for sheet in book for row in sheet.iter_rows() for cell in row if cell.value is not None)
    if "json" in kind or "html" in kind or "text" in kind:
        return _normalize_html(data.decode("utf-8", "replace"))
    return f"<{len(data)} bytes {kind}>"


_TIME_KEYS = {"updated_at", "created_at", "uploaded_at", "changed_at", "edited_at", "sheet_added_at", "hidden_at", "registrado_en"}


def _drop_times(value):
    """Saca las horas de guardado (cambian entre corridas) de lo que devuelve un escenario."""
    if hasattr(value, "keys") and not isinstance(value, dict):
        value = {k: value[k] for k in value.keys()}
    if isinstance(value, dict):
        return {k: _drop_times(v) for k, v in value.items() if k not in _TIME_KEYS}
    if isinstance(value, (list, tuple)):
        return [_drop_times(v) for v in value]
    return value


def _jsonable(value):
    """Filas (sqlite3.Row o db.Row) como dict; lo demás como texto."""
    if hasattr(value, "keys"):
        return {k: value[k] for k in value.keys()}
    if isinstance(value, (set, frozenset)):  # el orden de un set cambia en cada corrida de Python
        return sorted(value, key=repr)
    return str(value)


def _capture(modules, out_path):
    os.startfile = lambda *a, **k: None  # que ninguna ruta abra Excel en la PC
    code_dir = os.environ.get("BRADENTON_CODE_DIR") or APP_DIR
    sys.path.insert(0, code_dir)
    import flask_login.utils
    import webapp

    app = webapp.app
    environ = {"REMOTE_ADDR": "127.0.0.1", "HTTP_USER_AGENT": "comparador-bases"}
    with app.test_request_context(environ_base=environ):
        identifier = flask_login.utils._create_identifier()
    client = app.test_client()
    client.environ_base.update(environ)
    with client.session_transaction() as session:
        session.update({"_user_id": "admin", "_fresh": True, "_id": identifier, "_last_seen": time.time()})

    def pages(tag):
        found = {}
        for module in modules:
            for url in PAGES.get(module, []):
                response = client.get(url)
                found[f"{tag} {url}"] = {"status": response.status_code, "body": _body(response)}
        return found

    result = pages("antes")
    for module in modules:
        if module in SCENARIOS:
            scenario = SCENARIOS[module]
            output = scenario(client) if scenario.__code__.co_argcount else scenario()
            result[f"escritura {module}"] = {"status": 0, "body": json.dumps(_drop_times(output), ensure_ascii=False, indent=1, default=_jsonable)}
    result.update(pages("después"))
    with open(out_path, "w", encoding="utf-8") as handle:
        json.dump(result, handle, ensure_ascii=False)


# --- Orquestación ----------------------------------------------------------

def _copy_data(source, target):
    os.makedirs(target)
    for name in os.listdir(source):
        if name.endswith(".db"):
            import sqlite3

            src = sqlite3.connect(f"file:{os.path.join(source, name)}?mode=ro", uri=True)
            dst = sqlite3.connect(os.path.join(target, name))
            src.backup(dst)
            src.close()
            dst.close()
        elif name.endswith(".json"):
            shutil.copy2(os.path.join(source, name), target)
    for name in CONFIG_FILES:
        if os.path.isfile(os.path.join(APP_DIR, name)):
            shutil.copy2(os.path.join(APP_DIR, name), target)


def _run_capture(modules, data_dir, out_path, database_url, code_dir=None):
    env = {k: v for k, v in os.environ.items() if k not in ("DATABASE_URL", "BRADENTON_CODE_DIR")}
    env.update({"BRADENTON_DATA_DIR": data_dir, "PYTHONIOENCODING": "utf-8"})
    if code_dir:
        env["BRADENTON_CODE_DIR"] = code_dir
    if database_url:
        env["DATABASE_URL"] = database_url
    done = subprocess.run([sys.executable, os.path.abspath(__file__), "--capturar", out_path, *modules], env=env, cwd=APP_DIR)
    if done.returncode:
        sys.exit(f"La corrida {'Postgres' if database_url else 'SQLite'} falló (ver arriba).")
    with open(out_path, encoding="utf-8") as handle:
        return json.load(handle)


def main():
    parser = argparse.ArgumentParser(description="Compara la app sobre SQLite y sobre PostgreSQL con los mismos datos.")
    parser.add_argument("modulos", nargs="+")
    parser.add_argument("--datos", help="carpeta de datos de origen (por defecto, la de la app)")
    parser.add_argument("--original", metavar="COMMIT",
                        help="además, corre el código de ese commit (antes del pase) sobre SQLite: lo nuevo no tiene que cambiar nada")
    parser.add_argument("--capturar", help=argparse.SUPPRESS)
    args = parser.parse_args()
    if args.capturar:
        return _capture(args.modulos, args.capturar)

    sys.path.insert(0, APP_DIR)
    import app_paths
    import db

    if not db.uses_postgres():
        db.load_dotenv()
    url = db.database_url()
    import psycopg

    if not url or (psycopg.conninfo.conninfo_to_dict(url).get("host") or "localhost") not in LOCAL_HOSTS:
        sys.exit("El comparador vacía y recarga tablas: necesita DATABASE_URL apuntando a un Postgres de esta PC (127.0.0.1).")
    unknown = [m for m in args.modulos if m not in PAGES]
    if unknown:
        sys.exit(f"Sin páginas registradas para: {', '.join(unknown)}")

    with tempfile.TemporaryDirectory(prefix="bradenton_comparar_") as folder:
        lite_dir, pg_dir = os.path.join(folder, "sqlite"), os.path.join(folder, "postgres")
        source = args.datos or app_paths.DATA_DIR
        _copy_data(source, lite_dir)
        _copy_data(source, pg_dir)
        # Se recargan todos los módulos que ya están en Postgres, no solo los
        # que se comparan: las páginas leen también los otros.
        with psycopg.connect(url) as pg:
            migrated = [r[0] for r in pg.execute(
                "SELECT DISTINCT table_schema FROM information_schema.tables WHERE table_schema = ANY(%s)", (list(db.SCHEMAS),)
            ).fetchall()]
        missing = [m for m in args.modulos if m not in migrated]
        if missing:
            sys.exit(f"Sin tablas en Postgres para: {', '.join(missing)} (¿falta alembic upgrade head?)")
        copy = subprocess.run([sys.executable, os.path.join(APP_DIR, "migrar_sqlite_a_postgres.py"), *sorted(migrated), "--origen", pg_dir, "--vaciar",
                               "--config-origen", pg_dir],
                              cwd=APP_DIR, env={**os.environ, "DATABASE_URL": url, "PYTHONIOENCODING": "utf-8"})
        if copy.returncode:
            sys.exit("La copia a Postgres no quedó idéntica: no se compara.")
        lite = _run_capture(args.modulos, lite_dir, os.path.join(folder, "sqlite.json"), None)
        pg = _run_capture(args.modulos, pg_dir, os.path.join(folder, "postgres.json"), url)
        original = None
        if args.original:
            code_dir, orig_dir = os.path.join(folder, "codigo_original"), os.path.join(folder, "sqlite_original")
            subprocess.run(["git", "worktree", "add", "--detach", "--quiet", code_dir, args.original], cwd=APP_DIR, check=True)
            try:
                _copy_data(source, orig_dir)
                original = _run_capture(args.modulos, orig_dir, os.path.join(folder, "original.json"), None, code_dir)
            finally:
                subprocess.run(["git", "worktree", "remove", "--force", code_dir], cwd=APP_DIR)

    differences, bad_status = 0, []
    print()
    for key in lite:
        a, b = lite[key], pg.get(key)
        same = b is not None and a == b
        differences += not same
        if a["status"] not in (0, 200) or (b and b["status"] not in (0, 200)) or "<title>Iniciar sesión" in a["body"]:
            bad_status.append(f"{key} ({a['status']}/{b['status'] if b else '-'})")
        if not same:
            print(f"DISTINTO  {key}  ({a['status']}/{b['status'] if b else '-'})")
            if b is not None:
                diff = difflib.unified_diff(a["body"].splitlines(), b["body"].splitlines(), "SQLite", "Postgres", lineterm="", n=1)
                for line in list(diff)[:40]:
                    print("    " + line[:200])
    print(f"{len(lite) - differences} de {len(lite)} comparaciones idénticas (SQLite contra Postgres).")
    if original is not None:
        changed = 0
        for key in lite:
            if original.get(key) != lite[key]:
                changed += 1
                print(f"CAMBIÓ CONTRA EL ORIGINAL  {key}")
                if key in original:
                    diff = difflib.unified_diff(original[key]["body"].splitlines(), lite[key]["body"].splitlines(), "original", "nuevo", lineterm="", n=1)
                    for line in list(diff)[:30]:
                        print("    " + line[:200])
        print(f"{len(lite) - changed} de {len(lite)} iguales al código original sobre SQLite.")
        differences += changed
    if bad_status:
        print("Páginas que no respondieron bien (o pidieron login):\n  " + "\n  ".join(bad_status))
    print("Todo idéntico." if not differences and not bad_status else "HAY QUE REVISAR antes del commit.")
    sys.exit(1 if differences or bad_status else 0)


if __name__ == "__main__":
    main()
