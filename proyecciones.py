"""
Proyecciones (Reportes): el panel que miran los dueños día por día (pedido
del usuario, 2026-10-26; reemplaza a la página escondida del 2026-09-18).

Reproduce la hoja "Store Info" del Excel Cierre (filas Ventas / Margen /
Utilidad / Proyecciones) y el "PROYECTADO BRADENTON" del Drive:
  - Proyectado = vendido en los días cargados / días cargados * días del mes.
    El Excel usa siempre 30; acá van los días reales del mes.
  - Combustible: venta neta = Sales Fuel + Desc. Comb. El costo es el PPP de
    la hoja Físico: (inventario inicial + facturas de J.H.) / (galones
    iniciales + comprados), y la utilidad, venta - galones * PPP.
  - C-Store: venta de cada rubro * su % teórico (editables, los del Excel).

Control contra el Balance (hoja "BALANCE HISTORICO" del Bce): para un mes
cerrado se compara lo que proyectaba la app los días 7, 15 y 22 y a fin de mes
con lo real. Medido contra agosto 2026: "4.01 Ventas de combustible" es
Sales Fuel + Desc. Comb y "4.09 Ventas C-Store" son Tabacco + Soda + Beer/Wine
+ Resto (Lotto y Kia/Toyota van por otras cuentas), al centavo.
"""

import calendar
import datetime as dt
import os

import config_store
import fisico
import fisico_db
import reportes_db
from reporte_diario import DEPARTMENT_GROUPS, group_department_sales

# Los % teóricos del Excel Cierre (fila Margen, abajo de cada rubro). Son el
# default; el que se usa sale de reportes_db.get_department_margins.
DEFAULT_MARGINS = {
    "TABACCO": 0.15,
    "SODA": 0.40,
    "BEER/WINE": 0.25,
    "LOTERY/LOTTO": 0.06,
    "Gettel": 1.00,
    "RESTO": 0.40,
}

CATEGORY_DISPLAY_LABELS = {
    "TABACCO": "Tabacco",
    "SODA": "Soda",
    "BEER/WINE": "Beer / Wine",
    "LOTERY/LOTTO": "Lotto",
    "Gettel": "Kia / Toyota",
    "RESTO": "Resto",
}

# Los rubros que el Balance cuenta como "Ventas C-Store".
BALANCE_CSTORE = ("TABACCO", "SODA", "BEER/WINE", "RESTO")
CHECKPOINT_DAYS = (7, 15, 22)


def _project(actual, days_loaded, days_in_month):
    if not days_loaded or actual is None:
        return None
    return actual * days_in_month / days_loaded


def _ratio(num, den):
    return num / den if (num is not None and den) else None


def _ppp(year, month, upto=None):
    """
    Precio promedio ponderado del combustible: el inventario inicial de la hoja
    Físico más las facturas de J.H. del mes (hasta `upto`, para proyectar como
    se veía ese día). None si todavía no hay galones.
    """
    gallons, amount, _source = fisico._resolve_initial(year, month)
    invoices = fisico_db.get_month_invoices(year, month)
    if upto is not None:
        invoices = [inv for inv in invoices if (inv.get("invoice_date") or "") <= upto.isoformat()]
    gallons += sum(inv.get("gallons") or 0.0 for inv in invoices)
    amount += sum(inv.get("amount") or 0.0 for inv in invoices)
    return (amount / gallons if gallons else None), len(invoices)


def _month_days(year, month):
    """Un renglón por día cargado: combustible y C-Store por rubro."""
    store = {r["date"]: r for r in reportes_db.get_month_store_info(year, month) if r.get("volume") is not None}
    departments = reportes_db.get_month_departments_by_date(year, month)
    days = []
    for date_key in sorted(store):
        row = store[date_key]
        groups, _ = group_department_sales(departments.get(date_key, []))
        by_group = {g["label"]: g.get("amount") or 0.0 for g in groups}
        fuel_sales = (row.get("sales_fuel") or 0.0) + (row.get("desc_comb") or 0.0)
        cstore = sum(by_group.get(label, 0.0) for label in BALANCE_CSTORE)
        days.append({
            "date": dt.date.fromisoformat(date_key),
            "gallons": row.get("volume") or 0.0,
            "fuel_sales": fuel_sales,
            "price": _ratio(fuel_sales, row.get("volume")),
            "groups": by_group,
            "cstore": cstore,
            "lotto": by_group.get("LOTERY/LOTTO", 0.0),
            "kia": by_group.get("Gettel", 0.0),
            "total": fuel_sales + cstore,
        })
    return days


def _summary(year, month, days, margins, upto=None):
    """Vendido y proyectado con los días hasta `upto` (todos si es None)."""
    days_in_month = calendar.monthrange(year, month)[1]
    if upto is not None:
        days = [d for d in days if d["date"] <= upto]
    loaded = len(days)
    gallons = sum(d["gallons"] for d in days)
    fuel_sales = sum(d["fuel_sales"] for d in days)
    ppp, invoices = _ppp(year, month, upto)

    def fuel_block(gal, sales):
        if gal is None or sales is None:
            return {"gallons": None, "sales": None, "price": None, "cost": None, "margin_gal": None, "utilidad": None}
        cost = gal * ppp if ppp is not None else None
        price = _ratio(sales, gal)
        return {
            "gallons": gal,
            "sales": sales,
            "price": price,
            "cost": cost,
            "margin_gal": (price - ppp) if (price is not None and ppp is not None) else None,
            "utilidad": (sales - cost) if cost is not None else None,
        }

    rubros = []
    for label, _members in DEPARTMENT_GROUPS:
        amount = sum(d["groups"].get(label, 0.0) for d in days)
        margin = margins.get(label, DEFAULT_MARGINS.get(label, 0.0))
        proj = _project(amount, loaded, days_in_month)
        rubros.append({
            "label": label,
            "display": CATEGORY_DISPLAY_LABELS.get(label, label),
            "in_balance": label in BALANCE_CSTORE,
            "margin": margin,
            "amount": amount,
            "utilidad": amount * margin,
            "proj": proj,
            "utilidad_proj": proj * margin if proj is not None else None,
        })

    def totals(items, key):
        values = [r[key] for r in items]
        return None if any(v is None for v in values) else sum(values)

    balance_rubros = [r for r in rubros if r["in_balance"]]
    cstore = {
        "rubros": rubros,
        "total": {k: totals(rubros, k) for k in ("amount", "utilidad", "proj", "utilidad_proj")},
        "balance": {k: totals(balance_rubros, k) for k in ("amount", "utilidad", "proj", "utilidad_proj")},
    }
    for block in (cstore["total"], cstore["balance"]):
        block["margin"] = _ratio(block["utilidad"], block["amount"])
        block["margin_proj"] = _ratio(block["utilidad_proj"], block["proj"])

    return {
        "days_in_month": days_in_month,
        "days_loaded": loaded,
        "last_day": days[-1]["date"] if days else None,
        "ppp": ppp,
        "fuel_invoices": invoices,
        "fuel": fuel_block(gallons, fuel_sales) if loaded else fuel_block(None, None),
        "fuel_proj": fuel_block(_project(gallons, loaded, days_in_month), _project(fuel_sales, loaded, days_in_month)),
        "cstore": cstore,
    }


def build_projection(year, month):
    """Todo lo que muestra la página de Proyecciones para un mes."""
    margins = reportes_db.get_department_margins(DEFAULT_MARGINS)
    days = _month_days(year, month)
    result = _summary(year, month, days, margins)
    result.update({"year": year, "month": month, "days": days, "balance": None})
    # Cada día con el PPP del mes, como la hoja Físico: la suma de los días da
    # la utilidad del mes.
    ppp = result["ppp"]
    for d in days:
        d["margin_gal"] = (d["price"] - ppp) if (d["price"] is not None and ppp is not None) else None
        d["fuel_utilidad"] = (d["fuel_sales"] - d["gallons"] * ppp) if ppp is not None else None
        d["cstore_utilidad"] = sum(d["groups"].get(label, 0.0) * margins.get(label, DEFAULT_MARGINS[label])
                                   for label in BALANCE_CSTORE)
    totals = {key: sum(d[key] for d in days)
              for key in ("gallons", "fuel_sales", "cstore", "cstore_utilidad", "lotto", "kia", "total")}
    totals["groups"] = {label: sum(d["groups"].get(label, 0.0) for d in days) for label, _members in DEPARTMENT_GROUPS}
    totals["price"] = _ratio(totals["fuel_sales"], totals["gallons"])
    totals["margin_gal"] = (totals["price"] - ppp) if (totals["price"] is not None and ppp is not None) else None
    totals["fuel_utilidad"] = (totals["fuel_sales"] - totals["gallons"] * ppp) if ppp is not None else None
    result["days_totals"] = totals
    gallons, amount, source = fisico._resolve_initial(year, month)
    result["fuel_initial"] = {"gallons": gallons, "amount": amount, "source": source}

    real =balance_for_month(year, month)
    if real:
        result["balance"] = _balance_control(year, month, days, margins, real)
    return result


# --- Control contra el Balance ----------------------------------------------

BALANCE_FIELDS = ("galones", "ventas_combustible", "costo_combustible", "ventas_cstore", "costo_cstore")
_BALANCE_ROWS = {
    "galones vendidos": "galones",
    "4.01.00.00 - ventas de combustible": "ventas_combustible",
    "4.09.00.00 - ventas c-store": "ventas_cstore",
    "5.02.00.00 - costo de ventas de combustible": "costo_combustible",
    "5.15.00.00 - costo de venta c-store": "costo_cstore",
}


def read_balance_workbook(path):
    """
    {"AAAA-MM": {galones, ventas_combustible, costo_combustible, ventas_cstore,
    costo_cstore}} de la hoja "BALANCE HISTORICO" de un Bce: una columna por
    mes, con las fechas en el renglón "Nota 1 : Ventas". Solo los meses con al
    menos la venta de combustible; un dato vacío queda en None.
    """
    import openpyxl

    workbook = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        if "BALANCE HISTORICO" not in workbook.sheetnames:
            raise ValueError("El archivo no tiene la hoja \"BALANCE HISTORICO\": ¿es el Bce del mes?")
        rows = list(workbook["BALANCE HISTORICO"].iter_rows(values_only=True))
    finally:
        workbook.close()

    header = next((i for i, row in enumerate(rows)
                   if row and isinstance(row[0], str) and row[0].strip().lower().startswith("nota 1")), None)
    if header is None:
        raise ValueError("No encontré el renglón \"Nota 1 : Ventas\" con los meses en BALANCE HISTORICO.")
    months = {col: f"{value.year:04d}-{value.month:02d}"
              for col, value in enumerate(rows[header]) if isinstance(value, (dt.date, dt.datetime))}
    found = {}
    # La primera aparición de cada renglón después de los meses (más abajo se repiten).
    for row in rows[header + 1:]:
        label = (row[0] or "").strip().lower() if row and isinstance(row[0], str) else ""
        field = _BALANCE_ROWS.get(label)
        if field and field not in found:
            found[field] = row
    if "ventas_combustible" not in found:
        raise ValueError("No encontré \"4.01.00.00 - VENTAS DE COMBUSTIBLE\" en BALANCE HISTORICO.")

    result = {}
    for col, key in months.items():
        values = {}
        for field in BALANCE_FIELDS:
            row = found.get(field)
            value = row[col] if row is not None and col < len(row) else None
            values[field] = round(float(value), 2) if isinstance(value, (int, float)) else None
        if values["ventas_combustible"]:
            result[key] = values
    if not result:
        raise ValueError("El Balance no tiene ningún mes con ventas de combustible.")
    return result


def save_balance(months, filename):
    """
    Guarda los meses leídos. Un mes ya guardado solo se reemplaza si el archivo
    nuevo lo trae igual o más completo (el Bce de septiembre trae agosto
    completo y septiembre a medio cargar). Devuelve (guardados, salteados).
    """
    stored = config_store.read(config_store.PROYECCIONES_BALANCE) or {}
    saved, skipped = [], []
    for key, values in sorted(months.items()):
        filled = sum(1 for f in BALANCE_FIELDS if values.get(f) is not None)
        previous = stored.get(key) or {}
        if sum(1 for f in BALANCE_FIELDS if previous.get(f) is not None) > filled:
            skipped.append(key)
            continue
        stored[key] = dict(values, archivo=os.path.basename(filename),
                           cargado=dt.datetime.now().isoformat(timespec="seconds"))
        saved.append(key)
    config_store.write(config_store.PROYECCIONES_BALANCE, stored)
    return saved, skipped


def balance_for_month(year, month):
    stored = config_store.read(config_store.PROYECCIONES_BALANCE) or {}
    return stored.get(f"{year:04d}-{month:02d}")


def _balance_control(year, month, days, margins, real):
    """Lo real del Balance contra la proyección de los días 7, 15, 22 y fin de mes."""
    def utilidad(sales, cost):
        return (sales - cost) if (sales is not None and cost is not None) else None

    real_rows = {
        "galones": real.get("galones"),
        "ventas_combustible": real.get("ventas_combustible"),
        "utilidad_combustible": utilidad(real.get("ventas_combustible"), real.get("costo_combustible")),
        "ventas_cstore": real.get("ventas_cstore"),
        "utilidad_cstore": utilidad(real.get("ventas_cstore"), real.get("costo_cstore")),
    }
    days_in_month = calendar.monthrange(year, month)[1]
    cuts = [(f"Día {d}", dt.date(year, month, d)) for d in CHECKPOINT_DAYS if d < days_in_month]
    cuts.append(("Fin de mes", dt.date(year, month, days_in_month)))
    checkpoints = []
    for label, cut in cuts:
        summary = _summary(year, month, days, margins, upto=cut)
        if not summary["days_loaded"]:
            continue
        proj = {
            "galones": summary["fuel_proj"]["gallons"],
            "ventas_combustible": summary["fuel_proj"]["sales"],
            "utilidad_combustible": summary["fuel_proj"]["utilidad"],
            "ventas_cstore": summary["cstore"]["balance"]["proj"],
            "utilidad_cstore": summary["cstore"]["balance"]["utilidad_proj"],
        }
        checkpoints.append({
            "label": label,
            "days_loaded": summary["days_loaded"],
            "proj": proj,
            # Redondeado y + 0.0: que una coincidencia exacta no salga "-0.0 %".
            "deviation": {k: round(_ratio(proj[k] - real_rows[k], real_rows[k]), 4) + 0.0
                          if (proj[k] is not None and real_rows[k]) else None
                          for k in proj},
        })
    rows = [
        ("galones", "Galones vendidos", "gal"),
        ("ventas_combustible", "Venta de combustible", "$"),
        ("utilidad_combustible", "Utilidad de combustible", "$"),
        ("ventas_cstore", "Venta C-Store (sin Lotto ni Kia/Toyota)", "$"),
        ("utilidad_cstore", "Utilidad C-Store", "$"),
    ]
    return {
        "real": real_rows,
        "source": real.get("archivo"),
        "loaded_at": real.get("cargado"),
        "rows": rows,
        "checkpoints": checkpoints,
        "cstore_margin_real": _ratio(real_rows["utilidad_cstore"], real_rows["ventas_cstore"]),
        "fuel_margin_gal_real": _ratio(real_rows["utilidad_combustible"], real_rows["galones"]),
    }
