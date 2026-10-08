"""
Productos de las facturas de proveedores -- pedido explícito del usuario
(2026-09-28): "crear un sistema organizado en el que extraigamos los
productos que compramos con su precio, para así utilizar esa info para ver
cómo van cambiando los costos de los mismos productos en cada compra y
también poder compararlo con el CMV... a qué proveedor se le compra ese
producto por su UPC". Lectura/extracción + cálculo, sin UI -- el guardado
vive en proveedores_db (tabla supplier_invoice_lines) y las pantallas en la
sección Proveedores de webapp.py.

H.T. Hackney trae texto real. CEC y Colonial (2026-10-04, pedido del
usuario: "mejorar la precisión de los OCR de las facturas de los
proveedores... así se puedan extraer los productos al igual que H.T.")
son escaneos y se leen por OCR, con más controles (ver más abajo). Gold
Coast Eagle y Red Bull (2026-10-05) son tickets impresos escaneados y se
leen por lectura múltiple con votación (sección "Tickets impresos"). El
resto de los proveedores todavía no -- ver los relevamientos del
2026-09-28, 2026-10-04 y 2026-10-05 en HISTORIAL.md. Cada proveedor nuevo
se suma en LINE_EXTRACTORS.

Regla de oro: el detalle de una factura se guarda solo si cierra al
centavo -- cantidad x neto = total de cada renglón, la suma de renglones =
INVOICE SUBTOTAL impreso, y subtotal + cargos = Total impreso. Si algo no
cierra, ValueError y no se guarda ningún renglón (nunca un detalle a medias).
"""

import hashlib
import itertools
import os
import re
from datetime import datetime

import pdfplumber

import ocr_utils
from ocr_utils import (
    crop_relative,
    ensure_cv2,
    ensure_pdfplumber,
    ensure_pytesseract,
    extract_largest_page_image,
    remove_grid_lines,
)

try:
    import cv2
    import numpy as np
    from PIL import Image, ImageFilter
except ImportError:  # sin OpenCV/Pillow: solo fallan los escaneos (CEC, Colonial, GCE, Red Bull), con un error claro (ensure_cv2)
    cv2 = None  # type: ignore[assignment]
    np = None  # type: ignore[assignment]
    Image = None  # type: ignore[assignment]
    ImageFilter = None  # type: ignore[assignment]

# Columnas de la grilla de H.T., en el orden de los 16 bloques de guiones
# que la factura imprime debajo del encabezado ("- ------- -------------- ...").
_HT_COLUMNS = [
    "flag", "qty", "upc", "pack", "size", "description", "item_no", "srp_pct",
    "srp", "srp_ext", "price", "allowance", "tax", "net", "ext", "tx",
]
# Un cartón de cigarrillos (size "CTN") son 10 paquetes, y el POS tiene el
# costo por paquete: Marlboro Gold $102.95 el cartón = $10.30 en el POS.
_UNITS_PER_CARTON = 10


def normalize_upc(value):
    """Solo dígitos, sin ceros adelante -- así cruza con el UPC del POS (CMV)."""
    digits = re.sub(r"\D", "", str(value or ""))
    return digits.lstrip("0")


def product_key(supplier_key, upc, item_no, description=None):
    """
    Clave de producto: el UPC; sin UPC, el Item # propio del proveedor; en los
    que se identifican por el nombre (NAME_KEYED_SUPPLIERS: J.J. Taylor no trae
    código, el SKU de Midtown es de relleno a veces), el nombre normalizado.
    """
    if upc:
        return upc
    if supplier_key in NAME_KEYED_SUPPLIERS:
        return _name_product_key(supplier_key, name_key(description))
    return f"item-{supplier_key}-{item_no}"


def _name_product_key(supplier_key, key):
    return f"name-{supplier_key}-{hashlib.sha1(key.encode('utf-8')).hexdigest()[:12]}"


def _container_alternatives(supplier_key, description):
    """
    Claves del mismo nombre con o sin la letra del envase del final ("... 1/24/16 C"):
    en J.J. Taylor esa letra sale ilegible seguido. Solo para buscar la compra
    anterior cuando el nombre exacto no aparece.
    """
    words = name_key(description).split(" ")
    if words and words[-1] in ("B", "C"):
        options = [" ".join(words[:-1])]
    elif words and words[-1].isdigit():
        options = [" ".join(words + [letter]) for letter in ("B", "C")]
    else:
        options = []
    return [_name_product_key(supplier_key, option) for option in options]


def _num(text):
    text = (text or "").replace(",", "").replace("$", "").strip()
    if not text:
        return None
    negative = text.endswith("-") or text.startswith("-")
    text = text.strip("-").strip()
    try:
        value = float(text)
    except ValueError:
        return None
    return -value if negative else value


def _page_rows(page):
    """Renglones de la página: palabras agrupadas por altura, de izquierda a derecha."""
    rows = []
    last_top = None
    for word in sorted(page.extract_words(), key=lambda w: (round(w["top"]), w["x0"])):
        top = round(word["top"])
        if last_top is not None and top - last_top <= 2:
            rows[-1].append(word)
        else:
            rows.append([word])
            last_top = top
    for row in rows:
        row.sort(key=lambda w: w["x0"])
    return rows


def _column_index(dash_ranges, word):
    center = (word["x0"] + word["x1"]) / 2

    def distance(index):
        start, end = dash_ranges[index]
        if start - 1 <= center <= end + 1:
            return 0
        return min(abs(center - start), abs(center - end))

    return min(range(len(dash_ranges)), key=distance)


def extract_ht_hackney_lines(pdf_path, invoices=None):
    """
    Renglones de producto de una factura de H.T. Hackney. Las columnas se
    leen por posición (los bloques de guiones del encabezado marcan dónde
    empieza y termina cada una), porque varias vienen vacías según el
    producto (SRP%, Alw, Tax, y a veces el UPC) y leer el texto corrido las
    correría de lugar.
    """
    invoice_no = None
    lines = []
    charges = {}
    subtotal = None
    total = None
    category = None

    with pdfplumber.open(pdf_path) as pdf:
        for page in pdf.pages:
            dash_ranges = None
            for row in _page_rows(page):
                text = " ".join(w["text"] for w in row)
                if invoice_no is None:
                    match = re.search(r"Invoice #:\s*(\d+)", text)
                    if match:
                        invoice_no = match.group(1)
                # El encabezado repite el total en cada página ("**For
                # Reference Purposes Only** Total: ..."): no es el renglón final.
                if "For Reference Purposes Only" in text:
                    continue
                # Cargos, subtotal y total se buscan en toda la página, no solo
                # debajo de la grilla: a veces el Total cae solo en una última
                # página que ya no repite la grilla de columnas.
                match = re.match(r"^(FUEL SURCHARGE|DELIVERY CHARGE)\s+([\d,.]+-?)$", text)
                if match:
                    charges[match.group(1)] = _num(match.group(2))
                    continue
                # Al final del renglón y no al principio: a veces viene pegado
                # a un aviso de retiro de producto de la FDA ("...RECALLS Total: 4995.51").
                # El monto puede venir sin cero adelante ("Tax: .66").
                match = re.search(r"(?:Tax:\s*([\d,]*\.\d{2}-?)\s+)?Total:\s*([\d,]*\.\d{2}-?)$", text)
                if match:
                    if match.group(1):
                        charges["TAX"] = _num(match.group(1))
                    total = _num(match.group(2))
                    continue
                match = re.search(r"\*\*\* INVOICE SUBTOTAL \*\*\*.*?([\d,]+\.\d{2}-?)$", text)
                if match:
                    subtotal = _num(match.group(1))
                    continue
                if dash_ranges is None:
                    if len(row) >= 14 and all(set(w["text"]) == {"-"} for w in row):
                        if len(row) != len(_HT_COLUMNS):
                            raise ValueError(
                                "la grilla de productos de H.T. cambió de formato "
                                f"({len(row)} columnas en vez de {len(_HT_COLUMNS)})."
                            )
                        dash_ranges = [(w["x0"], w["x1"]) for w in row]
                    continue

                cells = {name: [] for name in _HT_COLUMNS}
                for word in row:
                    cells[_HT_COLUMNS[_column_index(dash_ranges, word)]].append(word["text"])
                cells = {name: " ".join(parts) for name, parts in cells.items()}
                if (re.fullmatch(r"\d+", cells["qty"]) and re.fullmatch(r"\d{5,8}", cells["item_no"])
                        and _num(cells["ext"]) is not None and _num(cells["net"]) is not None):
                    lines.append(_ht_line(cells, category, len(lines) + 1))
                    continue
                # Encabezado de sección ("CANDY", "CIGS, FULL PRICE"): en
                # mayúsculas y sin números. Las notas ("ABOVE IS ...", "THE
                # ABOVE ITEM WAS ...") no son sección.
                if text.isupper() and not re.search(r"\d", text) and "ABOVE" not in text and "***" not in text:
                    category = text

    if not lines:
        raise ValueError("no se encontró ningún renglón de producto en la factura de H.T.")
    for line in lines:
        if abs(line["qty"] * line["net"] - line["ext"]) > 0.011:
            raise ValueError(
                f"el renglón {line['line_no']} ({line['description']}) no cierra: "
                f"{line['qty']:g} x {line['net']:.2f} ≠ {line['ext']:.2f}."
            )
    lines_total = round(sum(line["ext"] for line in lines), 2)
    if subtotal is None or abs(lines_total - subtotal) >= 0.01:
        raise ValueError(
            f"los renglones suman ${lines_total:,.2f} y el INVOICE SUBTOTAL de la factura es "
            f"{'$' + format(subtotal, ',.2f') if subtotal is not None else 'ilegible'}."
        )
    if total is None or abs(subtotal + sum(charges.values()) - total) >= 0.01:
        raise ValueError("el subtotal más los cargos (fuel, delivery, tax) no da el Total de la factura.")
    if invoice_no is None:
        raise ValueError("no se encontró el N° de invoice en la factura de H.T.")
    return {"invoice_no": invoice_no, "lines": lines, "subtotal": subtotal, "total": total}


def _ht_line(cells, category, line_no):
    pack = int(_num(cells["pack"].lstrip("*")) or 1)
    size = cells["size"]
    units = pack * (_UNITS_PER_CARTON if size.upper() == "CTN" else 1)
    net = _num(cells["net"])
    return {
        "line_no": line_no,
        "upc": normalize_upc(cells["upc"]),
        "item_no": cells["item_no"],
        "description": cells["description"],
        "category": category,
        "qty": _num(cells["qty"]),
        "pack": pack,
        "size": size,
        "units": units,
        "price": _num(cells["price"]),
        "allowance": _num(cells["allowance"]),
        "tax": _num(cells["tax"]),
        "net": net,
        "ext": _num(cells["ext"]),
        "unit_cost": net / units,
        "srp": _num(cells["srp"]),
    }


# ---------------------------------------------------------------------------
# Facturas ESCANEADAS: CEC y Colonial (2026-10-04)
# ---------------------------------------------------------------------------
# Mismo contrato y misma regla de oro que H.T. Hackney: el detalle se guarda
# solo si cierra al centavo. Como estas facturas son escaneos (OCR), además:
# - cada renglón tiene que cerrar (cantidad x neto = total del renglón);
# - la cantidad se toma del OCR, y solo si ese número no cierra (las líneas
#   de la grilla lo ensucian: "|42" por 12) se deduce de total / precio,
#   siempre que dé un número entero Y lo confirme otra fuente (el conteo
#   impreso al pie o una relectura de esa celda sola);
# - los renglones tienen que sumar el subtotal/total impreso (o, en
#   Colonial, el importe que la app ya leyó del encabezado);
# - el UPC, cuando lo hay, tiene que pasar el dígito verificador.
# Validado sobre los PDFs reales del Drive (2025-2026): CEC 83 de 83, todas
# cerrando contra el Subtotal; Colonial 28 de 80 con el BALANCE DUE nuevo
# de proveedores.py (el resto da ValueError, y ninguna de las que pasan
# tiene un valor mal leído en la revisión a ojo).

_TOLERANCE = 0.005  # medio centavo: precio de 4 decimales x cantidad, redondeado


# ---------------------------------------------------------------------------
# Helpers comunes de OCR
# ---------------------------------------------------------------------------

def _pytesseract():
    ensure_pytesseract()
    return ocr_utils.pytesseract


def _page_images(pdf_path):
    """Imagen escaneada (ya orientada) de cada página, salteando las que no tienen."""
    ensure_pdfplumber()
    with pdfplumber.open(pdf_path) as pdf:
        images = [extract_largest_page_image(page) for page in pdf.pages]
    return [image for image in images if image is not None]


def _ocr_rows(image, config="--psm 6", box=None, scale=1, clean=False, prep=None):
    """
    Renglones del OCR con la posición de cada palabra. psm 6 lee la página
    como un único bloque, así cada renglón de la tabla sale entero (con psm 3
    Tesseract separa las columnas en bloques y desarma los renglones).
    Con box/scale se lee solo ese recorte agrandado, pero las posiciones
    vuelven en coordenadas de la página original. clean borra el sombreado
    de puntitos (renglones alternados de Colonial) antes de leer, de dos
    formas distintas (dos lecturas independientes): "blur" (desenfoque leve
    y corte fijo: los puntitos quedan más claros que la letra) o "median"
    (filtro de mediana y blanco/negro automático). prep: otro arreglo de la
    imagen ya agrandada (las pasadas de los tickets, más abajo).
    """
    pytesseract = _pytesseract()
    x_off = y_off = 0
    source = image
    if box is not None:
        x_off, y_off = max(0, int(box[0])), max(0, int(box[1]))
        source = image.crop((x_off, y_off, min(image.width, int(box[2])), min(image.height, int(box[3]))))
        source = source.convert("L")
    if scale != 1:
        source = source.convert("L").resize((source.width * scale, source.height * scale))
    if clean:
        ensure_cv2()
        gray = np.array(source.convert("L"))
        if clean == "blur":
            gray = np.where(cv2.GaussianBlur(gray, (3, 3), 0) < 120, 0, 255).astype(np.uint8)
        else:
            gray = cv2.medianBlur(gray, 7 if scale >= 3 else 5)
            _, gray = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)
        source = Image.fromarray(gray)
    if prep is not None:
        source = prep(source)
    data = pytesseract.image_to_data(source, config=config, output_type=pytesseract.Output.DICT)
    rows = {}
    for i, text in enumerate(data["text"]):
        text = (text or "").strip()
        if not text:
            continue
        key = (data["block_num"][i], data["par_num"][i], data["line_num"][i])
        rows.setdefault(key, []).append({
            "text": text,
            "x0": x_off + data["left"][i] / scale,
            "x1": x_off + (data["left"][i] + data["width"][i]) / scale,
            "top": y_off + data["top"][i] / scale,
            "bottom": y_off + (data["top"][i] + data["height"][i]) / scale,
        })
    result = []
    for words in rows.values():
        words.sort(key=lambda w: w["x0"])
        result.append(words)
    result.sort(key=lambda ws: min(w["top"] for w in ws))
    return result


def _row_text(words):
    return " ".join(w["text"] for w in words)


def _ocr_cell(image, box, whitelist="0123456789", psm=7):
    """Relee una celda sola, sin las líneas de la grilla (que el OCR confunde con 1, 4 o |)."""
    pytesseract = _pytesseract()
    x0, top, x1, bottom = box
    crop = image.crop((max(0, int(x0)), max(0, int(top)), min(image.width, int(x1)), min(image.height, int(bottom))))
    if crop.width < 4 or crop.height < 4:
        return ""
    cleaned = remove_grid_lines(crop, upscale=3)
    config = f"--psm {psm}" + (f" -c tessedit_char_whitelist={whitelist}" if whitelist else "")
    return pytesseract.image_to_string(cleaned, config=config).strip()


_DIGIT_LOOKALIKES = str.maketrans({"O": "0", "o": "0", "g": "9", "S": "5", "s": "5", "l": "1", "I": "1", "B": "8"})


def _fix_digits(text):
    """Letras que el OCR confunde con dígitos (g.0000 -> 9.0000); solo para celdas que son números."""
    return text.translate(_DIGIT_LOOKALIKES)


def _money(text):
    """'$1,234.56' / '1.234,56' mal leído / '34.9000' / '9,.0000' -> float, o None."""
    text = (text or "").replace("$", "").replace(" ", "").strip()
    if not text:
        return None
    # El OCR a veces lee la coma de miles como punto o el punto decimal como
    # coma: el separador decimal es siempre el ÚLTIMO, seguido de 2 o 4 dígitos.
    match = re.fullmatch(r"(-?)([\d.,]*?)[.,](\d{2}|\d{4})", text)
    if not match:
        return None
    sign, whole, decimals = match.groups()
    whole = re.sub(r"[.,]", "", whole) or "0"
    value = float(f"{whole}.{decimals}")
    return -value if sign else value


def _gtin_ok(digits):
    """Dígito verificador de un UPC-A/EAN (GTIN 8, 12, 13 o 14): detecta cualquier dígito mal leído."""
    if not re.fullmatch(r"\d{8}|\d{12,14}", digits or ""):
        return False
    body, check = digits[:-1], int(digits[-1])
    total = sum(int(d) * (3 if i % 2 == 0 else 1) for i, d in enumerate(reversed(body)))
    return (10 - total % 10) % 10 == check


def _closes(qty, net, ext):
    return abs(qty * net - ext) <= _TOLERANCE + 1e-9


def _whole_qty(net, ext):
    """La única cantidad entera que cierra total / precio, o None."""
    if not net or net <= 0:
        return None
    qty = round(ext / net)
    return qty if qty > 0 and _closes(qty, net, ext) else None


def _clean_token(text):
    return text.strip("|[](){}_—–~=‘’'\"`«»,;:!*")


def _line(line_no, *, upc="", item_no="", description="", category=None, qty, pack, size, units,
          price, allowance=None, tax=None, net, ext, srp=None):
    """Renglón con exactamente las claves de proveedores_productos._ht_line."""
    return {
        "line_no": line_no,
        "upc": normalize_upc(upc),
        "item_no": item_no,
        "description": description,
        "category": category,
        "qty": qty,
        "pack": pack,
        "size": size,
        "units": units,
        "price": price,
        "allowance": allowance,
        "tax": tax,
        "net": net,
        "ext": ext,
        "unit_cost": net / units,
        "srp": srp,
    }


def _check_sum(lines, printed, label):
    lines_total = round(sum(line["ext"] for line in lines), 2)
    if printed is None or abs(lines_total - printed) >= 0.01:
        raise ValueError(
            f"los renglones suman ${lines_total:,.2f} y el {label} de la factura es "
            f"{'$' + format(printed, ',.2f') if printed is not None else 'ilegible'}."
        )


def _median(values):
    values = sorted(values)
    return values[len(values) // 2] if values else None


def _upc_candidate(digits, lengths=(12,)):
    """
    UPC válido dentro de una tira de dígitos leída por OCR. La línea de la
    grilla a veces se lee como un dígito más pegado ("1709215205322"): se
    prueban las ventanas del largo esperado y vale solo si UNA pasa el
    dígito verificador.
    """
    found = set()
    for length in lengths:
        for start in range(0, len(digits) - length + 1):
            window = digits[start:start + length]
            if _gtin_ok(window):
                found.add(window)
    return found.pop() if len(found) == 1 else None


def _reread_digits(image, box, min_len, lengths=(12,)):
    """Relee una celda con solo dígitos y devuelve un UPC válido, o None."""
    digits = re.sub(r"\D", "", _ocr_cell(image, box))
    return _upc_candidate(digits, lengths) if len(digits) >= min_len else None


def _invoice_no_from_crop(image, box, pattern):
    """Relee el N° de factura en un recorte agrandado al doble (el OCR de página entera a veces pierde un dígito)."""
    pytesseract = _pytesseract()
    x0, top, x1, bottom = (int(v) for v in box)
    crop = image.crop((max(0, x0), max(0, top), min(image.width, x1), min(image.height, bottom)))
    crop = crop.convert("L").resize((crop.width * 2, crop.height * 2))
    for psm in (7, 6):
        text = pytesseract.image_to_string(crop, config=f"--psm {psm}")
        text = re.sub(r"(?<=\d) (?=\d)", "", text)
        match = re.search(pattern, text)
        if match:
            return match.group(1)
    return None


def _resolve_quantities(rows, printed_sum, supplier):
    """
    Cantidad y total de cada renglón (claves price, ext, qty_token). En orden:
    1) la cantidad leída cierra (cantidad x precio = total): listo;
    2) si no, total / precio da un entero: cantidad deducida (qty_derived),
       que después tiene que confirmar el conteo impreso o una relectura;
    3) si UN solo renglón tiene el total ilegible o mal leído ("$27200.",
       "204.90" por 204.00) pero su cantidad se leyó limpia, su total es el
       subtotal impreso menos los demás renglones, y vale solo si eso da
       exactamente cantidad x precio.
    Si algo no cierra, ValueError.
    """
    bad = []
    for row in rows:
        row["qty_derived"] = False
        price, ext, qty = row["price"], row["ext"], row["qty_token"]
        if price is None:
            raise ValueError(f"no se pudo leer el precio del renglón {row['line_no']} de {supplier}.")
        if qty is not None and ext is not None and _closes(qty, price, ext):
            row["qty"] = qty
            continue
        whole = _whole_qty(price, ext) if ext is not None else None
        if whole is not None:
            row["qty"] = whole
            row["qty_derived"] = True
            continue
        bad.append(row)
    if not bad:
        return
    row = bad[0]
    if len(bad) == 1 and printed_sum is not None and row["qty_token"]:
        rest = sum(other["ext"] for other in rows if other is not row)
        ext = round(printed_sum - rest, 2)
        if _closes(row["qty_token"], row["price"], ext):
            row["qty"] = row["qty_token"]
            row["ext"] = ext
            return
    raise ValueError(
        f"el renglón {row['line_no']} de {supplier} ({row.get('description', '')}) no cierra: "
        "cantidad x precio no da el total del renglón."
    )


def _qty_reread_ok(image, row):
    """
    Relee la celda de cantidad sola (sin grilla, a la altura del precio, que
    siempre se lee limpio) y confirma la cantidad deducida. Se prueba como
    renglón (psm 7) y como palabra suelta (psm 8).
    """
    box = (max(0, row["row_left"] - 40), row["cell_top"], row["qty_right"] - 4, row["cell_bottom"])
    for psm in (7, 8):
        text = _ocr_cell(image, box, psm=psm)
        if text.isdigit() and int(text) == row["qty"]:
            return True
    return False


# ---------------------------------------------------------------------------
# CEC (Chinook Enterprises Corp.)
# ---------------------------------------------------------------------------

# CEC vende solo por cartón: cigarrillos 305's y filter cigars 305's, ambos en
# cartones de 10 paquetes (el cuadro "Cartons" del pie los cuenta igual). El
# UPC de la factura es el del cartón, y el POS lo vende como cartón (CMV: 305
# GOLD $9.00, 305 PURPLE KING CARTON $34.90; el paquete suelto tiene otro
# UPC), así que el costo por unidad es el precio del cartón.
_CEC_UNITS_PER_CARTON = 1
# Precio por cartón con 4 decimales; el OCR a veces mete ",." o lee el 9 como "g".
_CEC_PRICE = re.compile(r"[\dgOoSlIB]{1,4}[.,]{1,2}\d{4}(?!\d)")
_CEC_MONEY = re.compile(r"\d[\d.,]*[.,]\d{2}(?!\d)")
_CEC_ITEM = re.compile(r"(?:FC\d{3}[A-Z]|\d{4})(?![\dA-Za-z])")


def _cec_kind(description, item_no):
    """
    'cig' (cigarrillos), 'other' (filter cigars, el "Others" del cuadro
    Cartons) o None si no se reconoce -- en ese caso no se guarda (no se
    sabe cuántas unidades trae).
    """
    text = description.upper()
    if "CIGAR" in text or item_no.startswith("FC"):
        return "other"
    # Los dos tipos son cartones de 10: confundirlos no cambia las unidades
    # (y el cuadro Cartons lo detectaría); lo que no es 305's no se guarda.
    if re.search(r"BOX|SOFT|5['’]S", text) or re.fullmatch(r"\d{4}", item_no):
        return "cig"
    return None


def extract_cec_lines(pdf_path, invoices=None):
    """
    Renglones de una factura de CEC: QTY, Product, Description, UPC, Price
    (4 decimales, por cartón) y Subtotal. Cada renglón se ancla en el precio
    de 4 decimales y el UPC (12 dígitos con verificador). La cantidad y el
    código de producto son las celdas más sucias (la línea de la grilla o
    el tilde del chofer pegados: "|42" por 12, "[4", "v//s08"): el código es
    solo informativo (la clave del producto es el UPC) y la cantidad, si no
    cierra contra precio x total, se deduce de total / precio y tiene que
    confirmarla el cuadro "Cartons" del pie (cigarrillos / otros / total) o
    una relectura de esa celda sola.

    Si la lectura de la página entera no cierra (un renglón que el OCR pegó
    al encabezado de la tabla, por ejemplo), se relee solo la tabla al doble
    de tamaño y se vuelve a validar todo; si tampoco cierra, ValueError.
    """
    images = _page_images(pdf_path)
    if not images:
        raise ValueError("no se encontró la imagen escaneada de la factura de CEC.")
    image = images[0]
    info = _cec_parse(_ocr_rows(image))
    try:
        return _cec_finish(image, info)
    except ValueError as first_error:
        if not info["rows"]:
            raise
        tops = [row["top"] for row in info["rows"]]
        height = _median([row["bottom"] - row["top"] for row in info["rows"]])
        bottom = info["subtotal_bottom"] or max(row["bottom"] for row in info["rows"]) + 3 * height
        retry = _cec_parse(_ocr_rows(image, box=(0, min(tops) - 3 * height, image.width, bottom + 4), scale=2))
        for key in ("invoice_no", "inv_word", "subtotal", "cartons", "total"):
            if retry[key] is None:
                retry[key] = info[key]
        try:
            return _cec_finish(image, retry)
        except ValueError:
            raise first_error from None


def _cec_parse(rows):
    """Encabezado, pie y renglones de producto (sin validar) de una lectura OCR."""
    info = {"invoice_no": None, "inv_word": None, "subtotal": None, "subtotal_bottom": None,
            "cartons": None, "total": None, "rows": []}
    for words in rows:
        text = _row_text(words)
        if info["invoice_no"] is None and info["inv_word"] is None:
            info["inv_word"] = next((w for w in words if w["text"].startswith("Inv")), None)
            if info["inv_word"] is not None:
                match = re.search(r"Inv\s*#?\s*(\d{7})\b", re.sub(r"(?<=\d) (?=\d)", "", text))
                info["invoice_no"] = match.group(1) if match else None
        match = re.search(r"Subtotal\W*(\d[\d.,]*[.,]\d{2})(?!\d)", text)
        if match:
            info["subtotal"] = _money(match.group(1))
            info["subtotal_bottom"] = max(w["bottom"] for w in words)
            continue
        match = re.search(r"Cartons?\W+(\d+)\W+(\S+?)\W+(\d+)\b", text)
        if match:
            cig, other_text, all_ = int(match.group(1)), match.group(2), int(match.group(3))
            other = int(re.sub(r"\D", "", other_text)) if re.search(r"\d", other_text) else None
            if other is None and cig <= all_:
                other = all_ - cig  # "(¢]": el 0 de "Others" ilegible
            info["cartons"] = (cig, other, all_) if other is not None and cig + other == all_ else None
            continue
        match = re.search(r"\bTotal\W*(\d[\d.,]*[.,]\d{2})(?!\d)", text)
        if match:
            info["total"] = _money(match.group(1))
            continue
        price_index = max((i for i, w in enumerate(words) if _CEC_PRICE.search(w["text"])), default=None)
        if price_index is not None:
            info["rows"].append(_cec_row(words, price_index, len(info["rows"]) + 1))
    return info


def _cec_finish(image, info):
    """Valida una lectura: N° de invoice, UPCs, cantidades, Subtotal y cuadro Cartons."""
    parsed = info["rows"]
    invoice_no = info["invoice_no"]
    if invoice_no is None:
        # Misma lectura que el encabezado de la app (página entera, psm 3):
        # el psm 6 a veces pierde un dígito del N° ("Inv #1/64904").
        text = _pytesseract().image_to_string(image)
        match = re.search(r"Inv\s*#\s*([\d\s]+?)\n", text)
        digits = re.sub(r"\s+", "", match.group(1)) if match else ""
        invoice_no = info["invoice_no"] = digits if re.fullmatch(r"\d{7}", digits) else None
    if invoice_no is None:
        inv_word = info["inv_word"]
        # Recorte angosto: a la derecha suele haber un número escrito a mano.
        box = ((inv_word["x0"] - 10, inv_word["top"] - 15, inv_word["x0"] + image.width * 0.24, inv_word["bottom"] + 15)
               if inv_word is not None else (image.width * 0.55, 0, image.width, image.height * 0.08))
        invoice_no = info["invoice_no"] = _invoice_no_from_crop(image, box, r"(?:Inv\s*#?\s*)?(\d{7})\b")
    if invoice_no is None:
        raise ValueError("no se encontró el N° de invoice en la factura de CEC.")
    if not parsed:
        raise ValueError("no se encontró ningún renglón de producto en la factura de CEC.")
    subtotal = info["subtotal"]

    # UPC ilegible (un dígito mal o la celda vacía): se relee la celda sola,
    # entre el borde de la columna UPC (tomado de los otros renglones) y el precio.
    upc_x = _median([row["upc_x"] for row in parsed if row["upc_x"] is not None])
    for row in parsed:
        if row["upc"] is None:
            left = row["upc_x"] if row["upc_x"] is not None else upc_x
            if left is not None:
                row["upc"] = _reread_digits(
                    image, (left - 12, row["cell_top"], row["price_x"] - 4, row["cell_bottom"]), 11
                )
        if row["upc"] is None:
            raise ValueError(f"no se pudo leer el UPC del renglón {row['line_no']} de CEC.")

    _resolve_quantities(parsed, subtotal, "CEC")
    lines = []
    for row in parsed:
        kind = _cec_kind(row["description"], row["item_no"])
        # El código de producto no lo valida ninguna cuenta ("(3403" por
        # 3103): se guarda solo si una relectura de la celda sola coincide.
        if row["item_no"] and _ocr_cell(
            image, (row["item_x"] - 6, row["cell_top"], row["desc_x"] - 4, row["cell_bottom"]),
            whitelist="ABCDEFGHIJKLMNOPQRSTUVWXYZ0123456789",
        ) != row["item_no"]:
            row["item_no"] = ""
        if kind is None:
            raise ValueError(
                f"el renglón {row['line_no']} de CEC no es un producto conocido ({row['description']})."
            )
        row["kind"] = kind
        lines.append(_line(
            row["line_no"], upc=row["upc"], item_no=row["item_no"], description=row["description"],
            category="CIGARETTES" if kind == "cig" else "FILTER CIGARS",
            qty=row["qty"], pack=_CEC_UNITS_PER_CARTON, size="CTN", units=_CEC_UNITS_PER_CARTON,
            price=row["price"], net=row["price"], ext=row["ext"],
        ))
    _check_sum(lines, subtotal, "Subtotal")

    # Cuadro "Cartons" (cigarrillos, otros, total): confirma las cantidades.
    cartons = info["cartons"]
    if cartons is not None:
        cig = sum(row["qty"] for row in parsed if row["kind"] == "cig")
        other = sum(row["qty"] for row in parsed if row["kind"] == "other")
        if (cig, other) != cartons[:2]:
            raise ValueError(
                f"los renglones suman {cig} cartones de cigarrillos y {other} de otros, y el cuadro "
                f"Cartons de la factura dice {cartons[0]} y {cartons[1]}."
            )
    for row in parsed:
        if row["qty_derived"] and cartons is None and not _qty_reread_ok(image, row):
            raise ValueError(
                f"la cantidad del renglón {row['line_no']} ({row['description']}) de CEC no se pudo "
                "leer con seguridad."
            )
    # El "Total" del pie es igual al Subtotal en todas las facturas vistas.
    total = info["total"]
    if total is not None and abs(total - subtotal) >= 0.01:
        total = None
    return {"invoice_no": invoice_no, "lines": lines, "subtotal": subtotal,
            "total": total if total is not None else subtotal}


def _cec_row(words, price_index, line_no):
    price_word = words[price_index]
    price = _money(_fix_digits(_CEC_PRICE.search(price_word["text"]).group(0)))
    exts = _CEC_MONEY.findall(" ".join(w["text"] for w in words[price_index + 1:]))
    ext = _money(exts[-1]) if exts else None
    if ext is None:
        # "$27200": el punto decimal no salió; los centavos son los 2 últimos
        # dígitos (las cuentas del renglón y del Subtotal lo confirman o no).
        match = re.search(r"\$\s*(\d{3,7})(?![\d.,])", " ".join(w["text"] for w in words[price_index + 1:]))
        ext = int(match.group(1)) / 100 if match else None

    # UPC: la palabra más a la derecha (antes del precio) con 11 o más dígitos.
    upc = None
    upc_index = None
    for i in range(price_index - 1, -1, -1):
        digits = re.sub(r"\D", "", words[i]["text"])
        if len(digits) >= 11:
            upc_index = i
            upc = _upc_candidate(digits)
            break
    left_words = words[:upc_index] if upc_index is not None else words[:price_index]

    # Cantidad, código y descripción. "40-3103": cantidad y código pegados.
    tokens = []
    for word in left_words:
        # "40-3103": cantidad y código pegados; "3306.~=S305's-": código y descripción.
        text = re.sub(r"(?<=\d)[-—](?=\d{4})", " ", word["text"])
        text = re.sub(r"(?<=[^\s\d])(?=(?:[38]05|05|5)['’]s)", " ", text)
        for part in text.split():
            part = _clean_token(part.strip("."))
            if part:
                tokens.append((part, word))
    desc_index = next((i for i, (t, _w) in enumerate(tokens) if re.search(r"5['’]s", t)), None)
    item_index = next(
        (i for i, (t, _w) in enumerate(tokens[:desc_index]) if _CEC_ITEM.match(t)), None
    )
    if desc_index is None:
        desc_index = item_index + 1 if item_index is not None else len(tokens)
    item_no = _CEC_ITEM.match(tokens[item_index][0]).group(0) if item_index is not None else ""
    description = " ".join(t for t, _w in tokens[desc_index:] if re.search(r"\w", t))
    qty_end = item_index if item_index is not None else desc_index
    qty_token = next((t for t, _w in tokens[:qty_end] if re.fullmatch(r"\d{1,3}", t)), None)
    first_kept = tokens[qty_end][1] if qty_end < len(tokens) else price_word
    upc_x = words[upc_index]["x0"] if upc_index is not None else None
    desc_word = tokens[desc_index][1] if desc_index < len(tokens) else None
    return {
        "line_no": line_no,
        "price": price,
        "ext": ext,
        "qty_token": int(qty_token) if qty_token else None,
        "upc": upc,
        "upc_x": upc_x,
        "item_x": tokens[item_index][1]["x0"] if item_index is not None else None,
        "desc_x": desc_word["x0"] if desc_word is not None else (upc_x or price_word["x0"]),
        "item_no": item_no,
        "description": description,
        "price_x": price_word["x0"],
        "qty_right": first_kept["x0"],
        "row_left": min(w["x0"] for w in words),
        "top": min(w["top"] for w in words),
        "bottom": max(w["bottom"] for w in words),
        # Alto de las celdas para releerlas: el del precio, que siempre sale
        # limpio (el renglón entero se estira con los tildes a mano).
        "cell_top": price_word["top"] - 10,
        "cell_bottom": price_word["bottom"] + 10,
    }


# ---------------------------------------------------------------------------
# Colonial Wholesale Dist. LLC
# ---------------------------------------------------------------------------

# Colonial factura por UNIT (BX caja, CT cartón, CS bulto, PK paquete, RL
# rollo, EA unidad) y la columna SIZE dice cuántas unidades de venta trae:
# "10CT" = 10, "10PK" = cartón de 10 paquetes, "500CT" = 500 sorbetes. Esas
# son las unidades del renglón (BLACK & MILD 5PK en BX 10CT = 10 paquetes de
# 5, costo por paquete de 5). No trae UPC: la clave del producto es el ITEM#
# de Colonial, así que el ITEM# y el SIZE se guardan solo si dos lecturas
# distintas coinciden (la de la tabla tal cual y otra al doble de tamaño,
# sin sombreado; si no coinciden, desempata la celda sola).
_COLONIAL_SIZE = re.compile(r"^([\dSOILB]{1,4})(CT|PK)$")
_COLONIAL_MONEY = re.compile(r"^\d[\d,]*[.,]\d{2}$")
_COLONIAL_PROMO = re.compile(r"^[-.,]?\d{2}$|^\d+[.,]\d{2}$")
# Delivery fee que Colonial suma al pie (BALANCE DUE = renglones + fee).
_COLONIAL_FEES = (7.99, 9.99, 11.99)
_ONE_LIKE = str.maketrans({"l": "1", "I": "1", "i": "1", "]": "1", "[": "1", "!": "1", "|": "1",
                           "O": "0", "o": "0", "S": "5", "s": "5", "B": "8", "Z": "2", "z": "2"})


def _colonial_int(token):
    """Cantidad o ITEM# con las letras que el OCR confunde con dígitos; None si no es un número."""
    text = _clean_token(token or "").translate(_ONE_LIKE)
    return text if re.fullmatch(r"\d+", text) else None


def _colonial_size(token):
    """'10CT' / 'SOCT' (50CT) / '1OPK' -> ('10CT', 10), o None."""
    text = re.sub(r"[^\w]", "", token or "").upper().replace("CCT", "CT")  # "6&CT", "6CcT"
    match = _COLONIAL_SIZE.match(text)
    if not match:
        return None
    count = match.group(1).replace("L", "1").translate(_ONE_LIKE)
    if not count.isdigit() or int(count) == 0:
        return None
    return f"{int(count)}{match.group(2).upper()}", int(count)


def _without_vertical_lines(image):
    """
    Copia en grises sin las líneas verticales de la grilla. El borde
    izquierdo de la tabla queda pegado al ITEM# y el OCR lo lee como un "1"
    de más (7098 -> 17098): las lecturas de confirmación se hacen sin esas
    líneas, así ese error no se repite en las dos y no pasa como confirmado.
    """
    ensure_cv2()
    gray = np.array(image.convert("L"))
    ink = cv2.adaptiveThreshold(gray, 255, cv2.ADAPTIVE_THRESH_MEAN_C, cv2.THRESH_BINARY_INV, 25, 15)
    # Más largo que cualquier letra (~1,2% del alto) pero corto para seguir
    # las líneas que el escaneo dobla.
    kernel = cv2.getStructuringElement(cv2.MORPH_RECT, (1, max(30, gray.shape[0] // 60)))
    lines = cv2.morphologyEx(ink, cv2.MORPH_OPEN, kernel)
    lines = cv2.dilate(lines, cv2.getStructuringElement(cv2.MORPH_RECT, (5, 3)))
    # Se rellena con el fondo de alrededor, no con blanco: una franja blanca
    # en un renglón sombreado deja dos bordes que el OCR lee como un "1"
    # (15CT -> 115CT) en todas las lecturas sin grilla a la vez.
    gray = cv2.inpaint(gray, lines, 3, cv2.INPAINT_TELEA)
    return Image.fromarray(gray)


def _colonial_row(words):
    """
    Un renglón de producto de Colonial, sin validar: ITEM#, ORDERED,
    SHIPPED, UNIT, SIZE, DESCRIPTION, PRICE, PROMO y AMOUNT. Se lee de
    derecha a izquierda (los 3 importes del final) y la izquierda se ancla en
    el SIZE ("10CT"), que siempre viene después de UNIT. None si no es un
    renglón de producto ("Box 1 of 4", notas).
    """
    text = re.sub(r"(?<![\d.,])(\d+)\s+([.,]\d{2})(?!\d)", r"\1\2", _row_text(words))  # "158 .88"
    tokens = [t for t in re.split(r"[\s|]+", text) if t]
    nums = [re.sub(r"^[^\d.,-]+|[^\d]+$", "", t) for t in tokens]
    digit_idx = [i for i, t in enumerate(tokens) if re.search(r"\d", t)]
    if len(digit_idx) < 3:
        return None
    amt_i, promo_i, price_i = digit_idx[-1], digit_idx[-2], digit_idx[-3]
    if not _COLONIAL_MONEY.match(nums[amt_i]):
        return None
    amount = _money(nums[amt_i])
    promo = None
    if _COLONIAL_PROMO.match(nums[promo_i]):
        promo_text = nums[promo_i].lstrip("-")
        promo = _money(promo_text if re.search(r"[.,]", promo_text) else "." + promo_text)
    price = None
    if _COLONIAL_MONEY.match(nums[price_i]):
        price = _money(nums[price_i])
    elif re.fullmatch(r"\W?\d{2}", tokens[price_i]):
        price = int(tokens[price_i][-2:]) / 100  # ".99" leído "‘99" (la cuenta lo confirma o no)
    if price is None and _COLONIAL_MONEY.match(nums[promo_i]):
        # Promo ilegible y pegada: "14.99 =| 14.99" -> el penúltimo es el precio.
        price, price_i, promo = _money(nums[promo_i]), promo_i, None
    if price is None:
        return None

    left = tokens[:price_i]
    # El SIZE va antes de la descripción: no se busca después de una palabra
    # ("BACKWOOD 5PK SWEET" no es un SIZE 5PK).
    first_word = next((i for i, t in enumerate(left) if re.search(r"[A-Za-z]{3}", t)), len(left))
    # Candidatos a SIZE antes de la primera palabra de la descripción. Si
    # hay dos ("2CT 10PK"), el primero es SHIPPED y UNIT pegados. Si hay uno
    # solo con el mismo número que la cantidad de al lado y sin UNIT antes
    # ("9059 2 2CT EDGEFIELD"), puede ser eso mismo: no se toma.
    candidates = [i for i in range(1, min(len(left), 8, first_word + 1)) if _colonial_size(left[i])]
    size_i = candidates[-1] if candidates else None
    if size_i is not None and len(candidates) == 1:
        count = _colonial_size(left[size_i])[1]
        previous = _colonial_int(left[size_i - 1])
        has_unit = any(re.fullmatch(r"[A-Za-z]{2}", _clean_token(t)) for t in left[:size_i])
        if previous is not None and int(previous) == count and not has_unit:
            size_i = None
    if size_i is None:
        size_text, before, desc_tokens = None, left[:4], left[4:]
    else:
        size_text = _colonial_size(left[size_i])[0]
        before = left[:size_i]
        desc_tokens = left[size_i + 1:]
    numbers = [_colonial_int(t) for t in before]
    item_i = next((i for i, n in enumerate(numbers) if n and 2 <= len(n) <= 6), None)
    item = numbers[item_i] if item_i is not None else None
    # La cantidad que cuenta es SHIPPED (la última antes de UNIT); si solo se
    # leyó una de las dos, igual se prueba: la cuenta del renglón decide.
    qtys = [n for n in numbers[item_i + 1:] if n is not None and len(n) <= 3] if item_i is not None else []
    shipped = int(qtys[-1]) if qtys else None
    description = " ".join(t for t in desc_tokens if re.search(r"\w", t))
    description = re.sub(r"^[^\w]+", "", description)
    if size_text is None and not (item and re.search(r"[A-Za-z]{3}", description)):
        return None  # el pie (ORD/SHIP/FEE/BALANCE) o una nota con números, no un producto
    item_word = next((w for w in words if _colonial_int(w["text"]) == item), None) if item else None
    size_word = next((w for w in words if size_text and _colonial_size(w["text"])), None)
    return {
        "item": item,
        "shipped": shipped,
        "size": size_text,
        "description": description,
        "price": price,
        "promo": promo,
        "amount": amount,
        "y": (min(w["top"] for w in words) + max(w["bottom"] for w in words)) / 2,
        "row_left": min(w["x0"] for w in words),
        "item_x1": item_word["x1"] if item_word else None,
        "size_x": (size_word["x0"], size_word["x1"]) if size_word else None,
        "cell_top": min(w["top"] for w in words) - 6,
        "cell_bottom": max(w["bottom"] for w in words) + 6,
    }


def _colonial_qty(row):
    """
    Cantidad y neto de un renglón: SHIPPED x (PRICE - PROMO) = AMOUNT. Si la
    promo no se leyó, se prueba sin promo (la cuenta lo confirma o no). Si
    SHIPPED no cierra, la cantidad se deduce del total (qty_derived) y hay
    que confirmarla después.
    """
    nets = []
    if row["promo"] is not None:
        nets.append(round(row["price"] - row["promo"], 2))
    if row["promo"] != 0:
        nets.append(row["price"])
    for net in nets:
        if row["shipped"] and net > 0 and _closes(row["shipped"], net, row["amount"]):
            return row["shipped"], net, False
    for net in nets:
        qty = _whole_qty(net, row["amount"]) if net > 0 else None
        if qty is not None:
            return qty, net, True
    return None, None, None


def _colonial_footer(image, rows):
    """
    Pie de la última página: ORD, SHIP, DELIVERY FEE y BALANCE DUE, en un
    recuadro que el OCR de página no lee (la grilla pegada a los números):
    se recorta la fila de valores, se borra la grilla y se prueban dos
    lecturas. Devuelve [(ship, fee, balance), ...] -- lecturas posibles, que
    valen solo si cierran con los renglones -- o [].
    """
    label = next((ws for ws in rows if re.search(r"DELIVER|BALANCE", _row_text(ws), re.I)), None)
    customer = next((ws for ws in rows if re.search(r"Cust\w*\W*Sig", _row_text(ws), re.I)), None)
    if label is None and customer is None:
        return []
    height = _median([w["bottom"] - w["top"] for ws in rows for w in ws if re.search(r"\d", w["text"])]) or 30
    if label is not None:
        top = max(w["bottom"] for w in label) + 2
        bottom = min(w["top"] for w in customer) - 2 if customer is not None else top + 3 * height
    else:
        bottom = min(w["top"] for w in customer) - 2
        top = bottom - 3 * height
    crop = image.crop((0, int(max(0, top)), image.width, int(min(image.height, bottom))))
    if crop.height < 8:
        return []
    cleaned = remove_grid_lines(crop, upscale=3)
    readings = []
    for psm in (6, 11):
        text = _pytesseract().image_to_string(cleaned, config=f"--psm {psm}")
        numbers = re.findall(r"\d[\d,]*[.,]\d{2}(?!\d)|\b\d{1,3}\b", text)
        ints = [int(n) for n in numbers if re.fullmatch(r"\d{1,3}", n)]
        money = [_money(n) for n in numbers if not re.fullmatch(r"\d{1,3}", n)]
        if len(ints) >= 2 and len(money) == 2:
            readings.append((ints[1], money[0], money[1]))
        elif len(ints) >= 2 and len(money) == 1:
            readings.append((ints[1], 0.0, money[0]))
    return readings


def extract_colonial_lines(pdf_path, invoices=None):
    """
    Renglones de una factura de Colonial (1 o 2 páginas). El OCR de página
    entera lee mal la tabla (renglones alternados con sombreado de puntitos),
    así que la tabla se lee cuatro veces (ver _colonial_page) y cada renglón
    junta sus lecturas:
    - ITEM# y SIZE valen solo con las reglas de _colonial_vote (no alcanza
      con "dos de tres"); si no, ValueError;
    - la cantidad es SHIPPED si cierra la cuenta SHIPPED x (PRICE - PROMO) =
      AMOUNT; si hay que deducirla del total, otra lectura o el pie tienen
      que confirmarla.

    Totales: los renglones más el delivery fee tienen que dar al centavo el
    importe que la app ya leyó del encabezado (`invoices`). El fee subió con
    el tiempo ($7.99 en 2025, $9.99 en mar-2026, $11.99 desde jun-2026): si
    el pie se puede leer (SHIP, fee y BALANCE DUE, que tienen que cerrar con
    los renglones), vale el fee leído; si no, uno de esos tres. Sin
    `invoices`, solo vale si el pie se leyó y cierra.
    """
    images = _page_images(pdf_path)
    if not images:
        raise ValueError("no se encontró la imagen escaneada de la factura de Colonial.")
    # N° de invoice: misma lectura que el encabezado de la app (cajita de
    # arriba a la derecha de la primera página).
    top_text = _pytesseract().image_to_string(crop_relative(images[0], 0.68, 0.0, 1.0, 0.22))
    match = re.search(r"\b(\d{7})\b", top_text)
    if not match:
        raise ValueError("no se encontró el N° de invoice en la factura de Colonial.")
    invoice_no = match.group(1)

    groups = []
    footer_readings = []
    for page in images:
        page_groups, page_footer = _colonial_page(page)
        groups.extend(page_groups)
        footer_readings = page_footer or footer_readings
    if not groups:
        raise ValueError("no se encontró ningún renglón de producto en la factura de Colonial.")

    lines = []
    unconfirmed_qty = []
    for group in groups:
        line_no = len(lines) + 1
        qty, net, amount, confirmed = _colonial_group_qty(group)
        if qty is None:
            description = next((r["description"] for r in group if r["description"]), "")
            raise ValueError(
                f"el renglón {line_no} de Colonial ({description}) no cierra: "
                "cantidad x precio no da el total del renglón."
            )
        if not confirmed:
            unconfirmed_qty.append(line_no)
        item = _colonial_vote(group, "item", line_no)
        size = _colonial_vote(group, "size", line_no)
        units = int(re.match(r"\d+", size).group(0))
        price = round(net + _colonial_promo_of(group, net), 2)
        lines.append(_line(
            line_no, item_no=item, description=_colonial_description(group),
            qty=qty, pack=units, size=size, units=units,
            price=price, allowance=round(price - net, 2) or None, net=net, ext=amount,
        ))

    lines_total = round(sum(line["ext"] for line in lines), 2)
    shipped_total = sum(line["qty"] for line in lines)
    # Pie leído (SHIP, fee y BALANCE DUE): vale la lectura que cierra con los renglones.
    footer = next(
        (f for f in footer_readings if f[0] == shipped_total and abs(lines_total + f[1] - f[2]) < 0.01), None
    )
    header_amounts = [inv["amount"] for inv in (invoices or []) if str(inv.get("invoice_no")) == invoice_no]
    if header_amounts:
        fees = [footer[1]] if footer else list(_COLONIAL_FEES)
        if not any(abs(lines_total + fee - amount) < 0.01 for fee in fees for amount in header_amounts):
            raise ValueError(
                f"los renglones suman ${lines_total:,.2f} y la factura es de ${header_amounts[0]:,.2f} "
                "(ni sumando el delivery fee da)."
            )
    elif footer is None:
        raise ValueError(
            f"los renglones suman ${lines_total:,.2f} y no hay un total de la factura para controlarlos."
        )
    if unconfirmed_qty and footer is None:
        raise ValueError(
            f"la cantidad del renglón {unconfirmed_qty[0]} de Colonial no se pudo leer con seguridad."
        )
    total = footer[2] if footer else next(
        a for a in header_amounts if any(abs(lines_total + fee - a) < 0.01 for fee in _COLONIAL_FEES)
    )
    return {"invoice_no": invoice_no, "lines": lines, "subtotal": lines_total, "total": total}


def _colonial_page(page):
    """
    Las cuatro lecturas de la tabla de una página, agrupadas por renglón, y
    las lecturas posibles del pie. Lecturas (medido contra el consenso en 20
    facturas): 0 = hoja tal cual (la mejor para SIZE, pero pega el borde de
    la tabla al ITEM#), 1 = sin líneas verticales (la mejor para ITEM#),
    2 = sin líneas al doble y sin sombreado, 3 = hoja tal cual al doble.
    """
    clean_page = _without_vertical_lines(page)
    lower = (0, page.height * 0.25, page.width, page.height)
    first_rows = _ocr_rows(page, box=lower)
    reads = [first_rows, _ocr_rows(clean_page, box=lower)]
    found = [r for rows in reads for r in (_colonial_row(ws) for ws in rows) if r]
    footer = _colonial_footer(page, first_rows)
    if not found:
        return [], footer
    height = _median([r["cell_bottom"] - r["cell_top"] for r in found])
    box = (0, min(r["cell_top"] for r in found) - height, page.width, max(r["cell_bottom"] for r in found) + height)
    reads.append(_ocr_rows(clean_page, box=box, scale=2, clean="blur"))
    reads.append(_ocr_rows(page, box=box, scale=2))

    # Renglones de las cuatro lecturas agrupados por altura (uno por lectura).
    groups = []
    for read_index, rows in enumerate(reads):
        for words in rows:
            row = _colonial_row(words)
            if row is None:
                continue
            row["read"] = read_index
            row["image"] = clean_page
            row["page"] = page
            group = min(groups, key=lambda g: abs(g[0]["y"] - row["y"]), default=None)
            if (group is not None and abs(group[0]["y"] - row["y"]) < height * 0.5
                    and all(r["read"] != read_index for r in group)):
                group.append(row)
            else:
                groups.append([row])
    groups.sort(key=lambda g: g[0]["y"])
    # El pie a veces sale en la lectura sin grilla como un renglón suelto de
    # números: "51 51 7.99 1436.71" (ORD, SHIP, DELIVERY FEE, BALANCE DUE).
    last_y = groups[-1][0]["y"] if groups else 0
    for rows in reads:
        for words in rows:
            if min(w["top"] for w in words) <= last_y:
                continue
            match = re.fullmatch(
                r"\W*(\d{1,3})\W+(\d{1,3})\W+(?:(\d+[.,]\d{2})\W+)?(\d[\d,]*[.,]\d{2})\W*", _row_text(words)
            )
            if match:
                fee = _money(match.group(3)) if match.group(3) else 0.0
                footer.append((int(match.group(2)), fee, _money(match.group(4))))
    return groups, footer


def _majority(values):
    """El valor que se repite al menos dos veces (el más repetido), o None."""
    counts = {}
    for value in values:
        if value is not None:
            counts[value] = counts.get(value, 0) + 1
    best = max(counts.items(), key=lambda kv: kv[1], default=(None, 0))
    return best[0] if best[1] >= 2 else None


def _colonial_group_qty(group):
    """
    (cantidad, neto, total, confirmada) de un renglón con varias lecturas.
    Vale la cantidad leída que cierra la cuenta (la que más lecturas
    repiten); si ninguna cierra, la deducida del total, y queda sin
    confirmar salvo que otra lectura haya leído esa misma cantidad.
    """
    closing = []
    derived = []
    amount = _majority([r["amount"] for r in group])
    for row in group:
        if amount is not None and row["amount"] != amount:
            continue  # la lectura que leyó otro total no cuenta
        qty, net, is_derived = _colonial_qty(row)
        if qty is None:
            continue
        (derived if is_derived else closing).append((qty, net, row["amount"]))
    for options, from_token in ((closing, True), (derived, False)):
        if not options:
            continue
        best = max(set(options), key=options.count)
        if options.count(best) == 1 and len(set(options)) > 1:
            return None, None, None, None  # lecturas que cierran con cuentas distintas
        qty = best[0]
        confirmed = from_token or sum(1 for r in group if r["shipped"] == qty) >= 1
        return best[0], best[1], best[2], confirmed
    return None, None, None, None


def _colonial_promo_of(group, net):
    """La promo (descuento por unidad) que explica el neto: la leída, o 0."""
    for row in group:
        if row["promo"] is not None and abs(row["price"] - row["promo"] - net) < 0.005:
            return row["promo"]
    return 0.0


def _colonial_description(group):
    """La descripción que más lecturas repiten; si no, la de la hoja tal cual (solo informativa)."""
    texts = [r["description"] for r in sorted(group, key=lambda r: r["read"]) if r["description"]]
    return _majority(texts) or (texts[0] if texts else "")


# Lectura principal de cada campo (la que nunca leyó mal en la medición) y
# de qué lecturas puede venir la confirmación.
_COLONIAL_PRIMARY = {"size": 0, "item": 1}


def _colonial_cell(row, field):
    """
    Relectura de la celda sola (sin grilla, agrandada): solo sirve para
    CONFIRMAR el valor de la lectura principal cuando las otras no lo
    leyeron. SIZE se relee de la hoja tal cual; ITEM#, de la hoja sin líneas.
    """
    if field == "size":
        if not row["size_x"]:
            return None
        box = (row["size_x"][0] - 10, row["cell_top"], row["size_x"][1] + 10, row["cell_bottom"])
        size = _colonial_size(_ocr_cell(row["page"], box, whitelist="0123456789CTPK"))
        return size[0] if size else None
    if not row["item_x1"]:
        return None
    box = (row["row_left"] - 20, row["cell_top"], row["item_x1"] + 10, row["cell_bottom"])
    text = _ocr_cell(row["image"], box)
    return text if re.fullmatch(r"\d{2,6}", text) else None


def _colonial_vote(group, field, line_no):
    """
    ITEM# / SIZE. No alcanza con "dos de tres" cualquiera: las lecturas sin
    grilla repiten a veces el mismo error (15CT -> "115CT" en dos) y un SIZE
    o un ITEM# mal leído es un costo por unidad o un producto equivocado.
    Vale el valor si:
    - la lectura principal del campo (la que nunca leyó mal en la medición:
      SIZE la hoja tal cual, ITEM# la hoja sin líneas) lo confirma otra
      lectura, o la celda sola si las demás no lo leyeron; o
    - la principal no lo leyó y lo dicen igual todas las demás que lo
      leyeron (al menos dos); o
    - lo dicen al menos tres lecturas y como mucho una dice otra cosa.
    Si no, ValueError.
    """
    primary_read = _COLONIAL_PRIMARY[field]
    by_read = {r["read"]: r[field] for r in group if r[field]}
    if field == "item":
        # El borde izquierdo pegado al ITEM# en las lecturas con grilla (0 y
        # 3): "17098" por 7098. No cuenta como lectura distinta.
        clean_values = {v for read, v in by_read.items() if read in (1, 2)}
        by_read = {read: v for read, v in by_read.items()
                   if not (read in (0, 3) and any(v[1:] == c for c in clean_values))}
    primary = by_read.get(primary_read)
    others = [v for read, v in by_read.items() if read != primary_read]
    values = list(by_read.values())
    value = None
    if primary is not None and primary in others:
        value = primary
    elif primary is not None and not others:
        row = next(r for r in group if r["read"] == primary_read)
        value = primary if _colonial_cell(row, field) == primary else None
    elif primary is None and len(others) >= 2 and len(set(others)) == 1:
        value = others[0]
    if value is None:
        best = max(set(values), key=values.count, default=None)
        if best is not None and values.count(best) >= 3 and len(values) - values.count(best) <= 1:
            value = best
    if value is None:
        label = "ITEM#" if field == "item" else "SIZE (unidades)"
        raise ValueError(f"el {label} del renglón {line_no} de Colonial no se pudo leer con seguridad.")
    return value


# ---------------------------------------------------------------------------
# Tickets impresos por lectura múltiple: Gold Coast Eagle y Red Bull (2026-10-05)
# ---------------------------------------------------------------------------
# Pedido del usuario (2026-10-05): sumar un par de proveedores más a la
# lectura de productos para comparar los precios con los de la factura
# anterior. Los dos son tickets impresos escaneados con poca resolución, a
# veces torcidos o curvados: el texto se lee bastante limpio, pero Tesseract
# confunde algunos dígitos de esas fuentes (en Red Bull el 5 sale 6 casi
# siempre; en GCE 8/6, 9/0 y 1/7), y cada forma de preparar la imagen se
# equivoca en renglones distintos. Entonces:
# - la página se endereza y se lee varias veces (_TICKET_PASSES, de a una y
#   solo hasta que la factura cierra); los renglones de las distintas
#   pasadas se emparejan por altura en la página, y cada monto va a la
#   columna cuyo título le queda encima;
# - de cada renglón se elige la combinación de lecturas que cierra sus
#   cuentas con más votos (_ticket_options: primero solo lo leído; si no
#   cierra nada, uno o dos montos cambiados por un dígito confundible de lo
#   leído; y por último un campo deducido de la cuenta con la cantidad
#   leída). Un renglón que no cierra se relee solo (su franja, enderezada);
# - la suma de los renglones tiene que dar el total impreso tal como lo leyó
#   alguna pasada (si no, se prueba cambiar UN renglón por otra combinación
#   que también cierra, y vale solo si exactamente un cambio da el total), y
#   los conteos del pie (cajas, unidades) confirman cantidades y packs.
# Misma regla de oro que el resto: si algo no cierra, ValueError y no se
# guarda ningún renglón.


def _sharpen(image):
    return image.filter(ImageFilter.UnsharpMask(radius=2, percent=150, threshold=3))


def _otsu(image):
    ensure_cv2()
    _, gray = cv2.threshold(np.array(image.convert("L")), 0, 255, cv2.THRESH_BINARY + cv2.THRESH_OTSU)
    return Image.fromarray(gray)


def _blur_threshold(image):
    ensure_cv2()
    gray = cv2.GaussianBlur(np.array(image.convert("L")), (5, 5), 0)
    return Image.fromarray(np.where(gray < 150, 0, 255).astype(np.uint8))


# (config de Tesseract, agrandado, arreglo de la imagen), de la que más
# renglones lee bien a la que menos (medido sobre facturas reales de GCE).
_TICKET_PASSES = (
    ("--psm 6", 2, None),
    ("--psm 6", 3, _sharpen),
    ("--psm 6", 3, None),
    ("--psm 4", 3, None),
    ("--psm 6", 3, _otsu),
    ("--psm 6", 4, _blur_threshold),
)
# Dígitos que Tesseract confunde en estas fuentes (visto en las facturas reales).
_TICKET_CONFUSIONS = {"5": "6", "6": "58", "8": "63", "3": "8", "0": "9", "9": "0", "1": "7", "7": "1"}
_TICKET_AMOUNT = re.compile(r"-?\d{1,3}(?:[.,]?\d{3})*[.,]\d{2}")


def _deskew(image):
    """
    Endereza un ticket escaneado torcido (hasta 3 grados): con la imagen
    achicada prueba ángulos y se queda con el que deja los renglones más
    horizontales (máxima varianza de la tinta por fila). Torcido, el total
    del renglón (a la derecha) queda a otra altura que el Item # y el OCR lo
    lee como otro renglón. Si ya está derecha, devuelve la misma imagen.
    """
    ensure_cv2()
    gray = np.array(image.convert("L"))
    factor = min(1.0, 500 / gray.shape[1])
    small = cv2.resize(gray, None, fx=factor, fy=factor, interpolation=cv2.INTER_AREA)
    _, ink = cv2.threshold(small, 0, 1, cv2.THRESH_BINARY_INV + cv2.THRESH_OTSU)
    center = (small.shape[1] / 2, small.shape[0] / 2)
    best_angle, best_score = 0.0, None
    for step in range(-30, 31):
        angle = step / 10
        matrix = cv2.getRotationMatrix2D(center, angle, 1.0)
        rotated = cv2.warpAffine(ink, matrix, (small.shape[1], small.shape[0]), flags=cv2.INTER_NEAREST)
        score = float(np.var(rotated.sum(axis=1)))
        if best_score is None or score > best_score:
            best_angle, best_score = angle, score
    if abs(best_angle) < 0.15:
        return image
    # El lienzo se agranda para que el giro no recorte las puntas (el total
    # del renglón está pegado al borde derecho).
    height, width = gray.shape
    matrix = cv2.getRotationMatrix2D((width / 2, height / 2), best_angle, 1.0)
    cos, sin = abs(matrix[0, 0]), abs(matrix[0, 1])
    new_width, new_height = int(height * sin + width * cos) + 2, int(height * cos + width * sin) + 2
    matrix[0, 2] += new_width / 2 - width / 2
    matrix[1, 2] += new_height / 2 - height / 2
    straight = cv2.warpAffine(gray, matrix, (new_width, new_height), flags=cv2.INTER_CUBIC,
                              borderMode=cv2.BORDER_CONSTANT, borderValue=255)
    return Image.fromarray(straight)


def _ticket_images(pdf_path):
    """Imágenes de las páginas, orientadas y enderezadas."""
    return [_deskew(image) for image in _page_images(pdf_path)]


def _ticket_amount(text):
    """'«1.80' / '39,20' / '$1,373.70' -> float; '33.117', '26.9%', '4).21' -> None (dudoso)."""
    text = _clean_token(text or "").strip("«»~+$§")
    if not _TICKET_AMOUNT.fullmatch(text):
        return None
    whole = re.sub(r"[.,]", "", text[:-3].lstrip("-")) or "0"
    value = float(f"{whole}.{text[-2:]}")
    return -value if text.startswith("-") else value


def _ticket_variants(value, confusions):
    """Montos a un dígito confundible de distancia: 3.26 -> 3.25, 3.28, ..."""
    text = f"{value:.2f}"
    found = set()
    for index, char in enumerate(text):
        for other in confusions.get(char, ""):
            found.add(round(float(text[:index] + other + text[index + 1:]), 2))
    found.discard(value)
    return found


def _ticket_upc(digits):
    """UPC-A válido dentro de los dígitos leídos (con el 0 de adelante perdido, también), o None."""
    upc = _upc_candidate(digits)
    if upc is None and len(digits) == 11 and _gtin_ok("0" + digits):
        upc = "0" + digits
    return upc


def _median_height(words):
    heights = sorted(w["bottom"] - w["top"] for w in words)
    return heights[len(heights) // 2]


def _ticket_row(page, words, y=None, height=None):
    return {
        "page": page,
        "y": y if y is not None else sum((w["top"] + w["bottom"]) / 2 for w in words) / len(words),
        "height": height if height is not None else _median_height(words),
        "words": words,
        "text": _row_text(words),
    }


def _ticket_readings(images, passes_done):
    """
    Una pasada más sobre todas las páginas: renglones con su página y altura
    ({"page", "y", "height", "words", "text"}), en orden. passes_done:
    cuántas pasadas ya se hicieron (la siguiente de _TICKET_PASSES).
    """
    config, scale, prep = _TICKET_PASSES[passes_done]
    rows = []
    for page_no, image in enumerate(images):
        rows.extend(_ticket_row(page_no, words) for words in _ocr_rows(image, config, scale=scale, prep=prep))
    return rows


def _ticket_columns(readings, labels):
    """
    Centro (x) de cada columna de montos según el renglón de títulos, por
    página (cada página es otro escaneo, con otro margen) y juntando todas
    las pasadas: {página: {columna: x}}. labels: [(columna, regex del
    título)]. Una página sin títulos legibles no figura (sus montos van por
    orden).
    """
    found = {}
    for reading in readings:
        for row in reading:
            hits = {}
            for word in row["words"]:
                for name, pattern in labels:
                    if name not in hits and re.fullmatch(pattern, word["text"], re.IGNORECASE):
                        hits[name] = (word["x0"] + word["x1"]) / 2
                        break
            if len(hits) >= len(labels) - 1:
                page = found.setdefault(row["page"], {name: [] for name, _ in labels})
                for name, x in hits.items():
                    page[name].append(x)
    return {page: {name: _median(xs) for name, xs in names.items()}
            for page, names in found.items() if all(names.values())}


def _ticket_assign(words, columns, names):
    """
    Montos de un renglón por columna. Si vinieron justo tantos montos como
    columnas, van en orden. Si falta o sobra alguno (el código de barras
    dibujado, un monto ilegible), cada uno va a la columna cuyo título le
    queda más cerca, descontando el corrimiento del renglón (la página
    curvada corre los renglones de abajo respecto de los títulos); lo que
    cae lejos de todas se descarta. Lo ilegible queda None.
    """
    tokens = [w for w in words if re.search(r"\d", w["text"])]
    if len(tokens) == len(names):
        return {name: _ticket_amount(w["text"]) for name, w in zip(names, tokens)}
    values = {name: None for name in names}
    if columns is None or not tokens:
        return values
    xs = sorted(columns[name] for name in names)
    gap = min(b - a for a, b in zip(xs, xs[1:]))
    centers = [(w["x0"] + w["x1"]) / 2 for w in tokens]
    shift = _median([x - min(xs, key=lambda c: abs(c - x)) for x in centers])
    distance = {}
    for word, x in zip(tokens, centers):
        name = min(names, key=lambda n: abs(columns[n] - (x - shift)))
        off = abs(columns[name] - (x - shift))
        if off <= gap * 0.5 and off < distance.get(name, float("inf")):
            values[name], distance[name] = _ticket_amount(word["text"]), off
    return values


def _ticket_clusters(candidates, tolerance):
    """
    Junta los renglones de producto de todas las pasadas: dos lecturas son el
    mismo renglón si están en la misma página y a menos de `tolerance` de
    altura. Devuelve los grupos en orden de página y altura.
    """
    clusters = []
    for cand in candidates:
        for cluster in clusters:
            if cluster[0]["page"] == cand["page"] and abs(cluster[0]["y"] - cand["y"]) <= tolerance:
                cluster.append(cand)
                break
        else:
            clusters.append([cand])
    clusters.sort(key=lambda c: (c[0]["page"], c[0]["y"]))
    return clusters


def _ticket_votes(cluster, field):
    counts = {}
    for cand in cluster:
        value = cand.get(field)
        if value is not None:
            counts[value] = counts.get(value, 0) + 1
    return counts


def _ticket_winner(counts):
    """El valor más votado; None si no hay o si dos empatan arriba."""
    if not counts:
        return None
    ranked = sorted(counts.items(), key=lambda kv: -kv[1])
    if len(ranked) > 1 and ranked[0][1] == ranked[1][1]:
        return None
    return ranked[0][0]


def _ticket_options(cluster, fields, close, derive, confusions, max_changed):
    """
    Combinaciones de lecturas que cierran el renglón, de la más votada a la
    menos: [(votos, valores)]. close(valores) -> los valores con "net" y
    "qty" si cierran, o None. Por niveles, y se pasa al siguiente solo si el
    anterior no dio ninguna:
    0) todos los montos tal como se leyeron;
    1) hasta max_changed montos cambiados por otro a un dígito confundible
       de lo leído ("3.26" por 3.25) -- los cambiados no suman votos. Este
       nivel se calcula siempre y queda detrás del 0 como alternativa para
       _ticket_fix_sum: en Red Bull "$64.70 / 3.25 / $61.45" se lee así en
       todas las pasadas y cierra, pero lo impreso es $54.70 y $51.45;
    2) un monto que ninguna pasada leyó bien sale de la cuenta del renglón
       (derive: [(campo, función(valores, cantidad), exige la cantidad
       leída)]); si el deducido es el total del renglón, ext_read=False y
       después lo tiene que confirmar el conteo de cajas del pie.
    La cantidad leída igual a la de la cuenta suma votos (qty_read). Cada
    opción es (votos, valores, nivel), ordenadas por nivel y votos.
    """
    votes = {field: sorted(_ticket_votes(cluster, field).items(), key=lambda kv: -kv[1])[:3] for field in fields}
    qty_votes = _ticket_votes(cluster, "qty")
    tiers = ({}, {}, {})

    def consider(tier, values, score, qty_needed=None, ext_read=True):
        result = close(dict(values))
        if result is None or (qty_needed is not None and result["qty"] != qty_needed):
            return
        n_qty = qty_votes.get(result["qty"], 0)
        key = tuple(values[f] for f in fields)
        total = score + n_qty
        if key not in tiers[tier] or tiers[tier][key][0] < total:
            tiers[tier][key] = (total, dict(result, qty_read=n_qty > 0, ext_read=ext_read), tier)

    def combos(lists):
        for choice in itertools.product(*lists):
            yield {f: v for f, (v, _) in zip(fields, choice)}, sum(n for _, n in choice)

    for values, score in combos([votes[f] for f in fields]):
        consider(0, values, score)
    for size in range(1, max_changed + 1):
        for changed in itertools.combinations(fields, size):
            lists = [[(v, 0) for read, _ in votes[f] for v in _ticket_variants(read, confusions)] if f in changed
                     else votes[f] for f in fields]
            for values, score in combos(lists):
                consider(1, values, score)
    if not tiers[0] and not tiers[1]:
        for field, deduce, needs_qty in derive:
            others = [f for f in fields if f != field]
            for choice in itertools.product(*[votes[f] for f in others]):
                values = {f: v for f, (v, _) in zip(others, choice)}
                score = sum(n for _, n in choice)
                for qty in [q for q in qty_votes if q > 0] if needs_qty else [None]:
                    values[field] = round(deduce(values, qty), 2)
                    consider(2, values, score, qty, ext_read=field != "ext")
    for key in tiers[0]:
        tiers[1].pop(key, None)
    if tiers[0] or tiers[1]:
        return sorted(tiers[0].values(), key=lambda o: -o[0]) + sorted(tiers[1].values(), key=lambda o: -o[0])
    return sorted(tiers[2].values(), key=lambda o: -o[0])


def _ticket_settled(options):
    """Hay una combinación ganadora: la primera no empata en nivel y votos con otra distinta."""
    return bool(options) and (len(options) == 1 or options[1][2] > options[0][2] or options[0][0] > options[1][0])


def _ticket_fix_sum(rows, printed_values, label, supplier):
    """
    rows: [{"options": [(votos, valores)], "choice": índice}]. Si la suma de
    los totales elegidos no da ninguno de los totales impresos leídos, prueba
    cambiar UN renglón por otra de sus combinaciones que cierran, y si no
    alcanza, dos o tres; vale solo si un único cambio da un total leído al
    centavo. Devuelve el total impreso que cerró.
    """
    def ext(row, index):
        return row["options"][index][1]["ext"]

    current = round(sum(ext(r, r["choice"]) for r in rows), 2)
    printed = {round(v, 2) for v in printed_values if v is not None}
    if current in printed:
        return current
    # Para la suma solo importa el total del renglón: de cada renglón, un cambio
    # por cada total distinto (el de la combinación mejor ubicada).
    changes = []
    for n, r in enumerate(rows):
        seen = {ext(r, r["choice"])}
        for alt in range(len(r["options"])):
            if ext(r, alt) not in seen:
                seen.add(ext(r, alt))
                changes.append((n, alt, round(ext(r, alt) - ext(r, r["choice"]), 2)))
    for size in (1, 2, 3):
        if size == 3 and len(changes) > 80:
            break
        fixes = [combo for combo in itertools.combinations(changes, size)
                 if len({n for n, _, _ in combo}) == size
                 and round(current + sum(d for _, _, d in combo), 2) in printed]
        if len(fixes) == 1:
            for n, alt, _ in fixes[0]:
                rows[n]["choice"] = alt
            return round(current + sum(d for _, _, d in fixes[0]), 2)
        if fixes:
            break
    shown = ", ".join(f"${v:,.2f}" for v in sorted(printed)) or "ilegible"
    raise ValueError(f"los renglones de {supplier} suman ${current:,.2f} y el {label} leído es {shown}.")


def _ticket_reread_row(images, cluster, parse, cache):
    """
    Relee solo la franja del renglón, enderezada por su cuenta (el ticket
    viene curvado: arriba se tuerce más que abajo y un solo ángulo para toda
    la página no alcanza), más agrandada y con varios arreglos de la imagen,
    y suma esas lecturas al grupo. cache: las relecturas ya hechas en
    pasadas anteriores, por página y altura.
    """
    page = cluster[0]["page"]
    y = _median([c["y"] for c in cluster])
    height = _median([c["height"] for c in cluster])
    key = next((k for k in cache if k[0] == "row" and k[1] == page and abs(k[2] - y) <= height / 2),
               ("row", page, y))
    if key not in cache:
        image = images[page]
        top = max(0, int(y - 1.8 * height))
        strip = _deskew(image.crop((0, top, image.width, min(image.height, int(y + 1.8 * height)))))
        found = []
        for config, scale, prep in (("--psm 6", 3, None), ("--psm 7", 3, None), ("--psm 6", 4, _sharpen),
                                    ("--psm 6", 3, _otsu)):
            for words in _ocr_rows(strip, config, scale=scale, prep=prep):
                candidate = parse(_ticket_row(page, words, y=y, height=height))
                if candidate is not None:
                    candidate["upc_box"] = None  # posición de la franja, no de la página
                    found.append(candidate)
        cache[key] = found
    return cluster + cache[key]


def _ticket_reread_upc(images, cluster, cache):
    """
    UPC de un renglón que ninguna pasada leyó con dígito verificador válido:
    se relee la celda sola, solo dígitos, donde la vieron las pasadas (hasta
    dos lugares distintos). Vale si todas las relecturas válidas dan el mismo.
    """
    found = set()
    boxes = []
    for cand in cluster:
        box = cand.get("upc_box")
        if box is not None and not any(abs(box[0] - b[0]) < 4 and abs(box[1] - b[1]) < 4 for b in boxes):
            boxes.append(box)
    page = cluster[0]["page"]
    for x0, top, x1, bottom in boxes[:2]:
        key = ("upc", page, round(x0), round(top))
        if key not in cache:
            pad = (bottom - top) * 0.6
            box = (x0 - pad, top - pad, x1 + pad, bottom + pad)
            cache[key] = {_ticket_upc(re.sub(r"\D", "", _ocr_cell(images[page], box, psm=psm)))
                          for psm in (7, 8)} - {None}
        found |= cache[key]
    return found.pop() if len(found) == 1 else None


def _ticket_upc_variants(cluster, confusions, prefix=""):
    """
    UPC de un renglón que ni las pasadas ni la relectura leyeron válido: si
    dos o más pasadas leyeron los mismos 12 dígitos (inválidos), se prueban
    los cambios de UN dígito confundible ("611269113670" -> 611269113570, el
    5 que en Red Bull sale 6), sin tocar el prefijo de la empresa si se
    conoce; vale si exactamente uno pasa el dígito verificador.
    """
    found = set()
    for raw, count in _ticket_votes(cluster, "upc_raw").items():
        if count < 2 or not raw.startswith(prefix):
            continue
        for index, char in enumerate(raw):
            if index < len(prefix):
                continue
            for other in confusions.get(char, ""):
                candidate = raw[:index] + other + raw[index + 1:]
                if _gtin_ok(candidate):
                    found.add(candidate)
    return found.pop() if len(found) == 1 else None


def _ticket_raw_upc(digits):
    """Los 12 dígitos leídos del UPC (con el 0 de adelante si se perdió), para _ticket_upc_variants."""
    if len(digits) == 11:
        digits = "0" + digits
    return digits if len(digits) == 12 else None


def _ticket_rows(clusters, images, cache, parse, supplier, rules):
    """
    Cada grupo de lecturas -> renglón con sus combinaciones que cierran, UPC y
    descripción. rules: las cuentas y confusiones del proveedor (_GCE_RULES,
    _RB_RULES).
    """
    rows = []
    for line_no, cluster in enumerate(clusters, start=1):
        options = _ticket_options(cluster, *rules["options"])
        upc = _ticket_winner(_ticket_votes(cluster, "upc"))
        if not _ticket_settled(options) or upc is None:
            cluster = _ticket_reread_row(images, cluster, parse, cache)
            options = _ticket_options(cluster, *rules["options"])
            upc = (_ticket_winner(_ticket_votes(cluster, "upc")) or _ticket_reread_upc(images, cluster, cache)
                   or _ticket_upc_variants(cluster, *rules["upc"]))
        description = (_ticket_winner(_ticket_votes(cluster, "description"))
                       or next((c["description"] for c in cluster if c.get("description")), ""))
        if not options:
            raise ValueError(f"el renglón {line_no} de {supplier} ({description}) no cierra: "
                             "ninguna lectura da cantidad x precio = total.")
        if not _ticket_settled(options):
            raise ValueError(f"el renglón {line_no} de {supplier} ({description}) tiene dos lecturas "
                             "posibles que cierran con los mismos votos.")
        if upc is None:
            raise ValueError(f"no se pudo leer el UPC del renglón {line_no} de {supplier} ({description}).")
        rows.append({"options": options, "choice": 0, "cluster": cluster, "upc": upc,
                     "description": re.sub(r"\s+", " ", description).strip(" |")})
    return rows


def _ticket_chosen(rows):
    return [dict(r["options"][r["choice"]][1], upc=r["upc"], description=r["description"], cluster=r["cluster"])
            for r in rows]


def _ticket_check_qty(chosen, confirmed, supplier, label):
    """
    Si el pie confirma las cantidades (confirmed: la suma de cantidades da el
    conteo de cajas impreso, o la de cantidad x unidades da el de unidades),
    listo; si no, cada cantidad deducida del total (o cada total deducido de
    la cantidad) tiene que haberla leído alguna pasada.
    """
    if confirmed:
        return
    unread = next((n for n, c in enumerate(chosen, 1) if c["qty"] and (not c["qty_read"] or not c["ext_read"])),
                  None)
    if unread is not None:
        raise ValueError(f"la cantidad del renglón {unread} de {supplier} no se pudo leer con seguridad "
                         f"(y el {label} del pie no la confirma).")


def _ticket_top(values):
    """
    Las lecturas más votadas de un número del pie (empatadas, todas): una
    lectura suelta que otras pasadas contradicen no cuenta -- con ella, la
    suma de renglones mal leídos podía "cerrar" ($593.90 leído una vez
    contra $553.90 en las otras cinco).
    """
    counts = _ticket_votes([{"v": v} for v in values], "v")
    top = max(counts.values(), default=0)
    return [v for v, n in counts.items() if n == top]


def _ticket_most_voted(values):
    """El valor que más se repite en una lista de lecturas del pie (None si no hay o empatan)."""
    return _ticket_winner(_ticket_votes([{"v": v} for v in values], "v"))


# --- Gold Coast Eagle --------------------------------------------------------
# Ticket angosto de 1 o 2 páginas (a veces dos facturas del mismo reparto en
# un PDF). Cada producto ocupa dos renglones:
#   ITEM# QTY U.P.C. PRICE DISC D.PRICE DEP EXT
#         Descripción con el pack al final ("Corona Extra 4/6/12 Ln")
# Cuentas: PRICE - DISC = D.PRICE; QTY x (D.PRICE + DEP) = EXT; el UPC pasa el
# dígito verificador; la suma de EXT = Total Sales; Total Sales - Total
# Credits = Invoice Total (o el importe de la línea "Inv# ... $..."); la suma
# de QTY = "Cases" y la de QTY x unidades = "Selling Units". La cantidad casi
# nunca se lee limpia (queda pegada al Item # o sale "]", "=", "i"): se deduce
# de EXT / D.PRICE y la confirma "Cases". Los renglones con EXT 0.00 son
# productos sin stock ("-1 Out of Stock") y no se guardan.

_GCE_FIELDS = ("price", "disc", "dprice", "dep", "ext")
_GCE_HEADERS = (("price", r"PRICE"), ("disc", r"DISC\W?"), ("dprice", r"D\W?PRICE"), ("dep", r"DEP"),
                ("ext", r"E[XA][TI1]"))
_GCE_PACK = re.compile(r"(\d{1,2})\s*/\s*(\d{1,3}(?:\s?\.\s?\d{1,2})?)\s*(ml|l\b|liter|pk|oz)?"
                       r"(?:\s*/\s*(\d{1,3}(?:\.\d)?))?", re.IGNORECASE)


def _gce_units(description):
    """
    (unidades de venta del POS por caja, pack impreso) según el pack de la
    descripción -- regla verificada contra "Selling Units" del pie y la lista
    de productos del POS:
    - tres niveles, "4/6/12": 4 paquetes (six-packs) por caja;
    - dos niveles, "15/25", "12/32", "24/200ml": unidades sueltas;
    - salvo envase de 16 oz o menos en caja de 15 o más ("24/12", "18/12",
      "24/7", "15/16"): esa caja es el pack que vende el POS ("BUSCH 24PK
      CANS", "18 PK CORONA", "BUDWEISER 15PK/16OZ" = 1 unidad).
    """
    matches = list(_GCE_PACK.finditer(description or ""))
    if not matches:
        return None, ""
    match = matches[-1]
    first = int(match.group(1))
    size = description[match.start():].strip(" |.,")
    if match.group(4):  # tres niveles
        return first, size
    try:
        second = float(match.group(2).replace(" ", ""))
    except ValueError:
        return None, size
    unit = (match.group(3) or "").lower()
    if unit not in ("ml", "l", "liter") and second <= 16 and first >= 15:
        return 1, size
    return first, size


def _gce_product(row, columns):
    """Renglón de producto leído en una pasada (None lo ilegible), o None si no es un renglón de producto."""
    words = row["words"]
    texts = [w["text"] for w in words]
    at = next((i for i, t in enumerate(texts[:5]) if len(re.sub(r"\D", "", _fix_digits(t))) >= 11), None)
    if at is None or len([t for t in texts[at + 1:] if re.search(r"\d", t)]) < 3:
        return None
    item = qty_text = ""
    before = [d for d in (re.sub(r"\D", "", _fix_digits(t)) for t in texts[:at]) if d]
    if before and len(before[0]) >= 5:
        item, qty_text = before[0][:5], before[0][5:] + "".join(before[1:])
    elif len(before) > 1:
        qty_text = "".join(before[1:])
    word = words[at]
    return dict(
        _ticket_assign(words[at + 1:], columns.get(row["page"]), _GCE_FIELDS),
        page=row["page"], y=row["y"], height=row["height"], item=item or None, description=None,
        upc=_ticket_upc(re.sub(r"\D", "", _fix_digits(texts[at]))),
        upc_raw=_ticket_raw_upc(re.sub(r"\D", "", _fix_digits(texts[at]))),
        upc_box=(word["x0"], word["top"], word["x1"], word["bottom"]),
        qty=int(qty_text) if qty_text and len(qty_text) <= 3 else None,
    )


def _gce_close(values):
    price, disc, dprice, dep, ext = (values[f] for f in _GCE_FIELDS)
    if abs(price - disc - dprice) > _TOLERANCE or dprice <= 0 or disc < 0 or not 0 <= dep < dprice or ext < 0:
        return None
    net = round(dprice + dep, 2)
    qty = 0 if ext == 0 else _whole_qty(net, ext)
    return None if qty is None else dict(values, net=net, qty=qty)


_GCE_DERIVE = (
    ("dprice", lambda v, qty: v["price"] - v["disc"], True),
    ("disc", lambda v, qty: v["price"] - v["dprice"], False),
    ("ext", lambda v, qty: qty * (v["dprice"] + v["dep"]), True),
)
# Confusiones en los dos sentidos (8/6, 9/0, 1/7, 3/8): hasta dos montos cambiados por renglón.
_GCE_RULES = {
    "options": (_GCE_FIELDS, _gce_close, _GCE_DERIVE, _TICKET_CONFUSIONS, 2),
    "upc": (_TICKET_CONFUSIONS,),
}


def _gce_footer(text):
    """Valores del pie que trae un renglón: {"cases"|"selling"|"sales"|"credits"|"total": valor}."""
    # "1, 385.99" (la coma separada) y "] 385.99" / "|, 385.99" (el 1 de adelante mal leído).
    text = re.sub(r"(?<![\d,.])[\]|](?:\s*,\s*|\s+)(\d{3}[.,]\d{2})(?!\d)", r"1,\1", text)
    text = re.sub(r"(\d)\s*,\s+(\d{3}[.,]\d{2})(?!\d)", r"\1,\2", text)
    found = {}
    patterns = (
        ("sales", r"Total\s*Sa\w*"), ("credits", r"Total\s*Cr\w*"), ("total", r"In\w{3,5}\s*T[o0]\w{2,3}"),
    )
    for key, label in patterns:
        match = re.search(label + r"\W*(-?[\d,]+[.,]\d{2})(?!\d)", text)
        if match:
            found[key] = _ticket_amount(match.group(1))
    match = re.match(r"^\W*Cases\s*\W?\s*(\d+)\s*$", text)
    if match:
        found["cases"] = int(match.group(1))
    match = re.search(r"Selling\s*Un\w*\W*(\d+)\s*$", text)
    if match:
        found["selling"] = int(match.group(1))
    return found


def _gce_header_no(text):
    """
    N° de la factura en el renglón del encabezado ("Account: 33994 Invoice#:
    777888 PO#:"; el OCR a veces lo parte en dos renglones), nunca en la
    línea de confirmación del pie ("Inv# 777888 $1,373.70").
    """
    if "$" in text:
        return None
    match = re.search(r"(\d{6})\s*P[O0]\w?\W", text)
    if match is None:
        match = re.search(r"(?:Inv|nvo)\w{0,5}\W{1,4}(\d{6})(?!\d)", text)
    return match.group(1) if match else None


_GCE_CONFIRM = re.compile(r"Inv\w{0,2}\W{0,4}(\d{6})\W{0,3}[$S§]\s*([\d,]+[.,]\d{2})(?!\d)")
_GCE_DATE = re.compile(r"\b(Mon|Tue|Wed|Thu|Fri|Sat|Sun)\w*\W+([A-Z][a-z]{2})\w*\W+(\d{1,2}),?\s*(\d{4})\b")


def _gce_date(text):
    """
    Fecha del encabezado ("Thu Jan 22, 2026 7:18 AM"), o None. Exige el día de
    la semana y que coincida con la fecha: así un dígito mal leído no pasa, y
    tampoco el "Expires Mar 31, 2027" del recuadro de la licencia.
    """
    match = _GCE_DATE.search(text)
    if match is None:
        return None
    try:
        found = datetime.strptime(f"{match.group(2)} {int(match.group(3))} {match.group(4)}", "%b %d %Y")
    except ValueError:
        return None
    return found if found.strftime("%a") == match.group(1) else None


def _gce_blocks(readings):
    """
    Arma las facturas del PDF con todas las pasadas hechas hasta ahora. Cada
    factura termina en su renglón "Invoice Total" (o "Total Sales"): los
    renglones de producto, el pie, la fecha y el N° del encabezado van a la
    factura cuyo cierre es el primero que viene después. Las líneas de
    confirmación del reparto ("Inv# 677092 $1,034.07", una por cada factura
    del reparto, repetidas al pie de cada ticket) van aparte. Devuelve
    (facturas, confirmaciones {N°: {importe}}, columnas).
    """
    all_rows = [row for reading in readings for row in reading]
    tolerance = _median([row["height"] for row in all_rows]) if all_rows else 10
    columns = _ticket_columns(readings, _GCE_HEADERS)
    ends = [{"page": row["page"], "y": row["y"]} for row in all_rows
            if {"total", "sales"} & set(_gce_footer(row["text"]))]
    boundaries = [(c[0]["page"], max(e["y"] for e in c) + tolerance) for c in _ticket_clusters(ends, tolerance * 5)]
    if not boundaries:
        raise ValueError("no se encontró el pie (Invoice Total) de la factura de Gold Coast.")

    def block_of(row):
        return next((i for i, (page, y) in enumerate(boundaries) if (row["page"], row["y"]) <= (page, y)), None)

    blocks = [{"end": end, "products": [], "footer": {}, "numbers": {}, "dates": {}} for end in boundaries]
    confirm_rows = []
    for reading in readings:
        started = set()  # facturas que en esta pasada ya pasaron el encabezado
        for position, row in enumerate(reading):
            match = _GCE_CONFIRM.search(row["text"])
            if match:
                confirm_rows.append({"page": row["page"], "y": row["y"], "number": match.group(1),
                                     "amount": _ticket_amount(match.group(2))})
                continue
            index = block_of(row)
            if index is None:
                continue
            block = blocks[index]
            product = _gce_product(row, columns)
            if product is not None:
                # Descripción: el renglón siguiente de esta misma pasada, si no es
                # otro producto, una nota ("-1 Out of Stock") ni el pie.
                following = reading[position + 1] if position + 1 < len(reading) else None
                if following is not None and _gce_product(following, columns) is None:
                    text = following["text"].strip(" |")
                    if (re.search(r"[A-Za-z]{3}", text) and not re.match(r"^\W*\d+\s+(Out|Picker)", text)
                            and not _gce_footer(text)):
                        product["description"] = text
                block["products"].append(product)
                started.add(index)
                continue
            for key, value in _gce_footer(row["text"]).items():
                block["footer"].setdefault(key, []).append(value)
            if index in started:
                continue
            number = _gce_header_no(row["text"])
            if number:
                block["numbers"][number] = block["numbers"].get(number, 0) + 1
            found = _gce_date(row["text"])
            if found:
                block["dates"][found] = block["dates"].get(found, 0) + 1
    for block in blocks:
        block["clusters"] = _ticket_clusters(block["products"], tolerance)
    # Cada línea de confirmación (la misma línea en todas las pasadas) vota su N°
    # y su importe; las dos líneas del reparto quedan muy juntas (16 px). Su
    # letra se lee peor que la del encabezado (0/8, 8/6: "633549 $6,020.99" por
    # 833549 $8,020.99): su N° es seguro ("sure") solo si todas las pasadas (3 o
    # más) lo leyeron igual.
    pairs = {}
    for cluster in _ticket_clusters(confirm_rows, max(4, tolerance * 0.35)):
        votes = _ticket_votes(cluster, "number")
        number = _ticket_winner(votes)
        amount = _ticket_winner(_ticket_votes(cluster, "amount"))
        if number is not None and amount is not None:
            entry = pairs.setdefault(number, {"amounts": set(), "sure": False})
            entry["amounts"].add(amount)
            entry["sure"] = entry["sure"] or (votes[number] >= 3 and len(votes) == 1)
    return blocks, pairs, columns


def _gce_resolve(block, images, cache, columns, confirmed):
    """
    Detalle de una factura de GCE ya armada, o ValueError si algo no cierra.
    confirmed: los importes de las líneas de confirmación del PDF (el Invoice
    Total de la factura puede estar ilegible).
    """
    rows = _ticket_rows(block["clusters"], images, cache, lambda row: _gce_product(row, columns), "Gold Coast",
                        _GCE_RULES)
    if not rows:
        raise ValueError("no se ven los renglones de producto de la factura.")
    footer = block["footer"]
    credit = _ticket_most_voted(footer.get("credits", [])) or 0.0
    totals = [t for t in _ticket_top(footer.get("total", [])) if t is not None]
    if not totals:
        # Invoice Total ilegible: el Total Sales menos los créditos o, sin pie
        # legible, los importes de las líneas de confirmación.
        totals = ([round(s - credit, 2) for s in _ticket_top(footer.get("sales", [])) if s is not None]
                  or sorted(confirmed))
    # La suma de los renglones tiene que dar el Total Sales o el Invoice Total más los créditos.
    printed = set(_ticket_top(footer.get("sales", []))) | {round(t + credit, 2) for t in totals}
    sales = _ticket_fix_sum(rows, printed, "Total Sales", "Gold Coast")
    total = next((t for t in totals if abs(sales - credit - t) < 0.005), None)
    if total is None:
        raise ValueError(f"el Total Sales (${sales:,.2f}) menos los créditos no da el Invoice Total leído.")
    chosen = [c for c in _ticket_chosen(rows) if c["qty"]]  # EXT 0.00: sin stock, no se entregó
    if not chosen:
        raise ValueError("la factura no tiene ningún producto entregado.")
    for c in chosen:
        # Se vota solo la cantidad por caja: el texto del pack varía con el OCR
        # ("6/4/16 Can" / "6/4/16 Can.") y partía los votos en un empate.
        packs = [_gce_units(cand.get("description")) for cand in c["cluster"]]
        per_case = _ticket_most_voted([p[0] for p in packs if p[0]])
        if per_case is None:
            raise ValueError(f"no se pudo leer el pack (unidades por caja) de {c['description'] or c['upc']}.")
        sizes = [p[1] for p in packs if p[0] == per_case]
        c["pack"] = (per_case, _ticket_most_voted(sizes) or sizes[0])
    selling = sum(c["qty"] * c["pack"][0] for c in chosen)
    selling_ok = selling in footer.get("selling", [])
    _ticket_check_qty(chosen, selling_ok or sum(c["qty"] for c in chosen) in footer.get("cases", []),
                      "Gold Coast", "total de cajas (Cases) o de unidades (Selling Units)")
    if not selling_ok:
        raise ValueError(f"las unidades de venta ({selling}) no coinciden con el Selling Units del pie.")
    lines = []
    for c in chosen:
        per_case, size = c["pack"]
        lines.append(_line(
            len(lines) + 1, upc=c["upc"], item_no=_ticket_winner(_ticket_votes(c["cluster"], "item")) or "",
            description=c["description"], qty=c["qty"], pack=per_case, size=size, units=per_case,
            price=c["price"], allowance=c["disc"] or None, net=c["net"], ext=c["ext"],
        ))
    return {"lines": lines, "subtotal": sales, "total": total}


def _gce_footer_total(block, pairs, number):
    """
    Invoice Total de una factura cuyos renglones no cerraron: vale si lo dicen
    dos lugares distintos del ticket -- el Invoice Total y la línea de
    confirmación de ese N°, o el Invoice Total y el Total Sales menos los
    créditos. None si no.
    """
    footer = block["footer"]
    totals = set(_ticket_top(footer.get("total", []))) - {None}
    agree = totals & pairs.get(number, {}).get("amounts", set())
    if len(agree) == 1:
        return agree.pop()
    total = _ticket_most_voted(footer.get("total", []))
    sales = _ticket_most_voted(footer.get("sales", []))
    credit = _ticket_most_voted(footer.get("credits", [])) or 0.0
    if total is not None and sales is not None and abs(sales - credit - total) < 0.005:
        return total
    return None


def _filename_numbers(filename, digits):
    return set(re.findall(rf"(?<!\d)(\d{{{digits}}})(?!\d)", filename or ""))


def _filename_dates(filename):
    """
    Fechas completas del nombre del archivo, en los dos órdenes (día-mes y
    mes-día, igual que el control de fecha de la carga en webapp.py):
    "Invoice 668752 22.01.2026.pdf" -> {22/01/2026}.
    """
    found = set()
    for match in re.finditer(r"(?<!\d)(\d{1,2})[.\-_](\d{1,2})[.\-_](\d{4})(?!\d)", filename or ""):
        a, b, year = (int(g) for g in match.groups())
        for day, month in ((a, b), (b, a)):
            try:
                found.add(datetime(year, month, day))
            except ValueError:
                pass
    return found


# Algunos repartos especiales (una sola marca, 1 renglón) vienen en otra
# plantilla: hoja completa impresa, no ticket ("LOAD SLSMN ACCT # DATE INV"
# arriba; el pie "8 CASE 232.48 BEER$ / 232.48 CONTENT$ / .00 DEPOSIT$" y el
# TOTAL a la derecha). 3 de 73 PDFs de 2025-2026. Renglón: descripción,
# CODE, CASE, PRICE (el neto), UPC sin el 0 de adelante ni el dígito
# verificador, DISC, AMOUNT.
_GCE_FORM_HEADER = re.compile(r"(?<!\d)\d{3,4}\s+\d{3}\s+\d{5}\s+(\d{1,2})/(\d{1,2})/(\d{2})\s+(\d{6})(?!\d)")
_GCE_FORM_LINE = re.compile(r"^(?P<description>.*?[A-Za-z].*?)\s+(?P<code>\d{5})\s+(?P<qty>\d{1,3})\s+"
                            r"(?P<price>\d[\d,]*\.\d{2})\s+(?P<upc>\d{10})\s+(?:\d*\.\d{2}\s+)?(?P<ext>\d[\d,]*\.\d{2})$")
_GCE_FORM_FOOTER = re.compile(r"(\d[\d,]*\.\d{2}|\.\d{2})\s*DEPOSIT\S*\s+(\d[\d,]*\.\d{2})\s*$")


def _gce_is_form(reading):
    return any("SLSMN" in row["text"] or re.search(r"DESCRIPTION\s+CODE\s+CASE", row["text"]) for row in reading)


def _form_amount(text):
    return _ticket_amount("0" + text if text.startswith(".") else text)


def _gce_form_read(readings):
    """Lo leído de la plantilla de hoja completa en todas las pasadas."""
    found = {"numbers": {}, "dates": {}, "contents": [], "deposits": [], "totals": [], "products": []}
    for reading in readings:
        for row in reading:
            text = row["text"].strip(" |")
            match = _GCE_FORM_HEADER.search(text)
            if match:
                found["numbers"][match.group(4)] = found["numbers"].get(match.group(4), 0) + 1
                try:
                    when = datetime(2000 + int(match.group(3)), int(match.group(1)), int(match.group(2)))
                    found["dates"][when] = found["dates"].get(when, 0) + 1
                except ValueError:
                    pass
            match = re.search(r"(\d[\d,]*\.\d{2}|\.\d{2})\s*CONTENT", text)
            if match:
                found["contents"].append(_form_amount(match.group(1)))
            match = _GCE_FORM_FOOTER.search(text)
            if match:
                found["deposits"].append(_form_amount(match.group(1)))
                found["totals"].append(_form_amount(match.group(2)))
            match = _GCE_FORM_LINE.match(text)
            if match:
                upc = "0" + match.group("upc")
                found["products"].append({
                    "page": row["page"], "y": row["y"], "description": match.group("description"),
                    "item": match.group("code"), "qty": int(match.group("qty")),
                    "price": _form_amount(match.group("price")), "ext": _form_amount(match.group("ext")),
                    "upc": upc + str((10 - sum(int(d) * (3 if i % 2 == 0 else 1) for i, d in enumerate(upc)) % 10) % 10),
                })
    return found


def _gce_form_lines(found, tolerance, content):
    """Renglones de la plantilla de hoja completa, o (None, motivo) si no cierran contra el CONTENT$."""
    lines = []
    for cluster in _ticket_clusters(found["products"], tolerance):
        closing = [c for c in cluster if c["price"] and c["ext"] is not None
                   and abs(c["qty"] * c["price"] - c["ext"]) < _TOLERANCE]
        best = _ticket_winner(_ticket_votes([{"v": (c["qty"], c["price"], c["ext"], c["upc"])} for c in closing], "v"))
        if best is None:
            return None, "un renglón no cierra: ninguna lectura da cantidad x precio = total."
        qty, price, ext, upc = best
        description = _ticket_winner(_ticket_votes(cluster, "description")) or cluster[0]["description"]
        per_case, size = _gce_units(description)
        if not per_case:  # "0/12" leído del OCR: dividía por cero (revisión 2026-10-08)
            return None, f"no se pudo leer el pack (unidades por caja) de {description}."
        lines.append(_line(len(lines) + 1, upc=upc, item_no=_ticket_winner(_ticket_votes(cluster, "item")) or "",
                           description=description, qty=qty, pack=per_case, size=size, units=per_case,
                           price=price, allowance=None, net=price, ext=ext))
    if not lines:
        return None, "no se ven los renglones de producto de la factura."
    if abs(round(sum(line["ext"] for line in lines), 2) - content) >= 0.005:
        return None, "los renglones no suman el CONTENT$ impreso."
    return lines, None


def _gce_form_invoices(images, reading, filename, readings):
    """
    La factura de la plantilla de hoja completa: N° y fecha del renglón
    LOAD/SLSMN/ACCT/DATE/INV; importe = el TOTAL, confirmado por el CONTENT$
    más el DEPOSIT$ (y por la suma de renglones, si cierran).
    """
    from_name = _filename_numbers(filename, 6)
    result = None
    for passes_done in range(len(readings), len(_TICKET_PASSES) + 1):
        found = _gce_form_read(readings)
        tolerance = _median([row["height"] for reading_ in readings for row in reading_]) or 10
        number = _ticket_winner(found["numbers"])
        if number is not None and not (number in from_name or (found["numbers"][number] >= 3
                                                              and len(found["numbers"]) == 1)):
            number = None
        total = _ticket_most_voted(found["totals"])
        content = _ticket_most_voted(found["contents"])
        deposit = _ticket_most_voted(found["deposits"]) or 0.0
        when = _ticket_winner(found["dates"])
        if number is None:
            result = {"error": "no se pudo leer el N° de invoice con seguridad.", "invoice_no": None}
        elif total is None or content is None or abs(round(content + deposit, 2) - total) >= 0.005:
            result = {"error": "no se pudo leer el TOTAL con seguridad.", "invoice_no": int(number)}
        elif when is None:
            result = {"error": "no se pudo leer la fecha de la factura.", "invoice_no": int(number)}
        else:
            lines, problem = _gce_form_lines(found, tolerance, content)
            result = {"invoice_no": int(number), "date": when, "amount": total, "lines": lines, "lines_error": problem}
            if lines is not None:
                break
        if passes_done == len(_TICKET_PASSES):
            break
        readings.append(reading(passes_done))
    return [result]


def _gce_invoices(images, reading, filename, cache=None):
    """
    Todas las facturas de GCE de un PDF (ver read_gce_invoices). reading(n):
    los renglones de la pasada n de _TICKET_PASSES; cache: las relecturas de
    renglones ya hechas. Se lee de a una pasada y se sigue mientras alguna
    factura no cierre o no tenga su encabezado seguro.
    """
    readings, resolved, failures = [], [], []
    cache = {} if cache is None else cache
    results = None
    for passes_done in range(len(_TICKET_PASSES)):
        readings.append(reading(passes_done))
        if passes_done == 0 and _gce_is_form(readings[0]):
            return _gce_form_invoices(images, reading, filename, readings)
        try:
            blocks, pairs, columns = _gce_blocks(readings)
        except ValueError:
            continue
        confirmed = {amount for entry in pairs.values() for amount in entry["amounts"]}
        for block in blocks:
            if _ticket_resolved(resolved, block["end"]) is None:
                try:
                    resolved.append((block["end"], _gce_resolve(block, images, cache, columns, confirmed)))
                except ValueError as exc:
                    failures.append((block["end"], str(exc)))
        details = [_ticket_resolved(resolved, block["end"]) for block in blocks]
        results = _gce_assign(blocks, pairs, details, failures, filename)
        if all(details) and not any("error" in invoice for invoice in results):
            break
    if results is None:
        raise ValueError("no se encontró el pie (Invoice Total) de ninguna factura de Gold Coast.")
    return results


def _gce_assign(blocks, pairs, details, failures, filename):
    """
    Encabezado de cada factura armada: N°, importe y fecha confirmados (ver
    read_gce_invoices), o {"error", "invoice_no"} si alguno no es seguro.
    """
    from_name = _filename_numbers(filename, 6)
    results, taken = [], set()
    # Primero las facturas con detalle (su total ya está confirmado por la suma).
    for index in sorted(range(len(blocks)), key=lambda i: details[i] is None):
        block, detail = blocks[index], details[index]
        footer = block["footer"]
        if detail is not None:
            amounts = {detail["total"]}
        else:
            credit = _ticket_most_voted(footer.get("credits", [])) or 0.0
            amounts = (set(_ticket_top(footer.get("total", [])))
                       | {round(s - credit, 2) for s in _ticket_top(footer.get("sales", [])) if s is not None}) - {None}
        agree = {n for n in pairs if n not in taken and pairs[n]["amounts"] & amounts}
        votes = {n: v for n, v in block["numbers"].items() if n not in taken}
        header = _ticket_winner(votes)
        if header is not None and (header in from_name or header in agree
                                   or (votes[header] >= 3 and len(votes) == 1)):
            # El del encabezado: lo confirma el nombre del archivo, la línea de
            # confirmación con el mismo importe, o todas las pasadas (3 o más).
            number = header
        elif votes:
            # Encabezado empatado ("616458" / "816458"): el leído que confirma el
            # nombre del archivo o la línea de confirmación con el mismo importe.
            named = {n for n in votes if n in from_name or n in agree}
            number = named.pop() if len(named) == 1 else None
        else:
            # Sin encabezado (el ticket cortado arriba): la línea de confirmación
            # con el importe de la factura, si su N° es seguro.
            sure = {n for n in agree if n in from_name or pairs[n]["sure"]}
            number = sure.pop() if len(sure) == 1 else None
        if number is None:
            if not block["products"] and not votes and (
                    not amounts or any(pairs[n]["amounts"] & amounts for n in taken if n in pairs)):
                continue  # un pie suelto, o repetido de una factura ya leída
            results.append((index, {"error": "no se pudo leer el N° de invoice con seguridad.", "invoice_no": None}))
            continue
        amount = detail["total"] if detail is not None else _gce_footer_total(block, pairs, number)
        if amount is None:
            results.append((index, {"error": "no se pudo leer el Invoice Total con seguridad.",
                                    "invoice_no": int(number)}))
            continue
        taken.add(number)
        results.append((index, {"invoice_no": int(number), "amount": amount,
                                "lines": detail["lines"] if detail is not None else None,
                                "lines_error": None if detail is not None else _ticket_resolved(failures, block["end"])}))
    # Fecha: la del encabezado de la factura; si no se ve (el ticket cortado
    # arriba), la del reparto (la de las otras facturas del PDF, si es una
    # sola) o la del nombre del archivo.
    dates = {_ticket_winner(block["dates"]) for block in blocks} - {None}
    named = _filename_dates(filename)
    for index, invoice in results:
        if "error" in invoice:
            continue
        found = _ticket_winner(blocks[index]["dates"])
        if found is None:
            found = next(iter(dates)) if len(dates) == 1 else (next(iter(named)) if len(named) == 1 else None)
        if found is None:
            number = invoice["invoice_no"]
            invoice.clear()
            invoice.update(error="no se pudo leer la fecha de la factura.", invoice_no=number)
            continue
        invoice["date"] = found
    return [invoice for _, invoice in sorted(results, key=lambda r: r[0])]


def _ticket_resolved(found, end):
    """
    Lo último guardado (el detalle, o el error) de la factura que termina en
    `end` (página, altura): de una pasada a otra, el cierre se mueve unos píxeles.
    """
    page, y = end
    return next((value for (p, other), value in reversed(found) if p == page and abs(other - y) < 60), None)


def read_gce_invoices(pdf_path):
    """
    Todas las facturas de Gold Coast Eagle de un PDF (puede traer dos del
    mismo reparto), cada una con su encabezado confirmado y, si cierra, su
    detalle de productos:
    [{"invoice_no", "date", "amount", "lines" (o None), "lines_error"}] y,
    por cada factura que no se pudo leer con seguridad,
    {"error", "invoice_no" (el leído, o None)}.
    - N°: el del encabezado, si una línea de confirmación ("Inv# 668752
      $1,831.05") lo trae con el mismo importe; si no, el de la única línea
      de confirmación con el importe de la factura.
    - Importe: la suma de los renglones (que además da el Total Sales y el
      Invoice Total impresos) o, si los renglones no cierran, el Invoice
      Total confirmado por la línea de confirmación o por el Total Sales.
    """
    images = _ticket_images(pdf_path)
    if not images:
        raise ValueError("no se encontró la imagen escaneada de la factura de Gold Coast.")
    return _gce_invoices(images, lambda n: _ticket_readings(images, n), os.path.basename(pdf_path))


def extract_gce_lines(pdf_path, invoices=None):
    """
    Renglones de una factura de Gold Coast Eagle (LINE_EXTRACTORS). La carga
    de Proveedores ya los trae de read_gce_invoices; esto queda para quien
    pida solo el detalle. Con `invoices`, el de la factura con ese N°.
    """
    found = [inv for inv in read_gce_invoices(pdf_path) if "error" not in inv]
    wanted = {str(inv["invoice_no"]) for inv in (invoices or [])}
    for invoice in found:
        if (not wanted or str(invoice["invoice_no"]) in wanted) and invoice["lines"] is not None:
            return {"invoice_no": str(invoice["invoice_no"]), "lines": invoice["lines"],
                    "subtotal": None, "total": invoice["amount"]}
    error = next((inv["lines_error"] for inv in found if inv.get("lines_error")), None)
    raise ValueError(error or "no se encontró en el PDF la factura de Gold Coast pedida.")


# --- Red Bull ------------------------------------------------------------------
# Ticket de una página. Cada producto ocupa dos renglones:
#   ID QTY UNITS DESCRIPTION PRICE DEP DISC SUGAR TOTAL
#   UPC (y el código de barras dibujado, que el OCR mezcla con los montos)
# Cuentas: TOTAL = QTY x (PRICE - DISC + DEP + SUGAR); UNITS = QTY x unidades
# por caja; el UPC pasa el dígito verificador; la suma de TOTAL = INVOICE (o
# "Subtotal", en el pie de jun-jul 2026); TOTAL DUE = eso más depósito,
# impuestos y cargos; "Cases Delivered" = suma de QTY y "Units Delivered" =
# suma de UNITS.

_RB_FIELDS = ("price", "dep", "disc", "sugar", "ext")
_RB_HEADERS = (("price", r"PRICE"), ("dep", r"DEP"), ("disc", r"DISC"), ("sugar", r"SUGA\w"), ("ext", r"TOTAL"))
_RB_UPC = re.compile(r"\d{11,13}")


# Tamaño de la lata pegado a "OZ", que el OCR lee "0Z", "02" o "07"
# ("RED BULL 1202 LS", "8.402 LS", "840ZLS"). Red Bull vende latas de 8.4,
# 12, 16 y 20 oz: "84" es 8.4 sin el punto y 6.4 no existe (el 8 leído 6).
_RB_SIZE = re.compile(r"(?<![\d.])(\d{1,2})(?:[.,](\d))?\s*[O0Q][Z27](?:\s*(L\S?\S?))?")
_RB_SIZE_FIX = {"84": "8.4", "6.4": "8.4"}


def _rb_description(description):
    """(descripción, tamaño) con el tamaño separado: "COCONUT 8.40ZLS" -> ("COCONUT 8.4 OZ LS", "8.4 OZ")."""
    match = _RB_SIZE.search(description)
    if match is None:
        return description, ""
    size = match.group(1) + (f".{match.group(2)}" if match.group(2) else "")
    size = _RB_SIZE_FIX.get(size, size) + " OZ"
    name = description[:match.start()].strip()
    rest = description[match.end():].strip()
    text = " ".join(part for part in (name, size, "LS" if match.group(3) else "", rest) if part)
    return text, size


def _rb_product(row, columns):
    """Renglón ID/QTY/UNITS/.../TOTAL leído en una pasada, o None si no es un renglón de producto."""
    words = [w for w in row["words"] if re.search(r"[0-9A-Za-z$]", w["text"])]
    if len(words) < 6 or not re.fullmatch(r"[A-Za-z|]{0,3}\d{3,9}", _clean_token(words[0]["text"])):
        return None
    first = next((i for i, w in enumerate(words) if i >= 3 and re.match(r"^[«|]*[$S§]\d", w["text"])), None)
    if first is None:
        return None
    # QTY y UNITS: los dos primeros números antes de la descripción.
    numbers, start = [], 1
    for index in range(1, first):
        text = words[index]["text"]
        if re.search(r"[A-Za-z]{2}", text) or len(numbers) == 2:
            break
        digits = re.sub(r"\D", "", _fix_digits(text))
        if digits:
            numbers.append(digits)
        start = index + 1
    money = [dict(w, text=re.sub(r"^[«|]*[S§](?=\d)", "$", w["text"])) for w in words[first:]]
    return dict(
        _ticket_assign([w for w in money if w["text"].startswith("$")], columns.get(row["page"]), _RB_FIELDS),
        page=row["page"], y=row["y"], height=row["height"], upc=None, upc_box=None,
        qty=int(numbers[0]) if numbers and len(numbers[0]) <= 2 else None,
        units_total=int(numbers[1]) if len(numbers) > 1 and len(numbers[1]) <= 4 else None,
        # Sin los pedazos del precio que perdió el "$" ("RED BULL 8.40ZLS 00").
        description=re.sub(r"(\s+[\d.,$]+)+$", "", " ".join(w["text"] for w in words[start:first])),
    )


def _rb_close(values):
    price, dep, disc, sugar, ext = (values[f] for f in _RB_FIELDS)
    net = round(price - disc + dep + sugar, 2)
    if net <= 0 or min(price, dep, disc, sugar) < 0 or dep + sugar >= price:
        return None
    qty = _whole_qty(net, ext)
    return None if qty is None else dict(values, net=net, qty=qty)


_RB_DERIVE = (
    ("sugar", lambda v, qty: v["ext"] / qty - (v["price"] - v["disc"] + v["dep"]), True),
    ("dep", lambda v, qty: v["ext"] / qty - (v["price"] - v["disc"] + v["sugar"]), True),
    ("disc", lambda v, qty: v["price"] + v["dep"] + v["sugar"] - v["ext"] / qty, True),
    ("ext", lambda v, qty: qty * (v["price"] - v["disc"] + v["dep"] + v["sugar"]), True),
)
# En Red Bull la confusión va en un solo sentido: un 5 impreso sale 6 (en
# montos y en el UPC; nunca al revés), y a veces en tres campos del mismo
# renglón ("$64.70 / 3.26 / 51.46" por $54.70 / 3.25 / 51.45). Como solo se
# prueban los 6, se pueden cambiar hasta tres montos. Todos los UPC de Red
# Bull empiezan con el prefijo de la empresa, 611269.
_RB_CONFUSIONS = {"6": "5"}
_RB_RULES = {
    "options": (_RB_FIELDS, _rb_close, _RB_DERIVE, _RB_CONFUSIONS, 3),
    "upc": (_RB_CONFUSIONS, "611269"),
}


def _rb_footer(text):
    found = {}
    patterns = (
        ("subtotal", r"(?:INVOICE|Subtotal)"), ("total_due", r"(?:TOTAL\s*DUE|Invoice\s*Total)"),
        ("deposit", r"(?:DEPOSIT|Can Deposit)"), ("tax", r"(?:TAX|Sales Tax)"), ("sugar", r"Sugar Tax"),
        ("fees", r"Fees"),
    )
    for key, label in patterns:
        match = re.search(label + r"\W*\$?\s*([\d,]+[.,]\d{2})(?!\d)", text)
        if match:
            found[key] = _ticket_amount(match.group(1))
    for key, label in (("cases", r"Cases\s*Del\w*"), ("units", r"Un\w{2,3}\s*Del\w*"), ("skus", r"SKU\W{0,2}s")):
        match = re.search(label + r"\W*(\d+)", text)
        if match:
            found[key] = int(match.group(1))
    return found


_RB_NUMBER = re.compile(r"\bInv\w{0,6}\W{1,3}(\d{9,11})(?!\d)")
_RB_DATE = re.compile(r"(?<!\d)(\d{2})/(\d{2})/(20\d{2})\s+\d{1,2}:\d{2}")


def _rb_date(text):
    """Fecha del renglón del vendedor ("Salesman: Ryan Burris 02/16/2026 3:11 PM"), o None."""
    match = _RB_DATE.search(text)
    if match is None:
        return None
    try:
        return datetime(int(match.group(3)), int(match.group(1)), int(match.group(2)))
    except ValueError:
        return None


def _rb_blocks(readings):
    """
    Arma las facturas del PDF (cada una termina en su TOTAL DUE) con todas las
    pasadas hechas hasta ahora: renglones de producto (agrupados entre
    pasadas), pie, N° y fecha leídos. Devuelve (facturas, columnas).
    """
    all_rows = [row for reading in readings for row in reading]
    tolerance = _median([row["height"] for row in all_rows]) if all_rows else 10
    columns = _ticket_columns(readings, _RB_HEADERS)
    ends = [{"page": row["page"], "y": row["y"]} for row in all_rows if "total_due" in _rb_footer(row["text"])]
    boundaries = [(c[0]["page"], max(e["y"] for e in c) + tolerance) for c in _ticket_clusters(ends, tolerance * 5)]
    if not boundaries:
        raise ValueError("no se encontró el pie (TOTAL DUE) de la factura de Red Bull.")

    def block_of(row):
        return next((i for i, (page, y) in enumerate(boundaries) if (row["page"], row["y"]) <= (page, y)), None)

    blocks = [{"end": end, "products": [], "footer": {}, "numbers": {}, "dates": {}} for end in boundaries]
    for reading in readings:
        for position, row in enumerate(reading):
            index = block_of(row)
            if index is None:
                continue
            block = blocks[index]
            product = _rb_product(row, columns)
            if product is not None:
                # UPC: en los dos renglones siguientes de esta pasada (antes del próximo producto).
                for following in reading[position + 1:position + 3]:
                    if _rb_product(following, columns) is not None:
                        break
                    word = next((w for w in following["words"]
                                 if _RB_UPC.fullmatch(re.sub(r"\D", "", _fix_digits(w["text"])))), None)
                    if word is not None:
                        digits = re.sub(r"\D", "", _fix_digits(word["text"]))
                        product["upc"], product["upc_raw"] = _ticket_upc(digits), _ticket_raw_upc(digits)
                        product["upc_box"] = (word["x0"], word["top"], word["x1"], word["bottom"])
                        break
                block["products"].append(product)
                continue
            for key, value in _rb_footer(row["text"]).items():
                block["footer"].setdefault(key, []).append(value)
            match = _RB_NUMBER.search(row["text"])
            if match:
                block["numbers"][match.group(1)] = block["numbers"].get(match.group(1), 0) + 1
            found = _rb_date(row["text"])
            if found:
                block["dates"][found] = block["dates"].get(found, 0) + 1
    for block in blocks:
        block["clusters"] = _ticket_clusters(block["products"], tolerance)
    return blocks, columns


def _rb_charges(footer):
    return sum(_ticket_most_voted(footer.get(key, [])) or 0.0 for key in ("deposit", "tax", "sugar", "fees"))


def _rb_resolve(block, images, cache, columns):
    """Detalle de una factura de Red Bull, o ValueError si algo no cierra."""
    clusters, footer = block["clusters"], block["footer"]
    rows = _ticket_rows(clusters, images, cache, lambda row: _rb_product(row, columns), "Red Bull", _RB_RULES)
    if not rows:
        raise ValueError("no se ven los renglones de producto de la factura.")
    charges = _rb_charges(footer)
    printed = _ticket_top(footer.get("subtotal", [])) + [round(t - charges, 2) for t in _ticket_top(footer.get("total_due", []))
                                                         if t is not None]
    try:
        subtotal = _ticket_fix_sum(rows, printed, "INVOICE/Subtotal", "Red Bull")
    except ValueError:
        # Los renglones, tal como cerraron cada uno, suman el total impreso con
        # 5 donde se leyó 6 ($670.15 leído $670.16 en todas las pasadas).
        subtotal = round(sum(r["options"][r["choice"]][1]["ext"] for r in rows), 2)
        if not any(_rb_same(f"{v:.2f}", f"{subtotal:.2f}") for v in printed if v is not None):
            raise
    chosen = _ticket_chosen(rows)
    for c in chosen:
        c["per_case"] = _ticket_most_voted([cand["units_total"] // c["qty"] for cand in c["cluster"]
                                            if cand.get("units_total") and cand["units_total"] % c["qty"] == 0])
    # Un solo renglón con UNITS ilegible (pegado a la cantidad: "3s72"): sus
    # unidades son el Units Delivered del pie menos las de los demás renglones,
    # si eso da un número entero de unidades por caja.
    missing = [c for c in chosen if c["per_case"] is None]
    delivered = _ticket_most_voted(footer.get("units", []))
    deduced = False
    if len(missing) == 1 and delivered:
        rest = delivered - sum(c["qty"] * c["per_case"] for c in chosen if c["per_case"] is not None)
        if rest > 0 and rest % missing[0]["qty"] == 0:
            missing[0]["per_case"], deduced = rest // missing[0]["qty"], True
    for c in chosen:
        if c["per_case"] is None:
            raise ValueError(f"no se pudieron leer las unidades por caja de {c['description']}.")
    units_sum = sum(c["qty"] * c["per_case"] for c in chosen)
    units_ok = units_sum in footer.get("units", [])
    # Con unas unidades deducidas del pie, el pie ya no puede confirmar las cantidades por unidades.
    _ticket_check_qty(chosen, (units_ok and not deduced) or sum(c["qty"] for c in chosen) in footer.get("cases", []),
                      "Red Bull", "Cases Delivered o Units Delivered")
    if not units_ok:
        raise ValueError(f"las unidades ({units_sum}) no coinciden con el Units Delivered del pie.")
    lines = []
    for c in chosen:
        per_case = c["per_case"]
        description, size = _rb_description(c["description"])
        lines.append(_line(
            len(lines) + 1, upc=c["upc"], description=description,
            qty=c["qty"], pack=per_case, size=size, units=per_case,
            price=c["price"], allowance=c["disc"] or None, net=c["net"], ext=c["ext"],
        ))
    total = next((round(subtotal + charges, 2) for t in _ticket_top(footer.get("total_due", []))
                  if t is not None and _rb_same(f"{t:.2f}", f"{subtotal + charges:.2f}")), None)
    if total is None:
        raise ValueError(f"el subtotal (${subtotal:,.2f}) más los cargos no da el TOTAL DUE leído.")
    return {"lines": lines, "subtotal": subtotal, "total": total}


def _rb_footer_total(footer):
    """
    TOTAL DUE de una factura cuyos renglones no cerraron: vale si el más
    leído coincide con el INVOICE/Subtotal más leído más los cargos (dos
    renglones distintos del pie) y no tiene ningún 6 (en Red Bull el 5
    impreso sale 6 en todas las lecturas: sin los renglones, $670.16 puede
    ser $670.15). None si no.
    """
    total = _ticket_most_voted(footer.get("total_due", []))
    subtotal = _ticket_most_voted(footer.get("subtotal", []))
    if (total is not None and subtotal is not None and "6" not in f"{total:.2f}"
            and abs(round(subtotal + _rb_charges(footer), 2) - total) < 0.005):
        return total
    return None


def _rb_same(read, printed):
    """El N° (o la fecha) leído puede ser el impreso: los mismos dígitos, salvo 6 leídos donde dice 5."""
    return len(read) == len(printed) and all(a == b or (a == "6" and b == "5") for a, b in zip(read, printed))


def _rb_number(block, from_name):
    """
    N° de la factura. En Red Bull el 5 impreso sale 6 casi siempre (también
    en el N°: 2035255085 se lee 2036256086), así que lo leído vale si lo
    confirma el nombre del archivo (el mismo N°, o con 5 donde se leyó 6), o
    si todas las pasadas (3 o más) leyeron igual un N° sin ningún 6. None si no.
    """
    votes = block["numbers"]
    named = {name for read in votes for name in from_name if _rb_same(read, name)}
    if len(named) == 1:
        return named.pop()
    header = _ticket_winner(votes)
    if header is not None and "6" not in header and votes[header] >= 3 and len(votes) == 1:
        return header
    return None


def _rb_date_of(block, filename):
    """
    Fecha de la factura, con el mismo cuidado que el N°: la del nombre del
    archivo si es la leída (o la leída con 6 donde dice 5, también en el año:
    2025 sale 2026, y a lo sumo otro dígito distinto: 29 sale 28); si no, la
    leída, siempre que no tenga ningún 6 en el día o el mes. None si no.
    """
    named = _filename_dates(filename)
    for read in sorted(block["dates"], key=lambda d: -block["dates"][d]):
        for name in named:
            pairs = list(zip(read.strftime("%m%d%Y"), name.strftime("%m%d%Y")))
            if sum(a != b and not (a == "6" and b == "5") for a, b in pairs) <= 1:
                return name
    found = _ticket_winner(block["dates"])
    if found is not None and "6" in found.strftime("%m%d"):
        return None
    return found


def _rb_invoices(images, reading, filename, cache=None):
    """Todas las facturas de Red Bull de un PDF (ver read_red_bull_invoices)."""
    readings, resolved, failures = [], [], []
    cache = {} if cache is None else cache
    results = None
    for passes_done in range(len(_TICKET_PASSES)):
        readings.append(reading(passes_done))
        try:
            blocks, columns = _rb_blocks(readings)
        except ValueError:
            continue
        for block in blocks:
            if _ticket_resolved(resolved, block["end"]) is None:
                try:
                    resolved.append((block["end"], _rb_resolve(block, images, cache, columns)))
                except ValueError as exc:
                    failures.append((block["end"], str(exc)))
        details = [_ticket_resolved(resolved, block["end"]) for block in blocks]
        results = _rb_assign(blocks, details, failures, filename)
        if all(details) and not any("error" in invoice for invoice in results):
            break
    if results is None:
        raise ValueError("no se encontró el pie (TOTAL DUE) de ninguna factura de Red Bull.")
    return results


def _rb_assign(blocks, details, failures, filename):
    """Encabezado de cada factura armada de Red Bull, como _gce_assign."""
    # El N° en el nombre del archivo, aunque tenga un dígito de más ("Invoice 22036941041").
    from_name = {run[i:i + 10] for run in re.findall(r"\d{10,12}", filename or "") for i in range(len(run) - 9)}
    results, taken = [], set()
    for block, detail in zip(blocks, details):
        number = _rb_number(block, from_name)
        if number is None or number in taken:
            if not block["products"] and not block["numbers"]:
                continue  # un pie suelto, sin nada de una factura
            results.append({"error": "no se pudo leer el N° de invoice con seguridad.", "invoice_no": None})
            continue
        amount = detail["total"] if detail is not None else _rb_footer_total(block["footer"])
        if amount is None:
            results.append({"error": "no se pudo leer el TOTAL DUE con seguridad.", "invoice_no": int(number)})
            continue
        found = _rb_date_of(block, filename)
        if found is None:
            results.append({"error": "no se pudo leer la fecha de la factura.", "invoice_no": int(number)})
            continue
        taken.add(number)
        results.append({"invoice_no": int(number), "date": found, "amount": amount,
                        "lines": detail["lines"] if detail is not None else None,
                        "lines_error": None if detail is not None else _ticket_resolved(failures, block["end"])})
    return results


def read_red_bull_invoices(pdf_path):
    """
    Todas las facturas de Red Bull de un PDF, con el mismo formato que
    read_gce_invoices. Importe: la suma de los renglones (que además da el
    INVOICE/Subtotal y el TOTAL DUE impresos) o, si los renglones no cierran,
    el TOTAL DUE confirmado por el INVOICE más los cargos.
    """
    images = _ticket_images(pdf_path)
    if not images:
        raise ValueError("no se encontró la imagen escaneada de la factura de Red Bull.")
    return _rb_invoices(images, lambda n: _ticket_readings(images, n), os.path.basename(pdf_path))


def extract_red_bull_lines(pdf_path, invoices=None):
    """Renglones de una factura de Red Bull (LINE_EXTRACTORS); ver extract_gce_lines."""
    found = [inv for inv in read_red_bull_invoices(pdf_path) if "error" not in inv]
    wanted = {str(inv["invoice_no"]) for inv in (invoices or [])}
    for invoice in found:
        if (not wanted or str(invoice["invoice_no"]) in wanted) and invoice["lines"] is not None:
            return {"invoice_no": str(invoice["invoice_no"]), "lines": invoice["lines"],
                    "subtotal": None, "total": invoice["amount"]}
    error = next((inv["lines_error"] for inv in found if inv.get("lines_error")), None)
    raise ValueError(error or "no se encontró en el PDF la factura de Red Bull pedida.")


# --- Frito-Lay (2026-10-05) ----------------------------------------------------
# Pedido del usuario (2026-10-05): sumar otro proveedor a la lectura de
# productos. Frito-Lay usa desde dic-2025 un ticket limpio ("E-CHECK SALE" /
# "CASH SALE") con un renglón por producto:
#   QTY UPC  SRP PKG-INFO  ITEM      COST    DISC-EACH DISC-TOTAL AMOUNT TAX
#   12 77063 2.79 DR FLM   00046073  1.9500  0.0000    0.00       23.40  *
# El UPC impreso es el código de producto de 5 dígitos: el UPC completo es
# 0 + 28400 (Frito-Lay) + código + dígito verificador (028400770637, el del
# POS). Lo que se lee limpio: ITEM y COST. La cantidad a veces queda pegada al
# UPC ("477018" = 4 x 77018) y el AMOUNT sale mal seguido ("23:40", "7280"),
# así que cada renglón se arma con COST x cantidad - descuento y se controla
# con el pie: GROSS SALES AMOUNT, TOTAL EXTENDED EACHES SOLD y el UNIT COST
# SUMMARY de la segunda página ("247 @ $1.9500": las unidades de cada costo).
# Particularidades vistas en las 15 facturas de dic-2025 a sep-2026:
# - renglones vendidos por caja: el ITEM baja a una segunda línea
#   ("6 26034 2.09 GH VANILLA 1.1900 ... / 1 6REG 00025294");
# - dips sin precio sugerido y con código corto ("2 26 NP FRO DIP"): el UPC
#   sale igual (028400000260, el del POS);
# - el código leído con basura ("7/060") no vota: igual cerraba la suma con
#   un dígito menos y daba un UPC equivocado.
# Las devoluciones (**RETURNS**, en el mismo documento o en otro con otro N°)
# no se cuentan. Los formatos viejos ("CHARGE SALES", hasta oct-2025) no
# tienen estas columnas y dan error (la factura entra sin productos).

_FL_ITEM = re.compile(r"0\d{7}")
_FL_COST = re.compile(r"\d{1,3}[.,]\d{4}")
_FL_SRP = re.compile(r"\d{1,2}[.,]\d{2}")
# "INVOICE #: 96889721", también "INVOICE 4: ..." o con un espacio en el medio ("968897 21").
_FL_DOC = re.compile(r"INVOICE\s*[#4]?\s*[:;]\s*(\d[\d ]{4,10}\d)|INVOICE\s*#\s*(\d{6,10})", re.IGNORECASE)
_FL_SUMMARY_TITLE = re.compile(r"UNIT\s*COST\s*SU\w{3,5}Y", re.IGNORECASE)
_FL_SUMMARY = re.compile(r"(\d{1,4})\s*@\s*\$\s*(\d{1,3}[.,]\d{4})")


def _fl_cost(text):
    """'1.8800' / '«4.0400.' -> float, o None."""
    match = re.fullmatch(r"\W*(\d{1,3}[.,]\d{4})\W*", _fix_digits(_clean_token(text or "")))
    return float(match.group(1).replace(",", ".")) if match else None


def _fl_upc(code):
    """Código de producto de Frito-Lay (2 a 5 dígitos) -> UPC-A completo con el prefijo 28400."""
    digits = "028400" + code.zfill(5)
    check = (10 - sum(int(d) * (3 if i % 2 == 0 else 1) for i, d in enumerate(digits)) % 10) % 10
    return digits + str(check)


def _fl_product(row):
    """Renglón de producto leído en una pasada, o None."""
    words = [w for w in row["words"] if re.search(r"[0-9A-Za-z]", w["text"])]
    texts = [_fix_digits(_clean_token(w["text"])) for w in words]
    at = next((i for i, t in enumerate(texts) if _FL_ITEM.fullmatch(t)), None)
    item = texts[at] if at is not None else None
    if at is None:
        # El ITEM a veces queda corrido al renglón de abajo: sin él, el
        # renglón se ancla en el COST (y tiene que tener precio sugerido).
        cost_at = next((i for i, t in enumerate(texts) if _fl_cost(t)), None)
        if cost_at is None or not (any(_FL_SRP.fullmatch(t) for t in texts[:cost_at])
                                   or (cost_at > 2 and texts[0].isdigit() and texts[1].isdigit())):
            return None
        at = cost_at - 1
    elif at == 2 and re.fullmatch(r"[\dGSBOIl]{1,3}REG", words[1]["text"]):
        return None  # "1 6REG 00025294": la segunda línea de un renglón vendido por caja
    if at < 2 or len(texts) < at + 2:
        return None
    srp_at = next((i for i in range(at) if _FL_SRP.fullmatch(texts[i])), None)
    has_srp = srp_at is not None
    if srp_at is None:
        # Sin precio sugerido ("2 26 NP FRO DIP"): la cantidad y el código son
        # los números de adelante; la descripción empieza en la primera palabra.
        srp_at = next((i for i in range(at) if not texts[i].isdigit()), at)
        if srp_at == 0:
            return None
    # Cantidad y código: solo números limpios ("7/060" no es 7060 ni 77060:
    # ese renglón no vota el código en esta pasada).
    lead = [t for t in texts[:srp_at] if re.search(r"\d", t)]
    pairs = []
    if any(not t.isdigit() for t in lead[-2:]):
        lead = []
    if len(lead) >= 2 and len(lead[-2]) <= 3 and 2 <= len(lead[-1]) <= 5:
        pairs.append((int(lead[-2]), lead[-1], True))  # cantidad y código separados: lectura limpia
    elif lead and 5 <= len(lead[-1]) <= 8:
        # Pegados ("477018"): cada corte posible (1 a 3 dígitos de cantidad, 4 o 5 de código).
        joined = lead[-1]
        pairs.extend((int(joined[:q]), joined[q:], False) for q in (1, 2, 3)
                     if 4 <= len(joined) - q <= 5 and int(joined[:q]) > 0)
    after = texts[at + 1:]
    if not pairs and not (after and _fl_cost(after[0]) is not None):
        return None
    return {
        "page": row["page"], "y": row["y"], "height": row["height"],
        "item": item, "srp": _ticket_amount(texts[srp_at]) if has_srp else None,
        "description": " ".join(w["text"] for w in words[srp_at + has_srp:at + (item is None)]),
        "cost": _fl_cost(after[0]) if after else None,
        "disc_each": _fl_cost(after[1]) if len(after) > 1 else None,
        "disc_total": _ticket_amount(after[2]) if len(after) > 2 else None,
        "amount": _ticket_amount(after[3]) if len(after) > 3 else None,
        "pairs": pairs,
    }


def _fl_doc(text):
    """N° de documento de un renglón ("INVOICE #: 96889721"), o None."""
    match = _FL_DOC.search(text)
    if not match:
        return None
    digits = re.sub(r"\s", "", match.group(1) or match.group(2))
    return int(digits) if 6 <= len(digits) <= 10 else None


def _fl_target(docs, wanted):
    """
    N° del documento de ventas en una pasada: el igual a `wanted`; si esta
    pasada lo leyó mal, el único a 1 o 2 dígitos de distancia (los otros
    documentos del PDF tienen otro N°); sin `wanted`, el primero.
    """
    if not docs:
        return None
    if wanted is None:
        return docs[0]
    if wanted in docs:
        return wanted
    near = {d for d in docs if len(str(d)) == len(str(wanted))
            and sum(a != b for a, b in zip(str(d), str(wanted))) <= 2}
    return near.pop() if len(near) == 1 else None


def _fl_collect(readings, wanted):
    """
    Renglones de producto (agrupados entre pasadas) y pie del documento de
    ventas `wanted` (el N° de la factura; None = el primero del PDF).
    """
    products, footer = [], {"gross": [], "eaches": [], "summary": {}}
    docs_by_reading = [[d for d in (_fl_doc(row["text"]) for row in reading) if d is not None] for reading in readings]
    found = {d for docs in docs_by_reading for d in docs}
    if wanted is None and found:
        # Sin el N° del encabezado: el documento que aparece primero en más pasadas.
        wanted = _ticket_winner(_ticket_votes([{"doc": docs[0]} for docs in docs_by_reading if docs], "doc"))
    single = wanted is not None and bool(found) and all(_fl_target([d], wanted) is not None for d in found)
    for reading, docs in zip(readings, docs_by_reading):
        if single and not docs:
            # Esta pasada no leyó ningún N° y el PDF trae un solo documento:
            # todo lo que lee es de esa factura.
            target = doc = wanted
        else:
            target, doc = _fl_target(docs, wanted), None
        section = summary_part = None
        summary = {}
        for row in reading:
            text = row["text"]
            number = _fl_doc(text)
            if number is not None:
                doc = number
            is_summary = bool(_FL_SUMMARY_TITLE.search(text))
            if re.search(r"RETURNS\s*\*", text) and not is_summary and summary_part is None:
                section = "returns"
            elif re.search(r"\*\s*SALES\s*\*", text) and not is_summary and summary_part is None:
                section = "sales"
            if doc is None or doc != target:
                continue
            # UNIT COST SUMMARY: "**SALES** 6 @ $1.9800" ... y después
            # "**RETURNS** 17 @ $1.9500": solo cuentan los de ventas.
            if is_summary or (summary_part and re.search(r"SALES\s*\*", text)):
                summary_part = "sales"
            if summary_part and re.search(r"RETURNS", text):
                summary_part = "returns"
            if re.search(r"SALES FOR WEEK|YTD", text, re.IGNORECASE):
                summary_part = None
            if summary_part == "sales":
                for qty, cost in _FL_SUMMARY.findall(text):
                    key = float(cost.replace(",", "."))
                    summary[key] = summary.get(key, 0) + int(qty)
            match = re.search(r"GROSS\s*SALES\s*A\w{3,6}\W*\$?\s*([\d,]+[.,]\d{2})", text, re.IGNORECASE)
            if match:
                footer["gross"].append(_ticket_amount(match.group(1)))
            # Las unidades de los renglones son las "extendidas": un renglón
            # vendido por caja ("26 76929 ... / 1 26REG") cuenta 26 acá y 0 en
            # el TOTAL EACHES SOLD.
            match = re.search(r"TOTAL\s*EXTENDED\s*EACHES\s*SOL\w\W*(\d+)", text, re.IGNORECASE)
            if match:
                footer["eaches"].append(int(match.group(1)))
            if section == "sales":
                product = _fl_product(row)
                if product is not None:
                    products.append(product)
        if summary:
            key = tuple(sorted(summary.items()))
            footer["summary"][key] = footer["summary"].get(key, 0) + 1
    all_rows = [row for reading in readings for row in reading]
    tolerance = _median([row["height"] for row in all_rows]) if all_rows else 10
    return _ticket_clusters(products, tolerance), footer


def _fl_options(cluster, costs=None):
    """
    Combinaciones (cantidad, código) de un renglón con su total
    (COST x cantidad - descuento), de la más creíble a la menos: primero las
    que dan un AMOUNT leído, después las lecturas limpias.
    """
    votes = _ticket_votes(cluster, "cost")
    if costs:
        # Solo los costos del UNIT COST SUMMARY (4.0400 leído 4.0800 no vale).
        votes = {cost: n for cost, n in votes.items() if cost in costs}
    cost = _ticket_winner(votes)
    if cost is None:
        return cost, []
    disc_total = _ticket_winner(_ticket_votes(cluster, "disc_total"))
    if disc_total is None:
        if _ticket_winner(_ticket_votes(cluster, "disc_each")) != 0.0:
            return cost, []
        disc_total = 0.0
    amounts = _ticket_votes(cluster, "amount")
    weight = {}
    for cand in cluster:
        for qty, code, clean in cand["pairs"]:
            weight[(qty, code)] = weight.get((qty, code), 0) + (2 if clean else 1)
    options = []
    for (qty, code), votes in weight.items():
        ext = round(qty * cost - disc_total, 2)
        options.append((amounts.get(ext, 0), votes, qty, code, ext))
    options.sort(key=lambda o: (-o[0], -o[1]))
    return cost, options


def _fl_resolve(clusters, footer):
    """Renglones de la factura de Frito-Lay, o ValueError si no cierran contra el pie."""
    if not clusters:
        raise ValueError("no se ven los renglones de producto de la factura (¿formato viejo de Frito-Lay?).")
    gross = _ticket_most_voted(footer["gross"])
    summary = _ticket_winner(footer["summary"])
    # El UNIT COST SUMMARY trae las unidades de cada costo: su suma son las
    # unidades vendidas (las mismas que el TOTAL EXTENDED EACHES SOLD).
    eaches = sum(qty for _, qty in summary) if summary else _ticket_most_voted(footer["eaches"])
    if gross is None or eaches is None:
        raise ValueError("no se pudo leer el GROSS SALES AMOUNT o el TOTAL EACHES SOLD de Frito-Lay.")
    rows = []
    for line_no, cluster in enumerate(clusters, start=1):
        cost, options = _fl_options(cluster, {c for c, _ in summary} if summary else None)
        item = _ticket_winner(_ticket_votes(cluster, "item"))
        if cost is None or not options:
            raise ValueError(f"el renglón {line_no} de Frito-Lay no se pudo leer (cantidad, costo o ITEM).")
        rows.append({"cluster": cluster, "cost": cost, "item": item, "options": options})

    def fits(choice):
        qty = sum(r["options"][i][2] for r, i in zip(rows, choice))
        ext = round(sum(r["options"][i][4] for r, i in zip(rows, choice)), 2)
        if qty != eaches or abs(ext - gross) >= 0.005:
            return False
        if summary is not None:
            by_cost = {}
            for r, i in zip(rows, choice):
                by_cost[r["cost"]] = by_cost.get(r["cost"], 0) + r["options"][i][2]
            if tuple(sorted(by_cost.items())) != summary:
                return False
        return True

    choice = [0] * len(rows)
    if not fits(choice):
        # Probar otra combinación en hasta dos renglones (la cantidad pegada al
        # código cortada en otro lugar); vale si una sola cierra todo el pie.
        alternatives = [(n, i) for n, r in enumerate(rows) for i in range(1, len(r["options"]))]
        found = []
        for size in (1, 2):
            for combo in itertools.combinations(alternatives, size):
                if len({n for n, _ in combo}) != size:
                    continue
                trial = list(choice)
                for n, i in combo:
                    trial[n] = i
                if fits(trial):
                    found.append(trial)
            if found:
                break
        if len(found) != 1:
            total = round(sum(r["options"][0][4] for r in rows), 2)
            raise ValueError(f"los renglones de Frito-Lay suman ${total:,.2f} y el GROSS SALES impreso es ${gross:,.2f} "
                             f"(o no dan las {eaches} unidades del pie).")
        choice = found[0]
    lines = []
    for r, i in zip(rows, choice):
        _, _, qty, code, ext = r["options"][i]
        cluster = r["cluster"]
        disc_each = _ticket_winner(_ticket_votes(cluster, "disc_each")) or 0.0
        description = (_ticket_winner(_ticket_votes(cluster, "description"))
                       or next((c["description"] for c in cluster if c.get("description")), ""))
        lines.append(_line(
            len(lines) + 1, upc=_fl_upc(code), item_no=r["item"], description=description,
            qty=qty, pack=1, size="", units=1, price=r["cost"], allowance=disc_each or None,
            net=round(r["cost"] - disc_each, 4), ext=ext,
            srp=_ticket_winner(_ticket_votes(cluster, "srp")),
        ))
    return lines, gross


def _fl_text_date(text):
    """'01 Jan 2026' -> datetime; el OCR cambia la M del mes por H o N ('25 Har 2026')."""
    match = re.search(r"DATE\W*(\d{1,2})\s+([A-Za-z]{3})\w*\s+(\d{4})", text, re.IGNORECASE)
    if match is None:
        return None
    day, month, year = match.groups()
    found = set()
    for first in {month[0], "M", "N"} if month[0] in "HMNhmn" else {month[0]}:
        try:
            found.add(datetime.strptime(f"{day} {first}{month[1:]} {year}", "%d %b %Y"))
        except ValueError:
            pass
    return found.pop() if len(found) == 1 else None


def read_frito_lay_header(pdf_path):
    """
    Respaldo del encabezado de Frito-Lay (proveedores._extract_frito_lay_invoice
    lee la página entera de una vez y a veces pierde el "INVOICE #" o el
    importe del "TOTAL DUE:", visto en la factura 60895948): N°, fecha y
    TOTAL DUE del primer documento del PDF, con las pasadas del lector de
    tickets. Cada dato vale si lo leen igual dos pasadas, o una y el nombre
    del archivo. ValueError si no.
    """
    images = _ticket_images(pdf_path)
    filename = os.path.basename(pdf_path)
    votes = {"invoice_no": {}, "date": {}, "amount": {}}
    for passes_done in range(min(3, len(_TICKET_PASSES))):
        doc = None
        found = {}
        for row in _ticket_readings(images, passes_done):
            number = _fl_doc(row["text"])
            if number is not None:
                if doc is not None and number != doc:
                    break  # otro documento (devoluciones): ya no es esta factura
                doc = found["invoice_no"] = number
            if found.get("date") is None:
                found["date"] = _fl_text_date(row["text"])
            match = re.search(r"TOTAL\s*DUE\W*\$?\s*([\d,]+[.,]\d{2})", row["text"], re.IGNORECASE)
            if match and "amount" not in found:
                found["amount"] = _ticket_amount(match.group(1))
        for field, value in found.items():
            if value is not None:
                votes[field][value] = votes[field].get(value, 0) + 1
        named_numbers = {int(n) for n in re.findall(r"\d{6,10}", filename)}
        named_dates = _filename_dates(filename)
        result = {}
        for field, counts in votes.items():
            for value, n in sorted(counts.items(), key=lambda kv: -kv[1]):
                if n >= 2 or (field == "invoice_no" and value in named_numbers) or (field == "date" and value in named_dates):
                    result[field] = value
                    break
        if len(result) == 3:
            return result
    raise ValueError("no se pudo leer invoice/fecha/total del PDF de Frito-Lay.")


def extract_frito_lay_lines(pdf_path, invoices=None):
    """
    Renglones de una factura de Frito-Lay (formato desde dic-2025). Con
    `invoices` (lo que leyó el encabezado de la app), los del documento con
    ese N°; la suma tiene que dar el GROSS SALES AMOUNT impreso.
    """
    images = _ticket_images(pdf_path)
    if not images:
        raise ValueError("no se encontró la imagen escaneada de la factura de Frito-Lay.")
    wanted = int(invoices[0]["invoice_no"]) if invoices else None
    readings, error, pages = [], None, range(len(images))
    for passes_done in range(len(_TICKET_PASSES)):
        config, scale, prep = _TICKET_PASSES[passes_done]
        readings.append([_ticket_row(page, words) for page in pages
                         for words in _ocr_rows(images[page], config, scale=scale, prep=prep)])
        if passes_done == 0:
            if not any(re.search(r"ITE[MN]\s+COST|EACH\s+UPC", row["text"]) for row in readings[0]):
                raise ValueError("la factura de Frito-Lay no tiene el formato con ITEM y COST (formato viejo).")
            # Las pasadas siguientes, solo sobre las páginas con renglones o con el pie
            # (no la foto del cheque ni el documento de devoluciones de otro N°).
            pages = sorted({row["page"] for row in readings[0]
                            if _fl_product(row) or re.search(r"GROSS|EACHES|@\s*\$", row["text"])})
        try:
            clusters, footer = _fl_collect(readings, wanted)
            lines, gross = _fl_resolve(clusters, footer)
            return {"invoice_no": str(wanted) if wanted is not None else None, "lines": lines,
                    "subtotal": gross, "total": gross}
        except ValueError as exc:
            error = exc
    raise error


# --- Productos sin UPC: el nombre como clave (2026-10-05) -------------------------
# Pedido del usuario: los proveedores que se leen bien pero no traen el UPC del
# POS (J.J. Taylor, Midtown) se comparan contra la factura anterior por el
# NOMBRE. El nombre tiene que salir igual en cada factura: se arma con las
# palabras en las que coinciden la mayoría de las lecturas de OCR (la tinta a
# mano, las tildes de control, sale distinta en cada lectura; la letra
# impresa, igual) y se compara exacto, sin parecidos aproximados ("White Claw
# Black Cherry" y "White Claw Blackberry" son productos distintos). Un nombre
# que no coincide con el de la factura anterior queda como producto "Nuevo",
# nunca comparado contra otro.

# Proveedores cuyos productos se identifican por el nombre (product_key).
NAME_KEYED_SUPPLIERS = {"jj_taylor", "midtown"}


def _name_consensus(texts):
    """Nombre de un renglón con las palabras que leyeron igual la mayoría de las lecturas (todas, si son dos)."""
    lists = [re.sub(r"\s+", " ", t).strip(" |").split(" ") for t in texts if t and t.strip(" |")]
    if not lists:
        return ""
    need = len(lists) // 2 + 1
    counts = {}
    for words in lists:
        for word in set(words):
            counts[word] = counts.get(word, 0) + 1
    kept = [[w for w in words if counts[w] >= need] for words in lists]
    best = max(range(len(lists)), key=lambda i: (len(kept[i]), -i))
    return " ".join(kept[best]) or " ".join(lists[0])


def name_key(description):
    """
    Nombre normalizado para comparar productos sin UPC: mayúsculas, solo
    letras y números, y en las palabras sin números las confusiones fijas de
    la letra de J.J. Taylor plegadas a una sola forma ("Claw"/"Gaw",
    "White"/"While", "Ice"/"Iee", "Smir"/"$mir"). Los números y el resto de
    las letras tienen que coincidir exacto.
    """
    words = re.findall(r"[A-Z0-9$]+", (description or "").upper())
    folded = []
    for word in words:
        if not re.search(r"\d", word):
            word = word.replace("$", "S").replace("CL", "G").translate(str.maketrans("TIE", "LLC"))
        folded.append(word.replace("$", ""))
    return " ".join(w for w in folded if w)


# --- J.J. Taylor (2026-10-05) --------------------------------------------------
# Pedido del usuario (2026-10-05): sacar a J.J. Taylor de la pausa y comparar
# sus productos contra la factura anterior por el NOMBRE (no trae UPC), para
# pasarle al manager los que cambiaron de precio. Ticket de reparto escaneado:
#   Date Invoice Load Sheet PO Number Route Deliveryman Salesman
#   09/10/2026 6000658 199396 ...
#   DEL PRODUCT                         PRICE  DISC  NET    TOTAL
#   4   Hein 4/6/12 B                   $37.60       $37.60 $150.40
# Cuentas: PRICE - DISC = NET; DEL x NET = TOTAL; la suma de TOTAL (más el
# impuesto por galón de cerveza, cuando lo cobran) = Total; la suma de DEL =
# "Total U without pick ups"; la de DEL x piezas = "Piece Count" (piezas de
# una caja según el pack: "4/6/12" = 4 six-packs, "1/24/16" = 24 latas).
# Todos los montos llevan "$": uno sin "$" puede ser el "$" leído como dígito
# ("446.00" por $46.00) o el "$" perdido ("37.60"); se votan las dos lecturas.

_JJ_FIELDS = ("price", "disc", "net", "ext")
_JJ_AMOUNT = re.compile(r"[$S§]?-?\d{1,3}(?:,?\d{3})*[.,]\d{2}")
_JJ_PACK = re.compile(r"(?<![\d.])(\d{1,2})\s*/\s*(\d{1,3})(?:\s*/\s*(\d{1,3}(?:\.\d{1,2})?))?")


def _jj_amounts(text):
    """
    Lecturas posibles de un monto de J.J. Taylor: (la leída, la sin el primer
    dígito si no tenía "$"). Los centavos sueltos del descuento ("$.80") salen
    "$.80" o "$80".
    """
    raw = _clean_token(text)
    match = re.fullmatch(r"[$S§](?:[.,])?(\d{2})|[.,](\d{2})", raw)
    if match:
        return float("0." + (match.group(1) or match.group(2))), None
    if not _JJ_AMOUNT.fullmatch(raw):
        return None, None
    value = _ticket_amount(raw.lstrip("$S§"))
    if raw[0] in "$S§":
        return value, None
    rest = raw[1:]
    alt = _ticket_amount(rest) if raw[0] in "45689" and _TICKET_AMOUNT.fullmatch(rest) else None
    return value, alt


def _jj_units(description):
    """
    (piezas por caja, pack impreso) según el pack de la descripción --
    regla verificada contra el "Piece Count" del pie: "4/6/12" = 4
    (six-packs), "2/12/12" = 2 (twelve-packs), "1/24/16" = 24 (latas
    sueltas), "1/12/19.2" = 12.
    """
    matches = list(_JJ_PACK.finditer(description or ""))
    if not matches:
        return None, ""
    match = matches[-1]
    first, second = int(match.group(1)), int(match.group(2))
    size = description[match.start():].strip(" |.,")
    return (second if first == 1 else first), size


def _jj_product(row):
    """
    Renglón de producto leído en una pasada (lista de lecturas: la principal
    y, si hay, una alternativa por cada monto sin "$"), o None. Antes de la
    cantidad puede haber marcas sueltas del margen (";", "*,", "Bo"): la
    cantidad es el número justo antes de la descripción.
    """
    words = [w for w in row["words"] if w["text"].strip("|!") != ""]
    amounts = []
    for word in reversed(words):
        value, alt = _jj_amounts(word["text"])
        if value is None:
            break
        amounts.insert(0, (value, alt))
    if len(amounts) not in (3, 4):
        return None
    head = words[:len(words) - len(amounts)]
    start = next((i for i, w in enumerate(head) if re.search(r"[A-Za-z]{2}", w["text"])), None)
    if not start:
        return None
    # Entre la cantidad y la descripción puede quedar un guion o una coma sueltos ("2 - Mikes").
    before = [w for w in head[:start] if re.search(r"[0-9A-Za-z]", w["text"])]
    if not before:
        return None
    qty_text = _fix_digits(_clean_token(before[-1]["text"]))
    if not re.fullmatch(r"\d{1,3}", qty_text):
        return None
    description_words = head[start:]
    # Un monto ilegible ("$3590") no es parte de la descripción.
    cut = next((i for i, w in enumerate(description_words) if re.fullmatch(r"[$S§]\d+", _clean_token(w["text"]))),
               len(description_words))
    description_words = description_words[:cut]
    fields = ("price", "net", "ext") if len(amounts) == 3 else _JJ_FIELDS
    main = {field: value for field, (value, _) in zip(fields, amounts)}
    if len(amounts) == 3:
        main["disc"] = 0.0
    base = {"page": row["page"], "y": row["y"], "height": row["height"]}
    found = [dict(base, **main, qty=int(qty_text),
                  description=" ".join(w["text"] for w in description_words).strip(" |"))]
    for field, (_, alt) in zip(fields, amounts):
        if alt is not None:
            found.append(dict(base, **{field: alt}))
    return found


def _jj_close(values):
    price, disc, net, ext = (values[f] for f in _JJ_FIELDS)
    if abs(price - disc - net) > _TOLERANCE or net <= 0 or disc < 0 or ext <= 0:
        return None
    qty = _whole_qty(net, ext)
    return None if qty is None else dict(values, qty=qty)


_JJ_DERIVE = (
    ("net", lambda v, qty: v["price"] - v["disc"], True),
    ("disc", lambda v, qty: v["price"] - v["net"], False),
    ("ext", lambda v, qty: qty * v["net"], True),
)
# La letra del ticket es la de Red Bull: el 5 impreso sale 6 (46.00 por 45.00,
# 28.96 por 28.95), hasta en los cuatro montos del renglón a la vez ("$35.90
# $6.95 $28.95 $28.95" leído 36.90 / 6.96 / 28.96 / 28.96 en todas las pasadas).
_JJ_CONFUSIONS = {"6": "5"}
_JJ_RULES = {"options": (_JJ_FIELDS, _jj_close, _JJ_DERIVE, _JJ_CONFUSIONS, 4)}


def _jj_fix_sum(rows, printed):
    """
    _ticket_fix_sum, y si no alcanza: con el 5 leído 6 en muchos renglones a
    la vez (5 en 26 renglones de la factura 5431210), hasta 6 renglones
    cambiados por su lectura con 5 -- vale solo si una única combinación da
    un total impreso al centavo.
    """
    try:
        return _ticket_fix_sum(rows, printed, "Total", "J.J. Taylor")
    except ValueError as exc:
        error = exc
    ext = lambda r, i: r["options"][i][1]["ext"]
    current = round(sum(ext(r, r["choice"]) for r in rows), 2)
    changes = [(n, alt, round(ext(r, alt) - ext(r, r["choice"]), 2))
               for n, r in enumerate(rows) for alt in range(len(r["options"]))
               if r["options"][alt][2] == 1 and ext(r, alt) != ext(r, r["choice"])]
    if len(changes) > 24:
        raise error
    for size in range(4, 7):
        fixes = [combo for combo in itertools.combinations(changes, size)
                 if len({n for n, _, _ in combo}) == size
                 and round(current + sum(d for _, _, d in combo), 2) in printed]
        if len(fixes) == 1:
            for n, alt, _ in fixes[0]:
                rows[n]["choice"] = alt
            return round(current + sum(d for _, _, d in fixes[0]), 2)
        if fixes:
            break
    raise error


def _jj_footer(text):
    """Valores del pie que trae un renglón: {"cases"|"pieces"|"total"|"charge": valor}."""
    found = {}
    match = re.search(r"Total\s*U\w*\s*with\w*\s*pick\s*ups?\W*(\d+)(?:\s+\$?([\d,]+[.,]\d{2}))?", text, re.IGNORECASE)
    if match:
        found["cases"] = int(match.group(1))
        if match.group(2):
            found["total"] = _ticket_amount(match.group(2))
    elif re.search(r"Pick\s*up", text, re.IGNORECASE):
        amounts = re.findall(r"\$\s*([\d,]+[.,]\d{2})", text)
        if amounts:
            found["total"] = _ticket_amount(amounts[-1])
    match = re.search(r"Piece\s*Count\W*(\d+)", text, re.IGNORECASE)
    if match:
        found["pieces"] = int(match.group(1))
    match = re.search(r"Gallons\W*\$?\s*([\d,]+[.,]\d{2})", text, re.IGNORECASE)
    if match:
        found["charge"] = _ticket_amount(match.group(1))
    return found


def _jj_documents(readings):
    """
    Documentos del PDF por página: cada ticket ocupa su página (además de la
    factura pueden venir tickets de cambio, "SWAP", con Total $0.00); una
    página sin encabezado ("Load Sheet") sigue al documento anterior.
    Devuelve [{"pages", "clusters", "footer", "numbers", "dates"}].
    """
    pages = sorted({row["page"] for reading in readings for row in reading})
    headed = {row["page"] for reading in readings for row in reading if re.search(r"Load\s*Sh\w*|PO\s*Num\w*", row["text"], re.IGNORECASE)}
    groups = []
    for page in pages:
        if page in headed or not groups:
            groups.append([page])
        else:
            groups[-1].append(page)
    all_rows = [row for reading in readings for row in reading]
    tolerance = _median([row["height"] for row in all_rows]) if all_rows else 10
    documents = []
    for group in groups:
        products, footer, numbers, dates = [], {}, {}, {}
        for reading in readings:
            after_titles = in_footer = False
            for row in reading:
                if row["page"] not in group:
                    continue
                text = row["text"]
                if after_titles:
                    after_titles = False
                    for number in re.findall(r"(?<![\d/])(\d{7})(?!\d)", text):
                        numbers[number] = numbers.get(number, 0) + 1
                    for month, day, year in re.findall(r"(\d{1,2})/(\d{1,2})/(20\d{2})", text):
                        try:
                            found = datetime(int(year), int(month), int(day))
                            dates[found] = dates.get(found, 0) + 1
                        except ValueError:
                            pass
                    continue
                if re.search(r"Load\s*Sh\w*|PO\s*Num\w*", text, re.IGNORECASE):
                    after_titles = True
                    continue
                footer_values = _jj_footer(text)
                if footer_values or re.search(r"Gallons", text, re.IGNORECASE):
                    in_footer = True
                    for key, value in footer_values.items():
                        footer.setdefault(key, []).append(value)
                    continue
                if not in_footer:
                    found = _jj_product(row)
                    if found:
                        products.extend(found)
        documents.append({"pages": group, "clusters": _ticket_clusters(products, tolerance),
                          "footer": footer, "numbers": numbers, "dates": dates})
    return documents


def _jj_close_digits(read, printed):
    """El leído puede ser el impreso: mismos dígitos salvo 6 donde dice 5 y, a lo sumo, otro dígito distinto."""
    return len(read) == len(printed) and sum(a != b and not (a == "6" and b == "5") for a, b in zip(read, printed)) <= 1


def _jj_number(numbers, filename):
    """
    N° de la factura: el del nombre del archivo si alguna pasada lo leyó igual
    o con a lo sumo un dígito distinto (el 5 sale 6: "6431210" por 5431210);
    sin N° en el nombre, el de 3+ pasadas unánimes. None si no.
    """
    named = _filename_numbers(filename, 7)
    found = {name for name in named for number in numbers if _jj_close_digits(number, name)}
    if len(found) == 1:
        return found.pop()
    if not named and len(numbers) == 1:
        number, votes = next(iter(numbers.items()))
        if votes >= 3:
            return number
    return None


def _jj_date(dates, filename):
    """Fecha: la del nombre del archivo si la leída difiere en a lo sumo un dígito (6 por 5); si no, la más votada con 2+ votos."""
    named = _filename_dates(filename)
    for read in sorted(dates, key=lambda d: -dates[d]):
        for name in named:
            if _jj_close_digits(read.strftime("%m%d%Y"), name.strftime("%m%d%Y")):
                return name
    found = _ticket_winner(dates)
    return found if found is not None and dates[found] >= 2 else None


# PRICE, NET y TOTAL son el mismo número cuando no hay descuento y la cantidad es 1.
_JJ_FAMILIES = (("amount", ("price", "net", "ext")), ("disc", ("disc",)))


def _jj_prefer_five(cluster):
    """
    El 5 impreso sale 6, nunca al revés: si en el renglón se leyó un monto con
    5 y, en otra pasada o en la columna hermana (PRICE y NET son el mismo
    número cuando no hay descuento), el mismo monto con 6 en ese lugar, vale
    el del 5. Devuelve el grupo con las lecturas corregidas.
    """
    seen = {}
    for cand in cluster:
        for family, fields in _JJ_FAMILIES:
            for field in fields:
                if cand.get(field) is not None:
                    seen.setdefault(family, set()).add(round(cand[field], 2))

    def fixed(value, family):
        text = f"{value:.2f}"
        for index, char in enumerate(text):
            if char == "6":
                other = round(float(text[:index] + "5" + text[index + 1:]), 2)
                if other in seen.get(family, ()):
                    return fixed(other, family)
        return value

    result = []
    for cand in cluster:
        cand = dict(cand)
        for family, fields in _JJ_FAMILIES:
            for field in fields:
                if cand.get(field) is not None:
                    cand[field] = fixed(cand[field], family)
        result.append(cand)
    return result


def _jj_rows(clusters, images, cache):
    rows = []
    for line_no, cluster in enumerate(clusters, start=1):
        cluster = _jj_prefer_five(cluster)
        options = _ticket_options(cluster, *_JJ_RULES["options"])
        if not _ticket_settled(options) and images is not None:
            parse = lambda row: (_jj_product(row) or [None])[0]
            cluster = _jj_prefer_five(_ticket_reread_row(images, cluster, parse, cache))
            options = _ticket_options(cluster, *_JJ_RULES["options"])
        description = _name_consensus([c.get("description") for c in cluster])
        if not options:
            raise ValueError(f"el renglón {line_no} de J.J. Taylor ({description}) no cierra: "
                             "ninguna lectura da cantidad x neto = total.")
        if not _ticket_settled(options):
            raise ValueError(f"el renglón {line_no} de J.J. Taylor ({description}) tiene dos lecturas "
                             "posibles que cierran con los mismos votos.")
        rows.append({"options": options, "choice": 0, "cluster": cluster,
                     "description": re.sub(r"\s+", " ", description).strip(" |")})
    return rows


def _jj_resolve(clusters, footer, images, cache):
    """Renglones de la factura de J.J. Taylor, o ValueError si algo no cierra."""
    if not clusters:
        raise ValueError("no se ven los renglones de producto de la factura de J.J. Taylor.")
    rows = _jj_rows(clusters, images, cache)
    # El Total leído, o el mismo con 5 donde se leyó 6 ($1,429.57 sale $1,429.67
    # en todas las pasadas): vale el que da la suma de los renglones.
    read = [t for t in _ticket_top(footer.get("total", [])) if t is not None]
    totals = set(read)
    for value in read:
        text = f"{value:.2f}"
        totals |= {round(float(text[:i] + "5" + text[i + 1:]), 2) for i, char in enumerate(text) if char == "6"}
    charges = {0.0} | set(_ticket_top(footer.get("charge", [])))
    printed = {round(t - c, 2) for t in totals for c in charges}
    subtotal = _jj_fix_sum(rows, printed)
    total = next(t for t in sorted(totals) if round(t - subtotal, 2) in charges)
    chosen = [dict(r["options"][r["choice"]][1], description=r["description"], cluster=r["cluster"]) for r in rows]
    for c in chosen:
        packs = [_jj_units(cand.get("description")) for cand in c["cluster"] if cand.get("description")]
        per_case = _ticket_most_voted([p[0] for p in packs if p[0]])
        if per_case is None:
            raise ValueError(f"no se pudo leer el pack (piezas por caja) de {c['description']}.")
        sizes = [p[1] for p in packs if p[0] == per_case]
        c["pack"] = (per_case, _ticket_most_voted(sizes) or sizes[0])
    # Confirman las cantidades: la suma de cajas ("Total U without pick ups")
    # o la de piezas ("Piece Count"). Las piezas no son un control exigido:
    # J.J. Taylor no cuenta siempre igual los "2/12/12" (2 o 12 piezas).
    pieces_ok = sum(c["qty"] * c["pack"][0] for c in chosen) in footer.get("pieces", [])
    _ticket_check_qty(chosen, pieces_ok or sum(c["qty"] for c in chosen) in footer.get("cases", []),
                      "J.J. Taylor", "total de cajas (Total U) o de piezas (Piece Count)")
    lines = []
    for c in chosen:
        per_case, size = c["pack"]
        lines.append(_line(
            len(lines) + 1, description=c["description"], qty=c["qty"], pack=per_case, size=size, units=per_case,
            price=c["price"], allowance=c["disc"] or None, net=c["net"], ext=c["ext"],
        ))
    return {"lines": lines, "subtotal": subtotal, "total": total}


def _jj_swap(document):
    """Ticket de cambio ("SWAP": se lleva un producto y se deja otro), con Total $0.00: no es una compra."""
    totals = [t for t in _ticket_top(document["footer"].get("total", [])) if t is not None]
    return totals == [0.0]


def _jj_footer_total(document):
    """
    Total de una factura cuyos renglones no cerraron: vale si dos o más
    pasadas lo leyeron igual, ninguna distinto, y no tiene ningún 6 (en esta
    letra el 5 impreso sale 6: $1,429.57 se lee $1,429.67 en todas). None si no.
    """
    totals = _ticket_votes([{"v": v} for v in document["footer"].get("total", []) if v is not None], "v")
    if len(totals) != 1:
        return None
    total, votes = next(iter(totals.items()))
    return total if votes >= 2 and "6" not in f"{total:.2f}" else None


def _jj_invoices(images, reading, filename, cache=None):
    """
    Las facturas de J.J. Taylor de un PDF (los tickets de cambio no cuentan):
    [{"invoice_no", "date", "amount", "lines", "lines_error"} o {"error",
    "invoice_no"}]. reading(n) -> las filas de la pasada n.
    """
    readings, cache = [], {} if cache is None else cache
    results = {}
    documents = []
    for passes_done in range(len(_TICKET_PASSES)):
        readings.append(reading(passes_done))
        documents = _jj_documents(readings)
        pending = False
        for index, document in enumerate(documents):
            if _jj_swap(document):
                continue
            if _jj_number(document["numbers"], filename) is None or _jj_date(document["dates"], filename) is None:
                pending = True
            if "lines" in results.get(index, {}):
                continue
            try:
                results[index] = _jj_resolve(document["clusters"], document["footer"], images, cache)
            except ValueError as exc:
                results[index] = {"error": str(exc)}
                pending = True
        if not pending:
            break
    invoices = []
    for index, document in enumerate(documents):
        if _jj_swap(document):
            continue
        number = _jj_number(document["numbers"], filename)
        if number is None:
            invoices.append({"error": "no se pudo leer el N° de invoice con seguridad.", "invoice_no": None})
            continue
        date = _jj_date(document["dates"], filename)
        if date is None and number in _filename_numbers(filename, 7):
            # Sin fecha legible: la del nombre del archivo, que ya confirmó el N°.
            date = next(iter(_filename_dates(filename)), None) if len(_filename_dates(filename)) == 1 else None
        if date is None:
            invoices.append({"error": "no se pudo leer la fecha de la factura.", "invoice_no": int(number)})
            continue
        result = results.get(index, {"error": "no se ven los renglones de producto de la factura."})
        if "lines" in result:
            invoices.append({"invoice_no": int(number), "date": date, "amount": result["total"],
                             "lines": result["lines"], "lines_error": None})
            continue
        total = _jj_footer_total(document)
        if total is None:
            invoices.append({"error": f"no se pudo leer el Total con seguridad ({result['error']})",
                             "invoice_no": int(number)})
        else:
            invoices.append({"invoice_no": int(number), "date": date, "amount": total,
                             "lines": None, "lines_error": result["error"]})
    return invoices


def _jj_upright(images, first):
    """
    (imágenes, primera pasada) con cada ticket derecho: una página en la que
    la primera pasada no ve ningún título del ticket vino al revés (la
    detección de orientación de Tesseract no siempre lo nota en estos
    tickets angostos) y se gira 180 grados. Si se giró alguna, first=None
    (la primera pasada se rehace).
    """
    seen = {row["page"] for row in first
            if re.search(r"PRODUCT|Piece\s*Count|Load\s*Sh\w*|PO\s*Num\w*|Pick\s*up", row["text"], re.IGNORECASE)}
    if all(page in seen for page in range(len(images))):
        return images, first
    return [image if page in seen else image.transpose(Image.ROTATE_180) for page, image in enumerate(images)], None


def read_jj_taylor_invoices(pdf_path):
    """Encabezado y renglones de cada factura de J.J. Taylor de un PDF (como read_gce_invoices)."""
    images = _ticket_images(pdf_path)
    if not images:
        raise ValueError("no se encontró la imagen escaneada de la factura de J.J. Taylor.")
    images, first = _jj_upright(images, _ticket_readings(images, 0))

    def reading(n):
        return first if n == 0 and first is not None else _ticket_readings(images, n)

    return _jj_invoices(images, reading, os.path.basename(pdf_path))


# --- Midtown Wholesale (2026-10-05) ---------------------------------------------
# Pedido del usuario (2026-10-05): comparar sus productos contra la factura
# anterior por el NOMBRE (el SKU es propio y a veces de relleno, "123456").
# Hoja completa, escaneo limpio, dos plantillas:
# - SALESGENT (desde 2026): SKU, descripción con la caja ("[30/BX]"), SO Qty,
#   Qty, Sold Price, Amount; pie "Total Quantity" y "Subtotal" (después
#   vienen el descuento "Line Item (D/C)" y el Grand Total).
# - SplitPOS (2025, "Receipt #"): descripción con la caja ("BX/50CT"), Qty,
#   precio y monto; la descripción larga sigue en el renglón de arriba o de
#   abajo; pie "Subtotal: 35 $958.67" (unidades y monto).
# Cuentas: Qty x precio = monto; la suma de montos = Subtotal; la de Qty = las
# unidades del pie. A la izquierda y entre la descripción y las cantidades
# quedan las tildes de control a mano ("~~", "7", "Y"): no son parte del nombre.

_MT_PASSES = (("--psm 6", 1, None), ("--psm 4", 1, None), ("--psm 6", 2, _sharpen), ("--psm 6", 1, _otsu))
_MT_FIELDS = ("price", "ext")
_MT_AMOUNT = re.compile(r"[$S§]\d{1,3}(?:,?\d{3})*[.,]\d{2}")


def _mt_units(description):
    """Unidades por caja según la descripción ("[30/BX]", "BX/50CT", "CS/24CT"); 1 si no dice."""
    text = (description or "").upper()
    match = re.search(r"\[\s*([0-9SOIl]{1,3})\s*/\s*B", text) or re.search(r"\b(?:BX|CS)\s*/\s*([0-9SOIl]{1,3})", text)
    if match:
        digits = _fix_digits(match.group(1))
        if digits.isdigit() and int(digits) > 0:
            return int(digits)
    return 1


def _mt_clean_description(words):
    """Descripción sin las tildes de control a mano del final ni la basura del margen."""
    texts = [w["text"] for w in words]
    while texts and (len(texts[-1]) <= 1 or not re.search(r"[0-9A-Za-z]{2}", texts[-1])):
        texts.pop()
    while texts and not re.search(r"[A-Za-z]", texts[0]) and not re.fullmatch(r"\d{4,6}", texts[0]):
        texts.pop(0)
    return " ".join(texts)


def _mt_product(row, new_template):
    """
    Renglón de producto leído en una pasada, o None. Al final, Sold Price y
    Amount (SALESGENT de ago-2026 trae además una columna Tax en el medio);
    antes, las cantidades: en SplitPOS la tilde a mano suele tapar la Qty
    ("AV)", "V4"), y entonces la cantidad sale de monto / precio.
    """
    # Sin la basura suelta ("-", "|", "~~") y con los montos partidos unidos ("$1 56.75").
    words = []
    for word in row["words"]:
        if not re.search(r"[0-9A-Za-z]", word["text"]):
            continue
        if (words and re.fullmatch(r"[$S§]\d{1,2}", _clean_token(words[-1]["text"]))
                and re.fullmatch(r"\d{2,3}[.,]\d{2}", _clean_token(word["text"]))):
            words[-1] = dict(words[-1], text=words[-1]["text"] + word["text"])
            continue
        words.append(word)
    amounts = []
    for position, word in enumerate(reversed(words)):
        text = _clean_token(word["text"]).rstrip(".,:;")
        if position == 0 and re.fullmatch(r"[$S§]\d{3,6}", text):
            amounts.insert(0, None)  # el monto sin el punto ("$2339"): sale de cantidad x precio
            continue
        if not _MT_AMOUNT.fullmatch(text) or len(amounts) == 3:
            break
        amounts.insert(0, _ticket_amount(text.lstrip("$S§")))
    consumed = len(amounts)
    if len(amounts) == 3:
        if amounts[1] not in (0, None):
            return None  # con impuesto en el renglón: no visto todavía
        amounts = [amounts[0], amounts[2]]
    if len(amounts) != 2:
        return None
    price, ext = amounts
    if price is None:
        return None
    head = words[:len(words) - consumed]
    qty_values = []
    while head and len(qty_values) < (2 if new_template else 1):
        text = _fix_digits(_clean_token(head[-1]["text"]))
        if not re.fullmatch(r"\d{1,3}", text):
            break
        qty_values.insert(0, int(text))
        head = head[:-1]
    sku = None
    if new_template:
        start = next((i for i, w in enumerate(head[:3]) if re.fullmatch(r"\d{4,6}", _clean_token(w["text"]))), None)
        if start is not None:
            sku, head = _clean_token(head[start]["text"]), head[start + 1:]
    description = _mt_clean_description(head)
    if not re.search(r"[A-Za-z]{2}", description):
        return None
    return {"page": row["page"], "y": row["y"], "height": row["height"], "sku": sku,
            "description": description, "qty": qty_values[-1] if qty_values else None, "price": price, "ext": ext}


def _mt_footer(text):
    found = {}
    match = re.search(r"Total\s*Quantity\W*(\d+)", text, re.IGNORECASE)
    if match:
        found["units"] = int(match.group(1))
    match = re.search(r"Subtotal\W*(?:(\d+)\s+)?\$\s*([\d,]+[.,]\d{2})", text, re.IGNORECASE)
    if match:
        if match.group(1):
            found["units"] = int(match.group(1))
        found["subtotal"] = _ticket_amount(match.group(2))
    match = (re.search(r"Grand\s*Total\W*\$\s*([\d,]+[.,]\d{2})", text, re.IGNORECASE)
             or re.search(r"(?<![A-Za-z])Total\s*\$\s*([\d,]+[.,]\d{2})", text))  # SplitPOS: "Total $958.67"
    if match:
        found["grand"] = _ticket_amount(match.group(1))
    match = re.search(r"Line\s*Item\W.{0,6}\$\s*([\d,]+[.,]\d{2})", text, re.IGNORECASE)
    if match:
        found["discount"] = _ticket_amount(match.group(1))
    return found


def _mt_collect(readings):
    """Renglones de producto (agrupados entre pasadas) y pie, con la descripción partida en dos renglones unida."""
    products, footer = [], {}
    for reading in readings:
        new_template = any(re.search(r"SALESGENT|Sold\s*Price|\bSKU\b", row["text"], re.IGNORECASE) for row in reading)
        found = []
        for position, row in enumerate(reading):
            values = _mt_footer(row["text"])
            for key, value in values.items():
                footer.setdefault(key, []).append(value)
            if "subtotal" in values:
                footer.setdefault("subtotal_at", []).append({"page": row["page"], "y": row["y"], "value": values["subtotal"]})
            product = None if values else _mt_product(row, new_template)
            if product is not None:
                found.append((position, product))
        if not new_template:
            # SplitPOS: la descripción larga sigue arriba o abajo del renglón con los montos.
            taken = {position for position, _ in found}
            for position, product in found:
                parts = [(product["y"], product["description"])]
                for near in (position - 1, position + 1):
                    if 0 <= near < len(reading) and near not in taken:
                        row = reading[near]
                        if (row["page"] == product["page"] and abs(row["y"] - product["y"]) < 2.5 * product["height"]
                                and re.search(r"[A-Za-z]{2}", row["text"]) and not _mt_footer(row["text"])
                                and not re.search(r"Receipt|Billed|Notes|Thank|Powered", row["text"], re.IGNORECASE)):
                            parts.append((row["y"], _mt_clean_description(row["words"])))
                product["description"] = " ".join(text for _, text in sorted(parts))
        products.extend(product for _, product in found)
    all_rows = [row for reading in readings for row in reading]
    tolerance = _median([row["height"] for row in all_rows]) if all_rows else 10
    return _ticket_clusters(products, tolerance), footer


def _mt_close(values):
    price, ext = values["price"], values["ext"]
    if price <= 0 or ext <= 0:
        return None
    qty = _whole_qty(price, ext)
    return None if qty is None else dict(values, net=price, qty=qty)


_MT_RULES = {"options": (_MT_FIELDS, _mt_close, (("ext", lambda v, qty: qty * v["price"], True),),
                         _TICKET_CONFUSIONS, 1)}


def _mt_resolve(clusters, footer):
    """Renglones de la factura de Midtown, o ValueError si no cierran contra el Subtotal y las unidades del pie."""
    if not clusters:
        raise ValueError("no se ven los renglones de producto de la factura de Midtown.")
    rows = []
    for line_no, cluster in enumerate(clusters, start=1):
        options = _ticket_options(cluster, *_MT_RULES["options"])
        description = _name_consensus([c.get("description") for c in cluster])
        if not options or not _ticket_settled(options):
            raise ValueError(f"el renglón {line_no} de Midtown ({description}) no cierra: cantidad x precio = monto.")
        rows.append({"options": options, "choice": 0, "cluster": cluster, "description": description})
    # La suma de los montos da el Subtotal o, cuando el descuento ya está en
    # un renglón (SALESGENT, "Line Item (D/C)"), el Grand Total.
    printed = set(_ticket_top(footer.get("subtotal", []))) | set(_ticket_top(footer.get("grand", [])))
    printed |= {round(s - d, 2) for s in _ticket_top(footer.get("subtotal", [])) for d in _ticket_top(footer.get("discount", []))}
    # Un PDF puede traer dos recibos del mismo reparto: la suma de sus Subtotales.
    receipts = [_ticket_winner(_ticket_votes(c, "value")) for c in _ticket_clusters(footer.get("subtotal_at", []), 40)]
    if len(receipts) > 1 and None not in receipts:
        printed.add(round(sum(receipts), 2))
    subtotal = _ticket_fix_sum(rows, printed - {None}, "Subtotal", "Midtown")
    chosen = [dict(r["options"][r["choice"]][1], description=r["description"], cluster=r["cluster"]) for r in rows]
    _ticket_check_qty(chosen, sum(c["qty"] for c in chosen) in footer.get("units", []), "Midtown",
                      "total de unidades (Total Quantity)")
    lines = []
    for c in chosen:
        units = _ticket_most_voted([_mt_units(cand.get("description")) for cand in c["cluster"] if cand.get("description")]) or 1
        lines.append(_line(
            len(lines) + 1, item_no=_ticket_winner(_ticket_votes(c["cluster"], "sku")) or "",
            description=c["description"], qty=c["qty"], pack=units, size="", units=units,
            price=c["price"], net=c["price"], ext=c["ext"],
        ))
    return lines, subtotal


def _mt_lines(images, reading):
    """(renglones, subtotal) con las pasadas de _MT_PASSES de a una hasta que cierran; ValueError si no."""
    readings, error = [], None
    for passes_done in range(len(_MT_PASSES)):
        readings.append(reading(passes_done))
        if passes_done == 0:
            continue  # el nombre sale del acuerdo entre al menos dos lecturas (_name_consensus)
        try:
            return _mt_resolve(*_mt_collect(readings))
        except ValueError as exc:
            error = exc
    raise error


def extract_midtown_lines(pdf_path, invoices=None):
    """
    Renglones de una factura de Midtown Wholesale (LINE_EXTRACTORS). La suma
    tiene que dar el Subtotal impreso (el Grand Total le resta el descuento).
    """
    images = _ticket_images(pdf_path)
    if not images:
        raise ValueError("no se encontró la imagen escaneada de la factura de Midtown.")

    def reading(n):
        config, scale, prep = _MT_PASSES[n]
        return [_ticket_row(page, words) for page in range(len(images))
                for words in _ocr_rows(images[page], config, scale=scale, prep=prep)]

    lines, subtotal = _mt_lines(images, reading)
    wanted = str(invoices[0]["invoice_no"]) if invoices else None
    return {"invoice_no": wanted, "lines": lines, "subtotal": subtotal, "total": subtotal}


# --- Coca-Cola (2026-10-06) ------------------------------------------------------
# Ticket largo de 1 a 3 hojas (a veces con la foto del cheque como primera
# página); la última hoja suele venir dada vuelta (180°). Unas pocas facturas
# de 2024 son PDF con texto (no escaneo). Cada producto ocupa dos renglones,
# agrupados por familia ("ADVANCED 20OZ 1-Ls 24" / "2/48 38.08": cajas,
# unidades y total del grupo):
#   DESCRIPTION MAT# QTY PRICE COND RATE NET EXTENDED
#   20ZPET24LS PA ORG 117687 1 26.34 ZCMA -3.95 19.04 19.04
#   049000032789 24 ZDCS -3.35      (UPC, unidades por caja, otro descuento)
# Cuentas: EXTENDED = QTY x NET (NET = PRICE menos los descuentos); la suma
# de EXTENDED = AMOUNT DUE = TOTAL PRODUCTS (suma de QTY x PRICE) + TOTAL
# ADJUSTMENTS; la suma de QTY = NET PRODUCT QTY y la de QTY x unidades = NET
# CONSUMER QTY (más NET SINGLES QTY, las unidades sueltas). Una devolución
# ("RETURNS") es otro documento con su propio AMOUNT DUE: no se carga (ver
# proveedores._extract_coca_invoices, AMOUNT PAID).

_CC_MONEY = re.compile(r"-?\d{1,3}(?:,?\d{3})*[.,]\d{2}")
_CC_COND = re.compile(r"^[Z2][CDO0Q][MHNDCS0O][A4]?$|^Z?D?CS$|^[Z2]?CMA$", re.IGNORECASE)
_CC_ANCHORS = re.compile(r"\b(SALES|DESCRIPTION|AMOUNT|TOTAL|OUTLET|INVOICE|DELIVERY|RECAP|PRODUCT|QTY|ZCMA|ZDCS|"
                         r"Ls|Pk|PET|CAN|COKE|SPRITE|DASANI)\b", re.IGNORECASE)


def _cc_money(text):
    text = _clean_token(text).strip("$S§")
    text = text.replace(",.", ".").replace(".,", ".")
    return _ticket_amount(text) if _CC_MONEY.fullmatch(text) else None


def _cc_product(row):
    """Renglón DESCRIPTION/MAT#/QTY/PRICE/.../NET/EXTENDED leído en una pasada, o None."""
    words = []
    for w in row["words"]:
        if not re.search(r"[0-9A-Za-z]", w["text"]):
            continue
        glued = re.fullmatch(r"(\d{1,2})-(\d{1,3}[.,]\d{2})", _clean_token(w["text"]))
        if glued:  # "1-25.02": cantidad y precio pegados
            words.extend([dict(w, text=glued.group(1)), dict(w, text=glued.group(2))])
        else:
            words.append(w)
    if len(words) < 5:
        return None
    # Desde la derecha: EXTENDED, NET, [RATE], [COND], PRICE.
    money, index = [], len(words) - 1
    while index >= 0 and len(money) < 4:
        text = _clean_token(words[index]["text"])
        value = _cc_money(text)
        if value is not None:
            money.append((index, value))
        elif not (_CC_COND.match(text) and len(money) >= 2):
            break
        index -= 1
    if len(money) < 3:
        return None
    ext, net = money[0][1], money[1][1]
    rest = money[2:]
    if len(rest) == 2:
        rate, (price_at, price) = rest[0][1], rest[1]
    elif rest[0][1] > 0 and rest[0][1] >= net:
        rate, (price_at, price) = None, rest[0]
    else:
        return None  # el precio no se leyó
    if price is not None and price < net:
        price = None
    if ext is None or net is None or ext <= 0 or net <= 0:
        return None
    if price_at < 2:
        return None
    qty_text = _fix_digits(_clean_token(words[price_at - 1]["text"]))
    mat_at = price_at - 2
    if re.fullmatch(r"\d{5,6}", qty_text):
        # La cantidad quedó pegada al precio ("136174 225,02" = 2 x 25.02): ninguno de los dos se sabe.
        qty, price, mat_at = None, None, price_at - 1
    else:
        qty = int(qty_text) if re.fullmatch(r"\d{1,3}", qty_text) and int(qty_text) > 0 else None
    mat_match = re.search(r"(\d{5,6})$", _fix_digits(_clean_token(words[mat_at]["text"])))
    head = [w["text"] for w in words[:mat_at]]
    if mat_match is None:
        return None
    glued = _clean_token(words[mat_at]["text"])[:-len(mat_match.group(1))]
    if glued:
        head.append(glued)
    description = " ".join(head).strip(" |")
    if not re.search(r"[A-Za-z]{2}", description):
        return None
    return {"page": row["page"], "y": row["y"], "height": row["height"], "upc": None, "upc_box": None,
            "item_no": mat_match.group(1), "qty": qty, "price": price, "rate": rate, "net": net, "ext": ext,
            "description": description}


def _cc_upc_row(row):
    """(UPC válido o None, dígitos leídos, unidades por caja, caja de la palabra) del renglón de abajo, o None."""
    words = [w for w in row["words"] if re.search(r"[0-9A-Za-z]", w["text"])]
    if not words:
        return None
    digits = re.sub(r"\D", "", _fix_digits(_clean_token(words[0]["text"])))
    if not 10 <= len(digits) <= 13:
        return None
    pack = None
    if len(words) > 1:
        text = _fix_digits(_clean_token(words[1]["text"]))
        if re.fullmatch(r"\d{1,3}", text) and int(text) > 0:
            pack = int(text)
    box = (words[0]["x0"], words[0]["top"], words[0]["x1"], words[0]["bottom"])
    return _ticket_upc(digits), _ticket_raw_upc(digits), pack, box


def _cc_close(values):
    net, ext = values["net"], values["ext"]
    if net is None or ext is None or net <= 0 or ext <= 0:
        return None
    qty = _whole_qty(net, ext)
    return None if qty is None else dict(values, qty=qty)


# Dígitos que el OCR confunde en la fuente de Coca-Cola (visto en las
# facturas reales: 20.65 sale 26.65, 38.75 sale 38.78, 75.78 sale 15.78, 20.52
# sale 10.52, 33.48 sale 13.48). Como se prueban muchos cambios, el que vale
# lo decide el total impreso de cada grupo (_cc_by_groups).
_CC_CONFUSIONS = {"0": "689", "6": "058", "8": "6309", "5": "86", "1": "7234", "7": "1", "2": "17", "3": "81",
                  "9": "08", "4": "1"}
_CC_RULES = {
    "options": (("net", "ext"), _cc_close, (("ext", lambda v, qty: qty * v["net"], True),
                                            ("net", lambda v, qty: v["ext"] / qty, True)), _CC_CONFUSIONS, 1),
    "upc": (_CC_CONFUSIONS, "0"),
}


def _cc_footer(text):
    found = {}
    for key, label in (("due", r"AMOUNT\s*DUE"), ("paid", r"AMOUNT\s*PAID"), ("products", r"TOTAL\s*PRODUCTS"),
                       ("adjustments", r"TOTAL\s*ADJUSTMENTS")):
        match = re.search(label + r"[^\w-]*(-?\s?[\d,]+[.,]\d{2})(?!\d)", text, re.IGNORECASE)
        if match:
            found[key] = _ticket_amount(match.group(1).replace(" ", ""))
    for key, label in (("cases", r"NET\s*(?:PRODUCT|BASE)\s*QTY"), ("consumer", r"(?:NET\s*)?CONSUMER\s*QTY"),
                       ("singles", r"NET\s*SINGLES\s*QTY")):
        match = re.search(label + r"\W*(\d+)\b", text, re.IGNORECASE)
        if match:
            found[key] = int(match.group(1))
    return found


_CC_GROUP = re.compile(r"^\W*(\d{1,3})\s*/\s*(\d{1,4})\s+(-?[\d,]+[.,]\d{2})\W*$")
_CC_NUMBER = re.compile(r"INV\w{0,4}\s*[#A-Z]?\s*#?\s*(\d{11})(?!\d)", re.IGNORECASE)
_CC_TOP_NUMBER = re.compile(r"^\W*(\d{11})\W*$")
_CC_SHIP = re.compile(r"SH\w{0,2}\W{0,3}\s*(\d{8})(?!\d)")
_CC_DATE = re.compile(r"DEL\s*DATE\W*(\d{1,2})/(\d{1,2})/(20\d{2})")


def _cc_page_score(text):
    return 3 * len(_CC_ANCHORS.findall(text)) + len(re.findall(r"\d[.,]\d{2}\b", text))


def _cc_upright(images):
    """
    Cada hoja derecha o girada 180°: la orientación automática deja dada
    vuelta la última hoja (la de los totales) en muchos escaneos. Se lee
    rápido de las dos formas y queda la que más texto reconocible da.
    """
    pytesseract = _pytesseract()
    result = []
    for image in images:
        small = image.convert("L")
        if small.width > 1400:
            small = small.resize((1400, int(small.height * 1400 / small.width)))
        straight = _cc_page_score(pytesseract.image_to_string(small, config="--psm 6"))
        flipped = _cc_page_score(pytesseract.image_to_string(small.transpose(Image.ROTATE_180), config="--psm 6"))
        result.append(image.transpose(Image.ROTATE_180) if flipped > straight * 1.5 else image)
    return result


def _cc_text_reading(pdf_path):
    """Los PDF con texto (algunos de 2024): una lectura exacta, renglón por renglón, sin OCR."""
    ensure_pdfplumber()
    rows = []
    with pdfplumber.open(pdf_path) as pdf:
        for page_no, page in enumerate(pdf.pages):
            for line in page.extract_text_lines():
                words = [{"text": t, "x0": 0, "x1": 0, "top": line["top"], "bottom": line["bottom"]}
                         for t in line["text"].split()]
                if words:
                    rows.append(_ticket_row(page_no, words, y=(line["top"] + line["bottom"]) / 2,
                                            height=line["bottom"] - line["top"]))
    # En el texto, la "A" final de ZCMA cae sola en el renglón siguiente.
    return [row for row in rows if row["text"].strip() not in ("A", "#")]


def _cc_is_text_pdf(pdf_path):
    ensure_pdfplumber()
    with pdfplumber.open(pdf_path) as pdf:
        return any("OUTLET" in (page.extract_text() or "") for page in pdf.pages)


def _cc_blocks(readings):
    """
    Facturas del PDF (cada documento termina en su AMOUNT DUE), con todas las
    pasadas hechas: renglones de producto (agrupados entre pasadas), pie,
    N°, fecha y grupos. Una devolución (RETURNS) queda marcada.
    """
    all_rows = [row for reading in readings for row in reading]
    tolerance = _median([row["height"] for row in all_rows]) if all_rows else 10
    # El documento termina en AMOUNT DUE y el AMOUNT PAID de abajo (que va en el
    # mismo grupo): si no, el PAID quedaba en la factura siguiente (50511242049).
    ends = [{"page": row["page"], "y": row["y"]} for row in all_rows
            if {"due", "paid"} & set(_cc_footer(row["text"]))]
    boundaries = [(c[0]["page"], max(e["y"] for e in c) + tolerance)
                  for c in _ticket_clusters(ends, tolerance * 6)]
    if not boundaries:
        raise ValueError("no se encontró el pie (AMOUNT DUE) de la factura de Coca-Cola.")

    def block_of(row):
        return next((i for i, (page, y) in enumerate(boundaries) if (row["page"], row["y"]) <= (page, y)), None)

    blocks = [{"end": end, "products": [], "footer": {}, "numbers": {}, "ships": {}, "dates": {}, "groups": [],
               "returns": 0, "sales": 0} for end in boundaries]
    for reading in readings:
        previous_text = ""
        for position, row in enumerate(reading):
            index = block_of(row)
            if index is None:
                continue
            block = blocks[index]
            text = row["text"]
            product = _cc_product(row)
            if product is not None:
                # El UPC va en el renglón de abajo (o en la hoja siguiente, después de su encabezado).
                for following in reading[position + 1:position + 14]:
                    if _cc_product(following) is not None:
                        break
                    found = _cc_upc_row(following)
                    if found is not None:
                        product["upc"], product["upc_raw"], product["pack"], product["upc_box"] = found
                        break
                block["products"].append(product)
                previous_text = text
                continue
            match = _CC_GROUP.match(text)
            if match:
                block["groups"].append({"page": row["page"], "y": row["y"], "label": previous_text.strip(" |"),
                                        "total": _ticket_amount(match.group(3).replace(",", ""))})
            if re.search(r"RETURNS", text, re.IGNORECASE):
                block["returns"] += 1
            if re.search(r"\bSALES\b", text):
                block["sales"] += 1
            for key, value in _cc_footer(text).items():
                block["footer"].setdefault(key, []).append(value)
            compact = re.sub(r"(?<=\d) (?=\d)", "", text)
            for pattern in (_CC_NUMBER, _CC_TOP_NUMBER):
                match = pattern.search(compact)
                if match:
                    block["numbers"][match.group(1)] = block["numbers"].get(match.group(1), 0) + 1
            match = _CC_SHIP.search(compact)
            if match:
                block["ships"][match.group(1)] = block["ships"].get(match.group(1), 0) + 1
            match = _CC_DATE.search(text)
            if match:
                try:
                    found = datetime(int(match.group(3)), int(match.group(1)), int(match.group(2)))
                    block["dates"][found] = block["dates"].get(found, 0) + 1
                except ValueError:
                    pass
            previous_text = text
    for block in blocks:
        block["clusters"] = _ticket_clusters(block["products"], tolerance)
    return blocks


def _cc_rows(clusters, images, cache):
    """
    Cada grupo de lecturas -> renglón con sus combinaciones que cierran
    (todas: el desempate lo hace el total del grupo, _cc_by_groups) y su UPC.
    """
    rows = []
    for line_no, cluster in enumerate(clusters, start=1):
        options = _ticket_options(cluster, *_CC_RULES["options"])
        upc = _ticket_winner(_ticket_votes(cluster, "upc"))
        if (not _ticket_settled(options) or upc is None) and images:
            cluster = _ticket_reread_row(images, cluster, _cc_product, cache)
            options = _ticket_options(cluster, *_CC_RULES["options"])
            upc = (_ticket_winner(_ticket_votes(cluster, "upc")) or _ticket_reread_upc(images, cluster, cache)
                   or _ticket_upc_variants(cluster, *_CC_RULES["upc"]))
        description = (_ticket_winner(_ticket_votes(cluster, "description"))
                       or next((c["description"] for c in cluster if c.get("description")), ""))
        if not options:
            raise ValueError(f"el renglón {line_no} de Coca-Cola ({description}) no cierra: "
                             "ninguna lectura da cantidad x NET = EXTENDED.")
        if upc is None:
            raise ValueError(f"no se pudo leer el UPC del renglón {line_no} de Coca-Cola ({description}).")
        # Una opción por cada total de renglón distinto (la mejor ubicada).
        distinct, seen = [], set()
        for option in options:
            if option[1]["ext"] not in seen:
                seen.add(option[1]["ext"])
                distinct.append(option)
        rows.append({"options": distinct[:6], "choice": 0, "cluster": cluster, "upc": upc, "line_no": line_no,
                     "description": re.sub(r"\s+", " ", description).strip(" |")})
    return rows


def _cc_combos(members, keep):
    total = 1
    for r in members:
        total *= min(keep, len(r["options"]))
    return total


def _cc_by_groups(rows, groups):
    """
    Elige la lectura de cada renglón con el total impreso de su grupo ("5/75
    111.43"): de las combinaciones que suman exactamente ese total, la de más
    votos (y menos dígitos cambiados); vale solo si no empata con otra. Un
    grupo cuyo total no se leyó, o que ninguna combinación da, queda con la
    lectura más votada de cada renglón (lo controla después la suma contra
    AMOUNT DUE). Un renglón con dos lecturas empatadas que ningún grupo
    decide es un error.
    """
    decided = set()
    for index, group in enumerate(groups):
        start = (group["page"], group["y"])
        end = (groups[index + 1]["page"], groups[index + 1]["y"]) if index + 1 < len(groups) else (99, 0)
        members = [r for r in rows if start < (r["cluster"][0]["page"], r["cluster"][0]["y"]) < end]
        if not members or group["total"] is None:
            continue
        # Para no probar millones de combinaciones: menos lecturas por renglón en los grupos grandes.
        keep = 6
        while keep > 1 and _cc_combos(members, keep) > 200000:
            keep -= 1
        ranked = []
        for combo in itertools.product(*[range(min(keep, len(r["options"]))) for r in members]):
            total = round(sum(r["options"][i][1]["ext"] for r, i in zip(members, combo)), 2)
            if abs(total - group["total"]) < 0.005:
                score = sum(r["options"][i][0] - 50 * r["options"][i][2] for r, i in zip(members, combo))
                ranked.append((score, combo))
        ranked.sort(key=lambda s: -s[0])
        if ranked and (len(ranked) == 1 or ranked[0][0] > ranked[1][0]):
            for r, i in zip(members, ranked[0][1]):
                r["choice"] = i
                decided.add(id(r))
    for r in rows:
        if id(r) not in decided and not _ticket_settled(r["options"]):
            raise ValueError(f"el renglón {r['line_no']} de Coca-Cola ({r['description']}) tiene dos lecturas "
                             "posibles y el total de su grupo no lo decide.")


def _cc_resolve(block, images, cache):
    """Detalle de una factura de Coca-Cola, o ValueError si algo no cierra."""
    footer = block["footer"]
    rows = _cc_rows(block["clusters"], images, cache)
    if not rows:
        raise ValueError("no se ven los renglones de producto de la factura de Coca-Cola.")
    _cc_by_groups(rows, _cc_groups(block))
    printed = list(_ticket_top(footer.get("due", [])))
    for products in _ticket_top(footer.get("products", [])):
        for adjustments in _ticket_top(footer.get("adjustments", [])):
            printed.append(round(products - abs(adjustments), 2))
    groups = [g["total"] for g in _cc_groups(block)]
    if groups:
        printed.append(round(sum(groups), 2))
    total = _ticket_fix_sum(rows, printed, "AMOUNT DUE", "Coca-Cola")
    chosen = [dict(r["options"][r["choice"]][1], upc=r["upc"], description=r["description"], cluster=r["cluster"])
              for r in rows]
    for c in chosen:
        # El número del renglón del UPC son las unidades de consumo del renglón (cajas x unidades por caja).
        c["pack"] = _ticket_most_voted([cand["pack"] // c["qty"] for cand in c["cluster"]
                                        if cand.get("pack") and cand["pack"] % c["qty"] == 0])
        c["price"] = _ticket_most_voted([cand["price"] for cand in c["cluster"] if cand.get("price")])
        c["item_no"] = _ticket_winner(_ticket_votes(c["cluster"], "item_no")) or ""
    confirmed = sum(c["qty"] for c in chosen) in footer.get("cases", [])
    products_read = _ticket_top(footer.get("products", []))
    if all(c["price"] for c in chosen) and round(sum(c["qty"] * c["price"] for c in chosen), 2) in products_read:
        confirmed = True
    _ticket_check_qty(chosen, confirmed, "Coca-Cola", "NET PRODUCT QTY o TOTAL PRODUCTS")
    for c in chosen:
        if not c["pack"]:
            raise ValueError(f"no se pudieron leer las unidades por caja de {c['description']}.")
    consumer = sum(c["qty"] * c["pack"] for c in chosen)
    singles = _ticket_most_voted(footer.get("singles", [])) or 0
    if footer.get("consumer") and consumer + singles not in footer["consumer"] and consumer not in footer["consumer"]:
        raise ValueError(f"las unidades ({consumer}) no coinciden con el NET CONSUMER QTY del pie.")
    lines = []
    for c in chosen:
        category = next((g["label"] for g in reversed(_cc_groups(block))
                         if (g["page"], g["y"]) < (c["cluster"][0]["page"], c["cluster"][0]["y"])), None)
        price = c["price"] if c["price"] and c["price"] >= c["net"] else c["net"]
        lines.append(_line(len(lines) + 1, upc=c["upc"], item_no=c["item_no"], description=c["description"],
                           category=category, qty=c["qty"], pack=c["pack"], size="", units=c["pack"],
                           price=price, allowance=round(price - c["net"], 2) or None, net=c["net"], ext=c["ext"]))
    return {"lines": lines, "total": total}


def _cc_groups(block):
    """Grupos ("2/48 38.08") del bloque, uno por altura (agrupados entre pasadas, el total más votado)."""
    clusters = _ticket_clusters(block["groups"], 15)
    return [{"page": c[0]["page"], "y": c[0]["y"], "label": _ticket_winner(_ticket_votes(c, "label")) or c[0]["label"],
             "total": _ticket_winner(_ticket_votes(c, "total"))} for c in clusters
            if _ticket_winner(_ticket_votes(c, "total")) is not None]


def _cc_number(block, filename):
    """N° de 11 dígitos: el más leído si lo confirma el SHP# (sus primeros 8), el nombre del archivo o 2+ pasadas."""
    votes = block["numbers"]
    from_name = set(re.findall(r"(?<!\d)\d{11}(?!\d)", filename or ""))
    ships = {s for s, n in block["ships"].items()}
    ranked = sorted(votes, key=lambda n: -votes[n])
    for number in ranked:
        if number in from_name or number[:8] in ships:
            return number
    winner = _ticket_winner(votes)
    return winner if winner is not None and votes[winner] >= 2 else None


def _cc_date(block, filename):
    votes = block["dates"]
    named = _filename_dates(filename)
    for read in sorted(votes, key=lambda d: -votes[d]):
        if read in named:
            return read
    if len(named) == 1:
        return next(iter(named))  # el nombre del archivo manda (ver _extract_coca_invoices)
    # "01.02.2024" da dos fechas posibles: los nombres de Coca-Cola van
    # día.mes.año (39703471016 01.02.2024 es del 1 de febrero).
    match = re.search(r"(?<!\d)(\d{1,2})\.(\d{1,2})\.(\d{4})(?!\d)", filename or "")
    if match:
        try:
            return datetime(int(match.group(3)), int(match.group(2)), int(match.group(1)))
        except ValueError:
            pass
    winner = _ticket_winner(votes)
    return winner if winner is not None and votes[winner] >= 2 else None


def _cc_solid(values):
    """Un número del pie leído sin dudas: en 3+ pasadas y el doble que cualquier otra lectura; o None."""
    counts = _ticket_votes([{"v": v} for v in values], "v")
    winner = _ticket_winner(counts)
    if winner is None or counts[winner] < 3:
        return None
    if any(2 * n > counts[winner] for value, n in counts.items() if value != winner):
        return None
    return winner


def _cc_amount(block, detail, returns=(), return_readings=()):
    """
    (importe, sale del PAID): AMOUNT DUE confirmado por otra parte del
    ticket y la regla de AMOUNT PAID: si lo pagado es otro importe (una
    devolución del mismo día ya neteada), vale lo pagado. (None, False) si
    no se puede confirmar.

    Confirma el DUE: la suma de los renglones; TOTAL PRODUCTS menos TOTAL
    ADJUSTMENTS (el OCR suele perder el signo menos), admitiendo un dígito
    confundido en uno de los dos (50709193040: 742.29 sale 142.29); la suma
    de los totales de grupo; o, con una devolución del PDF (`returns`), DUE
    más la devolución igual al PAID. Leer el DUE igual en todas las pasadas
    NO alcanza: la fuente grande confunde siempre igual el 5 con el 6
    (49119495039: $517.39 sale $617.39 en el DUE y en el PAID).
    """
    footer = block["footer"]
    due = detail["total"] if detail is not None else None
    solid_due = _cc_solid(footer.get("due", []))
    if due is not None and solid_due is not None and abs(solid_due - due) > 0.005:
        # Los renglones cierran contra los totales de grupo, pero el AMOUNT DUE
        # impreso dice otra cosa: no se sabe cuál vale (revisión 2026-10-08).
        return None, False
    paid_votes = _ticket_votes([{"v": v} for v in footer.get("paid", [])], "v")
    paid = _ticket_winner(paid_votes)
    if due is None:
        sums = set()
        products = _ticket_top(footer.get("products", []))
        adjustments = _ticket_top(footer.get("adjustments", []))
        if solid_due is not None:
            # Con el DUE leído sin dudas, vale cualquier lectura de TOTAL
            # PRODUCTS / ADJUSTMENTS que lo confirme: el pie chico se lee peor
            # (44371558036: ADJ 122.34 una vez de tres; 50709193040: 142.29
            # dos de seis, que con el 1 por 7 da 742.29 - 139.64 = 602.65).
            products = sorted(set(footer.get("products", [])))
            adjustments = sorted(set(footer.get("adjustments", [])))
        for p in products:
            for a in adjustments:
                a = abs(a)
                sums.add(round(p - a, 2))
                sums.update(round(v - a, 2) for v in _ticket_variants(p, _CC_CONFUSIONS))
                sums.update(round(p - abs(v), 2) for v in _ticket_variants(a, _CC_CONFUSIONS))
        groups = [g["total"] for g in _cc_groups(block)]
        if groups:
            sums.add(round(sum(groups), 2))
        candidates = [d for d in _ticket_top(footer.get("due", [])) if d in sums]
        due = candidates[0] if len(candidates) == 1 else None
    solid_paid = _cc_solid(footer.get("paid", []))
    if due is None and returns:
        if solid_due is not None and solid_paid is not None and abs(solid_due + sum(returns) - solid_paid) < 0.005:
            due = solid_due
    if (due is None and return_readings and solid_due is not None and solid_paid is not None
            and 0 < solid_paid < solid_due and round(solid_due - solid_paid, 2) in return_readings):
        # Devolución mal leída en su hoja, pero DUE - PAID es una de sus
        # lecturas: las tres cifras se confirman entre sí y la devolución ya
        # viene neteada en lo pagado (39399414007: 654.84 - 51.30 = 603.54).
        return solid_paid, True
    if due is None:
        return None, False
    # Lo pagado (una devolución ya neteada) vale solo leído sin dudas: con dos
    # lecturas de seis, un PAID mal leído pisaba un DUE confirmado (revisión
    # 2026-10-08).
    if solid_paid is not None and 0 < solid_paid < due:
        return solid_paid, True
    if paid is not None and 0 < paid < due:
        return None, False
    return due, False


def _cc_candidates(block, returns=()):
    """Importes posibles de una factura que no se pudo confirmar (para cruzarlos con Chase)."""
    footer = block["footer"]
    found = set(_ticket_top(footer.get("due", []))) | set(_ticket_top(footer.get("paid", [])))
    found |= {round(d + sum(returns), 2) for d in list(found)} if returns else set()
    return sorted(v for v in found if v > 0)


def _cc_invoices(images, reading, filename, passes, cache=None):
    readings, resolved, failures = [], {}, {}
    cache = {} if cache is None else cache
    blocks = []
    for passes_done in range(passes):
        readings.append(reading(passes_done))
        if passes_done == 0 and passes > 1:
            continue
        try:
            blocks = _cc_blocks(readings)
        except ValueError:
            continue
        pending = False
        for block in blocks:
            key = (block["end"][0], round(block["end"][1] / 60))
            if block["returns"] and not block["sales"] or key in resolved:
                continue
            try:
                resolved[key] = _cc_resolve(block, images, cache)
            except ValueError as exc:
                failures[key] = str(exc)
                pending = True
        if not pending:
            break
    if not blocks:
        raise ValueError("no se encontró el pie (AMOUNT DUE) de ninguna factura de Coca-Cola.")
    # Devoluciones (RETURNS) del PDF: si la factura no la trae ya neteada en
    # el AMOUNT PAID, se le resta (ver _cc_amount). Solo con una sola factura
    # con importe en el PDF; si no, no se sabe a cuál va.
    returns, unreadable_returns, return_readings = [], 0, set()
    for block in blocks:
        if block["returns"] and not block["sales"]:
            value = _cc_solid([abs(v) for v in block["footer"].get("due", [])])
            if value is None:
                unreadable_returns += 1
                return_readings = {round(abs(v), 2) for v in block["footer"].get("due", []) if v}
            elif value:
                returns.append(-value)
    sales = [b for b in blocks if not (b["returns"] and not b["sales"])
             and any(v > 0 for v in b["footer"].get("due", []) + b["footer"].get("paid", []))]
    results = []
    for block in blocks:
        if block["returns"] and not block["sales"]:
            continue  # devolución: se resta en la factura principal (o ya viene neteada en AMOUNT PAID)
        key = (block["end"][0], round(block["end"][1] / 60))
        detail = resolved.get(key)
        number = _cc_number(block, filename)
        if number is None:
            if block["products"] or block["numbers"]:
                results.append({"error": "no se pudo leer el N° de invoice con seguridad.", "invoice_no": None})
            continue
        date = _cc_date(block, filename)
        if date is None:
            results.append({"error": "no se pudo leer la fecha de la factura.", "invoice_no": int(number)})
            continue
        block_returns = returns if len(sales) == 1 else []
        single_unreadable = len(sales) == 1 and unreadable_returns == 1 and not returns
        amount, from_paid = _cc_amount(block, detail, block_returns, return_readings if single_unreadable else ())
        if amount and not from_paid and (returns or unreadable_returns):
            # La devolución no vino neteada en el AMOUNT PAID: se resta (pedido
            # del usuario, 2026-10-08: así está en el Excel de Proveedores,
            # 43485344087 $1,068.62 − $15.22 = $1,053.40).
            if unreadable_returns or len(sales) != 1:
                results.append({"error": "trae una devolución en el mismo PDF que no se pudo descontar con "
                                         "seguridad.", "invoice_no": int(number), "date": date,
                                "candidates": _cc_candidates(block, block_returns)})
                continue
            amount = round(amount + sum(returns), 2)
        if amount is None:
            results.append({"error": "no se pudo leer el AMOUNT DUE con seguridad.", "invoice_no": int(number),
                            "date": date, "candidates": _cc_candidates(block, block_returns),
                            "lines_error": failures.get(key)})
            continue
        if amount == 0:
            continue  # factura en $0.00: no deja deuda
        if amount < 0:
            results.append({"error": "dio un importe negativo.", "invoice_no": int(number)})
            continue
        results.append({"invoice_no": int(number), "date": date, "amount": amount,
                        "lines": detail["lines"] if detail is not None else None,
                        "lines_error": None if detail is not None else failures.get(key)})
    return results


def read_coca_invoices(pdf_path):
    filename = os.path.basename(pdf_path)
    if _cc_is_text_pdf(pdf_path):
        reading = _cc_text_reading(pdf_path)
        return _cc_invoices([], lambda n: reading, filename, 1)
    images = _ticket_images(pdf_path)
    if not images:
        raise ValueError("no se encontró la imagen escaneada de la factura de Coca-Cola.")
    images = _cc_upright(images)
    return _cc_invoices(images, lambda n: _ticket_readings(images, n), filename, len(_TICKET_PASSES))


# --- Sweetheart Ice Cream (2026-10-06) -------------------------------------------
# Ticket de una página (a veces fotografiado junto al cheque). Cada producto
# ocupa dos renglones:
#   Product# Descripción U/C Case/Unit Units Price ExtPrice
#   1002     GH Giant King Cone 12 1/ 0 12 2.4992 29.99
#            0-77567-00822-0  (el UPC, con guiones)
# Cuentas: Units x Price (4 decimales) = ExtPrice redondeado al centavo; la
# suma de ExtPrice = TOTAL SALES = BALANCE DUE; la suma de Units = las
# unidades del renglón TOTAL. Los productos van agrupados por categoría
# ("SUBTOTAL: Impulse", "Pints"); el recargo de combustible (9100 "Fuel
# Surcharge", UPC 0-00000-00000-0, "SUBTOTAL: Charges") suma al total pero no
# es un producto.

_SW_PRICE = re.compile(r"[\dSOIlB]{1,3}[.,]\d{4}")
_SW_EXT = re.compile(r"\d{1,4}[.,]\d{2}")
_SW_UPC = re.compile(r"(\d)\s*-\s*(\d{5})\s*-\s*(\d{5})\s*-\s*(\d)")


def _sw_product(row):
    """Renglón Product#/.../Units/Price/ExtPrice leído en una pasada, o None."""
    words = [w for w in row["words"] if re.search(r"[0-9A-Za-z]", w["text"])]
    # A la izquierda puede haber basura del borde de la foto ("; O ; 1003 GH ...").
    start = next((i for i, w in enumerate(words[:6]) if re.fullmatch(r"\d{4}\.?", _clean_token(w["text"]))
                  and i + 1 < len(words) and re.search(r"[A-Za-z]", words[i + 1]["text"])), None)
    if start is None:
        return None
    words = words[start:]
    if len(words) < 4:
        return None
    # Precio: el único número con 4 decimales (con basura pegada: "=��1.1996").
    price_at = price = None
    for i, w in enumerate(words):
        if i < 2:
            continue
        match = re.search(r"(\d{1,3})[.,](\d{4})(?!\d)", _fix_digits(w["text"]))
        if match:
            price_at, price = i, round(float(f"{match.group(1)}.{match.group(2)}"), 4)
            break
    if price_at is None or price_at + 1 >= len(words):
        return None
    match = re.search(r"(\d{1,4})[.,](\d{2})(?!\d)", words[price_at + 1]["text"])
    ext = float(f"{match.group(1)}.{match.group(2)}") if match else None
    units = _fix_digits(_clean_token(words[price_at - 1]["text"]))
    description, desc_words = [], []
    for word in words[1:price_at - 1]:
        if re.fullmatch(r"\d{1,3}", _clean_token(word["text"])) and description:
            break
        description.append(word["text"])
        desc_words.append(word)
    return {
        "page": row["page"], "y": row["y"], "height": row["height"], "upc": None, "upc_box": None,
        "item_no": _clean_token(words[0]["text"]).rstrip("."),
        "price": price, "ext": ext,
        "qty": int(units) if units.isdigit() and 0 < int(units) < 1000 else None,
        "uc": next((int(t) for t in (_clean_token(w["text"]) for w in words[2:price_at - 1])
                    if t.isdigit() and 0 < int(t) <= 48), None),
        "description": re.sub(r"^[\W_]+|[\W_]+$", "", " ".join(description)),
        "desc_box": (desc_words[0]["x0"], desc_words[-1]["x1"]) if desc_words else None,
    }


def _sw_close(values):
    price, ext = values["price"], values["ext"]
    if price is None or ext is None or price <= 0 or ext <= 0:
        return None
    qty = _whole_qty(price, ext)
    return None if qty is None else dict(values, net=price, qty=qty)


# Sin confusiones de dígitos: el precio tiene 4 decimales y la cuenta del
# renglón ya es muy exigente; si una pasada no lo lee bien, lo lee otra.
_SW_RULES = {
    "options": (("price", "ext"), _sw_close, (("ext", lambda v, qty: round(qty * v["price"], 2), True),), {}, 0),
    "upc": (_TICKET_CONFUSIONS, ""),
}


def _sw_footer(text):
    found = {}
    match = re.search(r"TOTAL\s*SALES\W*[$S§]?\s*([\d,]+[.,]\d{2})(?!\d)", text, re.IGNORECASE)
    if match:
        found["total_sales"] = _ticket_amount(match.group(1))
    match = re.search(r"BALANCE\s*DUE\W*[$S§]?\s*([\d,]+[.,]\d{2})(?!\d)", text, re.IGNORECASE)
    if match:
        found["balance_due"] = _ticket_amount(match.group(1))
    match = re.search(r"^\W*TOTAL\W+(\d{1,3})\s*[/I|l1]\s*(\d{1,3})\s+(\d{1,4})\b", text, re.IGNORECASE)
    if match:
        found["units"] = int(match.group(3))
    match = re.search(r"SUBTOTAL\W*([A-Za-z]+)", text, re.IGNORECASE)
    if match:
        found["category"] = match.group(1).capitalize()
    return found


_SW_NUMBER = re.compile(r"INVO\w{0,4}\W{0,3}\s*(1301\d{4,7})(?!\d)", re.IGNORECASE)
_SW_DATE = re.compile(r"Date\W{0,3}(\d{1,2})/(\d{1,2})/(20\d{2})(?!\d)")


def _sw_collect(readings):
    products, footer, numbers, dates, subtotals, footer_rows = [], {}, {}, {}, [], []
    for reading in readings:
        for position, row in enumerate(reading):
            product = _sw_product(row)
            if product is not None:
                for following in reading[position + 1:position + 3]:
                    if _sw_product(following) is not None:
                        break
                    match = _SW_UPC.search(_fix_digits(following["text"]))
                    if match:
                        digits = "".join(match.groups())
                        product["upc"], product["upc_raw"] = _ticket_upc(digits), _ticket_raw_upc(digits)
                        cells = [w for w in following["words"] if re.search(r"\d-\d|\d{5}", w["text"])]
                        if cells:
                            product["upc_box"] = (min(w["x0"] for w in cells), min(w["top"] for w in cells),
                                                  max(w["x1"] for w in cells), max(w["bottom"] for w in cells))
                        if digits == "000000000000":
                            product["upc"] = digits
                        break
                products.append(product)
                continue
            values = _sw_footer(row["text"])
            if "total_sales" in values or "balance_due" in values:
                footer_rows.append((row, "total_sales" if "total_sales" in values else "balance_due"))
            if "category" in values:
                subtotals.append({"page": row["page"], "y": row["y"], "category": values.pop("category")})
            for key, value in values.items():
                footer.setdefault(key, []).append(value)
            match = _SW_NUMBER.search(re.sub(r"(?<=\d) (?=\d)", "", row["text"]))
            if match:
                numbers[match.group(1)] = numbers.get(match.group(1), 0) + 1
            match = _SW_DATE.search(row["text"])
            if match:
                try:
                    found = datetime(int(match.group(3)), int(match.group(1)), int(match.group(2)))
                    dates[found] = dates.get(found, 0) + 1
                except ValueError:
                    pass
    all_rows = [row for reading in readings for row in reading]
    tolerance = _median([row["height"] for row in all_rows]) if all_rows else 10
    return {"clusters": _ticket_clusters(products, tolerance), "footer": footer, "numbers": numbers,
            "dates": dates, "subtotals": subtotals, "footer_rows": footer_rows}


def _sw_cell_votes(image, box, whitelist):
    """Lecturas de una celda sola, agrandada (x2, x3, x4) y con tres arreglos de la imagen, psm 7 y 8."""
    pytesseract = _pytesseract()
    box = (max(0, box[0]), max(0, box[1]), min(image.width, box[2]), min(image.height, box[3]))
    if box[2] - box[0] < 4 or box[3] - box[1] < 4:
        return []
    crop = image.crop(tuple(int(v) for v in box)).convert("L")
    texts = []
    for scale in (2, 3, 4):
        big = crop.resize((crop.width * scale, crop.height * scale), Image.LANCZOS)
        for prep in (None, _otsu, _sharpen):
            for psm in (7, 8):
                texts.append(pytesseract.image_to_string(
                    prep(big) if prep else big, config=f"--psm {psm} -c tessedit_char_whitelist={whitelist}").strip())
    return texts


def _sw_upc_cell(images, cluster, cache):
    """
    UPC de un renglón que ninguna pasada leyó con dígito verificador válido
    (en esta fuente el 8 sale 6: 0-77567-02874-7 se lee 02674): se relee la
    celda sola, donde la vio alguna pasada o, si no, debajo de la
    descripción. Vale el UPC válido más leído, con al menos 3 lecturas y el
    doble que cualquier otro válido.
    """
    page = cluster[0]["page"]
    y = _median([c["y"] for c in cluster])
    height = _median([c["height"] for c in cluster])
    key = ("upc-cell", page, round(y / 10))
    if key not in cache:
        boxes = [c["upc_box"] for c in cluster if c.get("upc_box")]
        if boxes:
            x0, top, x1, bottom = boxes[0]
            pad = (bottom - top) * 0.6
            box = (x0 - pad, top - pad, x1 + pad, bottom + pad)
        else:
            spans = [c["desc_box"] for c in cluster if c.get("desc_box")]
            if not spans:
                cache[key] = None
                return None
            x0, x1 = spans[0]
            box = (x0 - height, y + 0.5 * height, max(x1, x0 + 9 * height) + height, y + 2.0 * height)
        votes, printed = {}, {}
        for text in _sw_cell_votes(images[page], box, "0123456789-"):
            match = _SW_UPC.search(text)
            digits = "".join(match.groups()) if match else re.sub(r"\D", "", text)
            upc = _ticket_upc(digits)
            if upc is not None:
                votes[upc] = votes.get(upc, 0) + 1
            if match:
                printed[digits] = printed.get(digits, 0) + 1

        def clear(counts, minimum):
            ranked = sorted(counts.items(), key=lambda kv: -kv[1])
            return (ranked[0][0] if ranked and ranked[0][1] >= minimum and
                    (len(ranked) == 1 or ranked[0][1] >= 2 * ranked[1][1]) else None)

        # Algunos códigos vienen impresos con el dígito verificador mal (Talenti:
        # 1-86852-00109-0): valen tal cual si la celda los lee igual con mucha mayoría.
        cache[key] = clear(votes, 3) or clear(printed, 6)
    return cache[key]


def _sw_amount_cell(images, row, cache):
    """
    Relee el importe de TOTAL SALES / BALANCE DUE en su celda sola, agrandada
    y solo con dígitos, de varias formas (votos). En el ticket de Sweetheart
    el 8 se lee 6 en el renglón entero, en todas las pasadas ($236.15 por
    $238.15, las dos líneas del pie); la celda sola lo lee bien casi siempre.
    """
    word = row["words"][-1]
    key = ("cell", row["page"], round(word["top"] / 10))
    if key not in cache:
        pytesseract = _pytesseract()
        height = word["bottom"] - word["top"]
        box = (word["x0"] - height, word["top"] - height / 2, word["x1"] + height, word["bottom"] + height / 2)
        image = images[row["page"]]
        box = (max(0, box[0]), max(0, box[1]), min(image.width, box[2]), min(image.height, box[3]))
        if box[2] - box[0] < 4 or box[3] - box[1] < 4:
            cache[key] = []
            return cache[key]
        crop = image.crop(tuple(int(v) for v in box)).convert("L")
        votes = []
        for scale in (2, 3, 4):
            big = crop.resize((crop.width * scale, crop.height * scale), Image.LANCZOS)
            for prep in (None, _otsu, _sharpen):
                for psm in (7, 8):
                    text = pytesseract.image_to_string(prep(big) if prep else big,
                                                       config=f"--psm {psm} -c tessedit_char_whitelist=0123456789.,$")
                    match = re.search(r"(\d[\d,]*[.,]\d{2})\s*$", text.strip())
                    if match:
                        votes.append(_ticket_amount(match.group(1)))
        cache[key] = [v for v in votes if v is not None]
    return cache[key]


def _sw_rows(clusters, images, cache):
    """
    Como _ticket_rows, pero también se relee la franja del renglón cuando lo
    vio una sola pasada o cuando su total salió de la cuenta (no leído): un
    precio mal leído (1.4990 por 1.4900) da otro total que igual "cierra".
    """
    rows = []
    for line_no, cluster in enumerate(clusters, start=1):
        options = _ticket_options(cluster, *_SW_RULES["options"])
        upc = _ticket_winner(_ticket_votes(cluster, "upc"))
        if not _ticket_settled(options) or upc is None or options[0][2] == 2 or len(cluster) < 2:
            # Las relecturas de la franja traen posiciones de la franja, no de la página.
            first, cluster = cluster, _ticket_reread_row(images, cluster, _sw_product, cache)
            options = _ticket_options(cluster, *_SW_RULES["options"])
            upc = _ticket_winner(_ticket_votes(cluster, "upc")) or _sw_upc_cell(images, first, cache)
        description = (_ticket_winner(_ticket_votes(cluster, "description"))
                       or next((c["description"] for c in cluster if c.get("description")), ""))
        if not options:
            raise ValueError(f"el renglón {line_no} de Sweetheart ({description}) no cierra: "
                             "ninguna lectura da unidades x precio = total.")
        if not _ticket_settled(options):
            raise ValueError(f"el renglón {line_no} de Sweetheart ({description}) tiene dos lecturas "
                             "posibles que cierran con los mismos votos.")
        if upc is None:
            raise ValueError(f"no se pudo leer el UPC del renglón {line_no} de Sweetheart ({description}).")
        rows.append({"options": options, "choice": 0, "cluster": cluster, "upc": upc,
                     "description": re.sub(r"\s+", " ", description).strip(" |")})
    return rows


def _sw_is_charge(c):
    return c["upc"] == "000000000000" or re.search(r"surcharge|fuel", c["description"], re.IGNORECASE)


def _sw_resolve(found, images, cache):
    """Detalle de la factura de Sweetheart, o ValueError si algo no cierra."""
    footer = found["footer"]
    rows = _sw_rows(found["clusters"], images, cache)
    if not rows:
        raise ValueError("no se ven los renglones de producto de la factura de Sweetheart.")
    printed = _ticket_top(footer.get("total_sales", []) + footer.get("balance_due", []))
    total = _ticket_fix_sum(rows, printed, "TOTAL SALES/BALANCE DUE", "Sweetheart")
    chosen = _ticket_chosen(rows)
    for c in chosen:
        c["uc"] = _ticket_most_voted([cand["uc"] for cand in c["cluster"] if cand.get("uc")])
    units = sum(c["qty"] for c in chosen)
    confirmed = (units in footer.get("units", []) or units == sum(footer.get("subtotal_units") or [-1])
                 or all(c["uc"] and c["qty"] % c["uc"] == 0 for c in chosen if not _sw_is_charge(c)))
    _ticket_check_qty(chosen, confirmed, "Sweetheart", "total de unidades (TOTAL)")
    lines = []
    for c in chosen:
        if _sw_is_charge(c):
            continue
        uc = c["uc"]
        below = [s for s in found["subtotals"] if s["page"] == c["cluster"][0]["page"] and s["y"] > c["cluster"][0]["y"]]
        category = min(below, key=lambda s: s["y"])["category"] if below else None
        item = _ticket_winner(_ticket_votes(c["cluster"], "item_no")) or ""
        if uc and c["qty"] % uc == 0:
            qty, pack, net = c["qty"] // uc, uc, round(c["price"] * uc, 4)
        else:
            qty, pack, net = c["qty"], 1, c["price"]
        lines.append(_line(len(lines) + 1, upc=c["upc"], item_no=item, description=c["description"],
                           category=category, qty=qty, pack=pack, size="", units=pack, price=net, net=net,
                           ext=c["ext"]))
    if not lines:
        raise ValueError("la factura de Sweetheart no tiene renglones de producto.")
    subtotal = round(sum(line["ext"] for line in lines), 2)
    return {"lines": lines, "subtotal": subtotal, "total": total}


def _sw_number(found, filename):
    """N° más leído (al menos dos pasadas, sin empate); si no, el del nombre del archivo si alguna pasada lo leyó."""
    votes = found["numbers"]
    winner = _ticket_winner(votes)
    if winner is not None and votes[winner] >= 2:
        return winner
    named = {n for n in re.findall(r"1301\d{4,7}", filename or "") if n in votes}
    return named.pop() if len(named) == 1 else None


def _sw_date(found, filename):
    votes = found["dates"]
    named = [d for d in _filename_dates(filename) if d in votes]
    if len(named) == 1:
        return named[0]
    winner = _ticket_winner(votes)
    return winner if winner is not None and votes[winner] >= 2 else None


def _sw_footer_total(footer):
    """
    Importe sin renglones: TOTAL SALES y BALANCE DUE son el mismo importe
    impreso en dos lugares, así que se votan juntos (en los escaneos viejos el
    8 sale 6 seguido: $268.04 / $286.04 / $288.04); vale el más leído si lo
    leyeron los dos lugares.
    """
    sales, due = footer.get("total_sales", []), footer.get("balance_due", [])
    counts = _ticket_votes([{"v": v} for v in sales + due], "v")
    ranked = sorted(counts.values(), reverse=True)
    winner = _ticket_most_voted(sales + due)
    if winner is None:
        return None
    # Leído en los dos lugares, o con el doble de votos que la segunda lectura.
    return winner if (winner in sales and winner in due) or (len(ranked) > 1 and ranked[0] >= 2 * ranked[1]) else None


def _sw_invoice(images, reading, filename, cache=None, pages=None):
    readings, cache = [], {} if cache is None else cache
    detail = error = None
    found = None
    for passes_done in range(len(_TICKET_PASSES)):
        readings.append(reading(passes_done))
        if passes_done == 0:
            continue  # N°, fecha y montos: con al menos dos lecturas que coincidan
        found = _sw_collect(readings)
        for row, key in found["footer_rows"][:4]:
            found["footer"].setdefault(key, []).extend(_sw_amount_cell(images, row, cache))
        try:
            detail, error = _sw_resolve(found, images, cache), None
        except ValueError as exc:
            error = str(exc)
            continue
        if _sw_number(found, filename) and _sw_date(found, filename):
            break
    number = _sw_number(found, filename)
    if number is None:
        return {"error": "no se pudo leer el N° de invoice con seguridad.", "invoice_no": None}
    amount = detail["total"] if detail is not None else _sw_footer_total(found["footer"])
    if amount is None:
        return {"error": "no se pudo leer el TOTAL SALES con seguridad.", "invoice_no": int(number)}
    date = _sw_date(found, filename) or _sw_page_date(pages or images, filename)
    if date is None:
        return {"error": "no se pudo leer la fecha de la factura.", "invoice_no": int(number)}
    return {"invoice_no": int(number), "date": date, "amount": amount,
            "lines": detail["lines"] if detail is not None else None, "lines_error": error}


def _sw_page_date(images, filename):
    """
    Respaldo de la fecha (2026-10-07): en algunos tickets ninguna lectura por
    renglones ve "Date: MM/DD/AAAA" y un OCR común de la página entera (sin
    recortar el ticket) sí: así la leía el lector viejo. Vale solo si es la misma fecha del nombre del archivo.
    """
    named = _filename_dates(filename)
    pytesseract = _pytesseract()
    seen = set()
    for image in images:
        text = pytesseract.image_to_string(image)
        for match in re.finditer(r"Date\W{0,3}(\d{1,2})/(\d{1,2})/(20\d{2})(?!\d)", text):
            try:
                seen.add(datetime(int(match.group(3)), int(match.group(1)), int(match.group(2))))
            except ValueError:
                pass
    matches = seen & named
    return next(iter(matches)) if len(matches) == 1 else None


def _sw_ticket_only(image):
    """
    Recorta el ticket: muchas veces la foto trae el cheque al lado y el OCR
    mezcla sus trazos con los montos del renglón. Los bordes salen del
    renglón de títulos (Product# ... ExtPrice), con margen; sin títulos
    legibles, la imagen queda igual.
    """
    for words in _ocr_rows(image, "--psm 6"):
        texts = [w["text"].upper() for w in words]
        product = next((w for w, t in zip(words, texts) if t.startswith("PRODUCT")), None)
        price = [w for w, t in zip(words, texts) if t.endswith("PRICE") or t.endswith("PRIGE")]
        if product is not None and price:
            right = max(w["x1"] for w in price)
            width = right - product["x0"]
            if width < image.width * 0.3:
                continue
            return image.crop((max(0, int(product["x0"] - width * 0.05)), 0,
                               min(image.width, int(right + width * 0.05)), image.height))
    return image


def read_sweetheart_invoices(pdf_path):
    images = _ticket_images(pdf_path)
    if not images:
        raise ValueError("no se encontró la imagen escaneada de la factura de Sweetheart.")
    pages = images
    images = [_sw_ticket_only(image) for image in images]
    return [_sw_invoice(images, lambda n: _ticket_readings(images, n), os.path.basename(pdf_path), pages=pages)]


def extract_coca_lines(pdf_path, invoices=None):
    """Renglones de una factura de Coca-Cola (LINE_EXTRACTORS); ver extract_gce_lines."""
    return _ticket_lines(read_coca_invoices(pdf_path), invoices, "Coca-Cola")


def extract_sweetheart_lines(pdf_path, invoices=None):
    """Renglones de una factura de Sweetheart (LINE_EXTRACTORS); ver extract_gce_lines."""
    return _ticket_lines(read_sweetheart_invoices(pdf_path), invoices, "Sweetheart")


def _ticket_lines(found, invoices, label):
    found = [inv for inv in found if "error" not in inv]
    wanted = {str(inv["invoice_no"]) for inv in (invoices or [])}
    for invoice in found:
        if (not wanted or str(invoice["invoice_no"]) in wanted) and invoice["lines"] is not None:
            return {"invoice_no": str(invoice["invoice_no"]), "lines": invoice["lines"],
                    "subtotal": None, "total": invoice["amount"]}
    error = next((inv["lines_error"] for inv in found if inv.get("lines_error")), None)
    raise ValueError(error or f"no se encontró en el PDF la factura de {label} pedida.")


# supplier_key (el de SUPPLIER_REGISTRY en proveedores.py) -> extractor de renglones.
LINE_EXTRACTORS = {
    "ht_hackney": extract_ht_hackney_lines,
    "cec": extract_cec_lines,
    "colonial": extract_colonial_lines,
    "gce": extract_gce_lines,
    "red_bull": extract_red_bull_lines,
    "frito_lay": extract_frito_lay_lines,
    "midtown": extract_midtown_lines,
    "coca": extract_coca_lines,
    "sweetheart": extract_sweetheart_lines,
}


def extract_lines(supplier_key, pdf_path, invoices=None):
    """
    Detalle de productos de la factura, o None si el proveedor todavía no
    tiene extractor. `invoices`: los encabezados que la app ya leyó de ese
    PDF ({"invoice_no", "date", "amount"}) -- Colonial controla sus renglones
    contra ese importe, porque el total impreso de su factura no se lee por OCR.
    """
    extractor = LINE_EXTRACTORS.get(supplier_key)
    return extractor(pdf_path, invoices=invoices) if extractor else None


# ---------------------------------------------------------------------------
# Resumen por producto y cruce con el POS (CMV)
# ---------------------------------------------------------------------------

def _shown_change(cost, previous_cost):
    """
    El cambio que se muestra: la resta de los dos costos tal como se ven (2
    decimales cortados), para que la cuenta dé a ojo ($1.91 - $1.77 = $0.14).
    El estado Subió/Bajó sigue saliendo del cambio exacto.
    """
    return round(_cut(cost, 2) - _cut(previous_cost, 2), 2)


def _pct(part, whole):
    return round(part / whole * 100, 1) if whole else None


def _pos_diff(pos_cost, unit_cost):
    if pos_cost is None:
        return None
    diff = pos_cost - unit_cost
    return 0.0 if abs(diff) <= 0.0051 else round(diff, 2)


def _group_lines(lines):
    """{clave de producto: [renglones del más viejo al más nuevo]}."""
    by_product = {}
    for line in lines:
        key = product_key(line["supplier_key"], line["upc"], line["item_no"], line["description"])
        by_product.setdefault(key, []).append(line)
    return by_product


def _pos_index(pos_costs):
    return {normalize_upc(row.get("upc")): row for row in pos_costs if normalize_upc(row.get("upc"))}


def _summary(key, product_lines, supplier_labels, pos_by_upc):
    last = product_lines[-1]
    # Compra anterior = el último renglón de OTRA factura (el mismo producto
    # puede venir dos veces en una misma factura).
    previous = next(
        (line for line in reversed(product_lines[:-1])
         if (line["supplier_key"], line["invoice_no"]) != (last["supplier_key"], last["invoice_no"])),
        None,
    )
    change = round(last["unit_cost"] - previous["unit_cost"], 9) if previous else None
    pos = pos_by_upc.get(last["upc"]) if last["upc"] else None
    pos_cost = pos.get("cost") if pos else None
    pos_price = pos.get("price") if pos else None
    costs = [line["unit_cost"] for line in product_lines]
    return {
        "key": key,
        "upc": last["upc"],
        "item_no": last["item_no"],
        "description": last["description"],
        "category": last["category"],
        "size": last["size"],
        "units_per_case": last["units"],
        "suppliers": sorted({supplier_labels.get(line["supplier_key"], line["supplier_key"]) for line in product_lines}),
        "purchases": len({(line["supplier_key"], line["invoice_no"]) for line in product_lines}),
        "units_bought": sum((line["qty"] or 0) * (line["units"] or 0) for line in product_lines),
        "first_date": product_lines[0]["invoice_date"],
        "last_date": last["invoice_date"],
        "last_cost": last["unit_cost"],
        "previous_cost": previous["unit_cost"] if previous else None,
        "previous_date": previous["invoice_date"] if previous else None,
        "change": change,
        "shown_change": _shown_change(last["unit_cost"], previous["unit_cost"]) if previous else None,
        "change_pct": _pct(change, previous["unit_cost"]) if previous else None,
        "min_cost": min(costs),
        "max_cost": max(costs),
        "srp": last["srp"],
        "pos_name": pos.get("name") if pos else None,
        "pos_dept": pos.get("dept_name") if pos else None,
        "pos_cost": pos_cost,
        "pos_price": pos_price,
        # Diferencia entre el costo que tiene cargado el POS y el costo real
        # de la última compra -- si no es 0, el CMV está calculando el margen
        # con un costo viejo. Medio centavo o menos es redondeo del POS
        # (ej. $12.35 la caja de 10 = $1.235 la unidad, en el POS $1.24).
        "pos_cost_diff": _pos_diff(pos_cost, last["unit_cost"]),
        "margin_pct": _pct(pos_price - last["unit_cost"], pos_price) if pos_price else None,
    }


def build_product_list(lines, supplier_labels, pos_costs):
    """
    Un renglón por producto comprado (clave = UPC, o Item # si la factura no
    trae UPC), del comprado más recientemente al más viejo.
    """
    pos_by_upc = _pos_index(pos_costs)
    products = [
        _summary(key, product_lines, supplier_labels, pos_by_upc)
        for key, product_lines in _group_lines(lines).items()
    ]
    products.sort(key=lambda p: (p["last_date"], p["description"] or ""), reverse=True)
    return products


# ---------------------------------------------------------------------------
# Carpetas: proveedor -> facturas por fecha -> productos (2026-10-02)
# ---------------------------------------------------------------------------
# Pedido del usuario: "distintos módulos con fechas dentro de ellos como si
# fueran carpetas", solo de los proveedores que se leen bien (los de
# LINE_EXTRACTORS -- los únicos que guardan renglones), y cada factura
# compara sus precios contra la factura ANTERIOR POR FECHA, no por orden de
# carga: si se sube una del medio después, la más nueva pasa a compararse
# contra esa. Por eso todo se recalcula al mostrar, ordenando por fecha.

def _compare_supplier_invoices(supplier_lines):
    """
    Facturas de un proveedor del más vieja a la más nueva (fecha y N°), cada
    renglón con el costo de la última compra ANTERIOR de ese producto (de
    una factura con fecha previa; dos renglones de una misma factura nunca
    se comparan entre sí). Devuelve [(fecha, N°, [renglones])].
    """
    by_invoice = {}
    for line in sorted(supplier_lines, key=lambda l: (l["invoice_date"], str(l["invoice_no"]), l["line_no"] or 0)):
        by_invoice.setdefault((line["invoice_date"], line["invoice_no"]), []).append(line)
    last_by_product = {}
    result = []
    for (invoice_date, invoice_no), invoice_lines in sorted(by_invoice.items(), key=lambda item: (item[0][0], str(item[0][1]))):
        rows = []
        seen = {}
        for line in invoice_lines:
            key = product_key(line["supplier_key"], line["upc"], line["item_no"], line["description"])
            previous = last_by_product.get(key)
            if previous is None and not line["upc"] and line["supplier_key"] in NAME_KEYED_SUPPLIERS:
                # Sin la letra del envase (o con ella, si la anterior no la tenía): vale
                # la compra anterior solo si hay una única con ese nombre.
                found = [k for k in _container_alternatives(line["supplier_key"], line["description"])
                         if k in last_by_product]
                if len(found) == 1:
                    key, previous = found[0], last_by_product[found[0]]
            change = round(line["unit_cost"] - previous["unit_cost"], 9) if previous else None
            if change is None:
                state = "Nuevo"
            elif change > 1e-7:
                state = "Subió"
            elif change < -1e-7:
                state = "Bajó"
            else:
                state = "Igual"
            rows.append({
                **line,
                "product_key": key,
                "previous_cost": previous["unit_cost"] if previous else None,
                "previous_date": previous["invoice_date"] if previous else None,
                "previous_invoice_no": previous["invoice_no"] if previous else None,
                "change": change,
                "shown_change": _shown_change(line["unit_cost"], previous["unit_cost"]) if previous else None,
                "change_pct": _pct(change, previous["unit_cost"]) if previous else None,
                "state": state,
            })
            seen[key] = line
        last_by_product.update(seen)
        result.append((invoice_date, invoice_no, rows))
    return result


def _invoice_summary(invoice_date, invoice_no, rows, previous):
    count = lambda state: sum(1 for row in rows if row["state"] == state)
    return {
        "invoice_date": invoice_date,
        "invoice_no": invoice_no,
        "lines": len(rows),
        "total": round(sum(row["ext"] or 0.0 for row in rows), 2),
        "up": count("Subió"),
        "down": count("Bajó"),
        "same": count("Igual"),
        "new": count("Nuevo"),
        "previous_date": previous[0] if previous else None,
        "previous_invoice_no": previous[1] if previous else None,
    }


def build_supplier_invoices(supplier_key, lines):
    """Facturas del proveedor con su resumen contra la anterior, de la más nueva a la más vieja."""
    compared = _compare_supplier_invoices([line for line in lines if line["supplier_key"] == supplier_key])
    summaries = []
    for index, (invoice_date, invoice_no, rows) in enumerate(compared):
        previous = compared[index - 1][:2] if index > 0 else None
        summaries.append(_invoice_summary(invoice_date, invoice_no, rows, previous))
    summaries.reverse()
    return summaries


def build_invoice_products(supplier_key, invoice_date, invoice_no, lines):
    """(resumen, renglones) de una factura con cada producto contra su compra anterior; None si no existe."""
    compared = _compare_supplier_invoices([line for line in lines if line["supplier_key"] == supplier_key])
    for index, (date_value, number, rows) in enumerate(compared):
        if date_value == invoice_date and str(number) == str(invoice_no):
            previous = compared[index - 1][:2] if index > 0 else None
            return _invoice_summary(date_value, number, rows, previous), rows
    return None


def build_supplier_folders(lines, supplier_labels):
    """Una carpeta por proveedor con renglones guardados: cuántas facturas, de qué fechas y cómo vino la última."""
    keys = sorted({line["supplier_key"] for line in lines}, key=lambda k: supplier_labels.get(k, k).lower())
    folders = []
    for key in keys:
        invoices = build_supplier_invoices(key, lines)
        folders.append({
            "key": key,
            "label": supplier_labels.get(key, key),
            "invoices": len(invoices),
            "first_date": invoices[-1]["invoice_date"] if invoices else None,
            "last": invoices[0] if invoices else None,
            "years": sorted({inv["invoice_date"][:4] for inv in invoices}, reverse=True),
        })
    return folders


def build_product_detail(key, lines, supplier_labels, pos_costs, monthly_sales):
    """
    Resumen + historial de compras de un producto, y sus ventas por mes del
    POS (CMV) con el costo real de ese momento -- para ver a cuánto se vende
    contra lo que costó. None si ese producto nunca se compró.
    """
    product_lines = _group_lines(lines).get(key)
    if not product_lines:
        return None
    summary = _summary(key, product_lines, supplier_labels, _pos_index(pos_costs))

    purchases = []
    previous_cost = None
    for line in product_lines:
        change = round(line["unit_cost"] - previous_cost, 9) if previous_cost is not None else None
        purchases.append({
            **line,
            "supplier_label": supplier_labels.get(line["supplier_key"], line["supplier_key"]),
            "change": change,
            "shown_change": _shown_change(line["unit_cost"], previous_cost) if change is not None else None,
            "change_pct": _pct(change, previous_cost) if change is not None else None,
        })
        previous_cost = line["unit_cost"]
    purchases.reverse()

    sales = []
    if summary["upc"]:
        for row in monthly_sales:
            if normalize_upc(row.get("upc")) != summary["upc"] or not row.get("count"):
                continue
            month_end = f"{row['year']:04d}-{row['month']:02d}-31"
            # Costo vigente en ese mes = el de la última compra hasta fin de mes.
            cost_lines = [line for line in product_lines if line["invoice_date"] <= month_end]
            cost = cost_lines[-1]["unit_cost"] if cost_lines else None
            avg_price = row["amount"] / row["count"]
            sales.append({
                "year": row["year"],
                "month": row["month"],
                "count": row["count"],
                "amount": row["amount"],
                "avg_price": round(avg_price, 2),
                "cost": cost,
                "margin_pct": _pct(avg_price - cost, avg_price) if cost is not None and avg_price else None,
            })
        sales.reverse()

    return {"product": summary, "purchases": purchases, "sales": sales}


# ---------------------------------------------------------------------------
# Ventas y costos por mes (pedido del usuario, 2026-10-08): las ventas por
# producto de Elistar (CMV Ventas, Sales Insights) mes por mes, con el costo
# de la última compra a proveedor hasta fin de cada mes (el POS solo guarda
# el costo de hoy, no el de una fecha pasada), para ver qué precios se
# mantienen y si un costo que cambió se pasó o no al precio de venta.
# El precio promedio es importe / cantidad: Elistar ya le restó los
# descuentos Mix & Match, así que se mueve un poco aunque el precio de
# lista no cambie; por eso un cambio de precio cuenta desde 5 centavos y 2 %.
# ---------------------------------------------------------------------------

_PRICE_MIN_CHANGE = 0.05
_PRICE_MIN_PCT = 2.0

SALES_COST_STATES = (
    "Subió el costo, no el precio",
    "Subió el costo y el precio",
    "Bajó el costo",
    "Cambió el precio",
    "Sin cambios",
    "Sin compra cargada",
)


def _price_changed(first, last):
    return abs(last - first) >= _PRICE_MIN_CHANGE and first and abs(last - first) / first * 100 >= _PRICE_MIN_PCT


def build_sales_cost_view(sales_rows, lines, year_months):
    """
    Un renglón por producto vendido (departamento + UPC) en los meses
    `year_months` (de más viejo a más nuevo): por mes, cantidad, importe,
    precio promedio, costo de la última compra hasta fin de mes y margen;
    y el estado comparando el primer mes con venta contra el último.
    """
    purchases = {}
    for line in lines:
        upc = normalize_upc(line.get("upc"))
        if upc and line.get("unit_cost") is not None:
            purchases.setdefault(upc, []).append(line)  # vienen del más viejo al más nuevo

    products = {}
    for row in sales_rows:
        upc = normalize_upc(row.get("upc"))
        if not upc or not row.get("count"):
            continue
        product = products.setdefault((row["dept_name"], upc), {
            "dept_name": row["dept_name"], "upc": upc, "name": row.get("name") or "", "months": {}})
        cell = product["months"].setdefault((row["year"], row["month"]), {"count": 0, "amount": 0.0})
        cell["count"] += row["count"]
        cell["amount"] += row["amount"] or 0.0

    result = []
    for product in products.values():
        bought = purchases.get(product["upc"], [])
        cells = []
        for year, month in year_months:
            cell = product["months"].get((year, month))
            if cell is None:
                cells.append(None)
                continue
            month_end = f"{year:04d}-{month:02d}-31"
            known = [line for line in bought if line["invoice_date"] <= month_end]
            cost = known[-1]["unit_cost"] if known else None
            avg_price = cell["amount"] / cell["count"]
            cells.append({
                "count": cell["count"], "amount": round(cell["amount"], 2), "avg_price": round(avg_price, 2),
                "cost": cost, "margin_pct": _pct(avg_price - cost, avg_price) if cost is not None and avg_price else None,
            })
        sold = [c for c in cells if c is not None]
        with_cost = [c for c in sold if c["cost"] is not None]
        price_change = round(sold[-1]["avg_price"] - sold[0]["avg_price"], 2)
        cost_change = round(with_cost[-1]["cost"] - with_cost[0]["cost"], 9) if with_cost else None
        price_moved = _price_changed(sold[0]["avg_price"], sold[-1]["avg_price"])
        if cost_change is None:
            state = "Sin compra cargada"
        elif cost_change > 0.005:
            state = "Subió el costo y el precio" if price_moved and price_change > 0 else "Subió el costo, no el precio"
        elif cost_change < -0.005:
            state = "Bajó el costo"
        elif price_moved:
            state = "Cambió el precio"
        else:
            state = "Sin cambios"
        last_line = bought[-1] if bought else None
        result.append({
            **{k: product[k] for k in ("dept_name", "upc", "name")},
            "cells": cells,
            "count": sum(c["count"] for c in sold),
            "amount": round(sum(c["amount"] for c in sold), 2),
            "price_change": price_change if price_moved else 0.0,
            "cost_change": cost_change,
            "state": state,
            "last_purchase": last_line["invoice_date"] if last_line else None,
            "product_key": product_key(last_line["supplier_key"], last_line["upc"], last_line["item_no"],
                                       last_line.get("description")) if last_line else None,
        })
    result.sort(key=lambda p: (p["dept_name"].lower(), -p["amount"]))
    return result


# ---------------------------------------------------------------------------
# Reporte para el manager (pedido del usuario, 2026-10-02): "enviar un
# reporte en Excel o PDF directamente al manager de la estación de servicio
# con los productos que cambiaron de precio". Solo los que subieron o
# bajaron contra su compra anterior; con el precio del POS y el margen que
# deja el costo nuevo, para que sepa qué precio revisar.
# ---------------------------------------------------------------------------

def with_departments(rows, pos_costs):
    """
    Cada renglón con el departamento del POS de su UPC (pedido del usuario,
    2026-10-06: "así se los identifica más fácil"); sin UPC en el POS, la
    categoría de la factura.
    """
    pos_by_upc = _pos_index(pos_costs)
    result = []
    for row in rows:
        pos = pos_by_upc.get(row["upc"]) if row["upc"] else None
        result.append({**row, "department": (pos.get("dept_name") if pos else None) or row.get("category")})
    return result


def price_change_rows(rows, pos_costs):
    """Renglones de una factura (build_invoice_products) que cambiaron de costo: primero los que subieron."""
    pos_by_upc = _pos_index(pos_costs)
    changed = []
    for row in with_departments(rows, pos_costs):
        if row["state"] not in ("Subió", "Bajó"):
            continue
        if not row.get("shown_change"):
            # De $1.912 a $1.915: con 2 decimales cortados no se ve ningún cambio
            # y el reporte decía "Up +$0.00" (revisión 2026-10-08).
            continue
        pos = pos_by_upc.get(row["upc"]) if row["upc"] else None
        pos_price = pos.get("price") if pos else None
        margin = _pct(pos_price - row["unit_cost"], pos_price) if pos_price else None
        changed.append({**row, "pos_price": pos_price, "margin_pct": margin})
    changed.sort(key=lambda r: (0 if r["state"] == "Subió" else 1, r.get("department") or "", r.get("description") or ""))
    return changed


def _pack_label(row):
    return " ".join(str(part) for part in (row.get("pack"), row.get("size")) if part not in (None, ""))


# El reporte va en inglés (es para el manager, pedido del usuario 2026-10-06),
# con fechas MM/DD/YYYY, sin la fecha del costo anterior, el precio del POS
# como "Elistar Price" y el departamento. Los costos van sin redondear
# (pedido del usuario 2026-10-07), igual que en la pantalla.
_STATE_EN = {"Subió": "Up", "Bajó": "Down"}


def exact_money(value, places=6):
    """
    Costo por unidad sin redondear: $1.91625, no $1.92. Se cortan los ceros
    de más (mínimo 2 decimales) y, si la división no termina (34/12), se
    corta en el decimal `places` sin redondear. El costo y el costo anterior
    se muestran con places=2 (pedido del usuario 2026-10-07): $1.91625 -> $1.91.
    """
    from decimal import Decimal, ROUND_DOWN

    if value is None:
        return "—"
    digits = Decimal(f"{abs(value):.10f}").quantize(Decimal(1).scaleb(-places), rounding=ROUND_DOWN)
    text = f"{digits:,.{places}f}".rstrip("0")
    if len(text.split(".")[1]) < 2:
        text = f"{digits:,.2f}"
    return ("-$" if value < 0 and digits else "$") + text


def _mmddyyyy(iso):
    return f"{iso[5:7]}/{iso[8:10]}/{iso[0:4]}" if iso else ""


def build_price_change_pdf(supplier_label, invoice, changed, dest_path):
    from pdf_export import build_simple_table_pdf

    def money(value):
        return "" if value is None else f"{'-' if value < 0 else ''}${abs(value):,.2f}"

    table_rows = []
    for row in changed:
        sign = "+" if row["change"] > 0 else ""
        table_rows.append([
            row["upc"] or "",
            row.get("description") or "",
            row.get("department") or "",
            _pack_label(row),
            exact_money(row["previous_cost"], 2),
            exact_money(row["unit_cost"], 2),
            f"{sign}{exact_money(row['shown_change'], 2)}" + (f" ({sign}{row['change_pct']}%)" if row.get("change_pct") is not None else ""),
            money(row.get("srp")),
            money(row.get("pos_price")),
            "" if row.get("margin_pct") is None else f"{row['margin_pct']}%",
        ])
    up = sum(1 for row in changed if row["state"] == "Subió")
    build_simple_table_pdf(
        dest_path,
        f"{supplier_label} — Price Changes — Invoice #{invoice['invoice_no']}",
        ["UPC", "Product", "Department", "Pack", "Previous Cost", "New Cost", "Change", "Invoice SRP",
         "Elistar Price", "Margin"],
        table_rows,
        col_widths_mm=[30, 78, 34, 20, 24, 24, 34, 22, 24, 20],
        company_header=True,
        period_label=(
            f"Invoice dated {_mmddyyyy(invoice['invoice_date'])} — {up} went up, {len(changed) - up} went down"
            + (" — compared with each product's previous purchase" if changed else " — no product changed price")
        ),
        footer_note="Cost per unit. Margin = (Elistar price − new cost) / Elistar price.",
    )
    return dest_path


def _cut(value, places=6):
    """Valor cortado (no redondeado) en el decimal `places`: el formato de Excel redondearía 1.7704166 a 1.770417."""
    from decimal import Decimal, ROUND_DOWN

    if value is None:
        return None
    return float(Decimal(f"{value:.10f}").quantize(Decimal(1).scaleb(-places), rounding=ROUND_DOWN))


def build_price_change_workbook(supplier_label, invoice, changed, dest_path):
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = openpyxl.Workbook()
    sheet = workbook.active
    sheet.title = "Price Changes"
    sheet.append([f"{supplier_label} — Invoice #{invoice['invoice_no']} dated {_mmddyyyy(invoice['invoice_date'])}"])
    sheet["A1"].font = Font(bold=True, size=12)
    sheet.append([])
    headers = ["Change", "UPC", "Product", "Department", "Pack", "Previous Cost", "New Cost",
               "Difference", "Difference %", "Invoice SRP", "Elistar Price", "Margin %"]
    sheet.append(headers)
    for cell in sheet[3]:
        cell.font = Font(bold=True)
    up_fill = PatternFill("solid", fgColor="FEE2E2")
    down_fill = PatternFill("solid", fgColor="DCFCE7")
    money_format = '"$"#,##0.00'
    for row in changed:
        sheet.append([
            _STATE_EN[row["state"]], row["upc"], row.get("description"), row.get("department"), _pack_label(row),
            _cut(row["previous_cost"], 2), _cut(row["unit_cost"], 2), row["shown_change"],
            row["change_pct"], row.get("srp"), row.get("pos_price"), row.get("margin_pct"),
        ])
        new_row = sheet[sheet.max_row]
        new_row[0].fill = up_fill if row["state"] == "Subió" else down_fill
        for index in (5, 6, 7, 9, 10):
            new_row[index].number_format = money_format
        for cell in new_row[5:]:
            cell.alignment = Alignment(horizontal="center")
    for col_letter, width in zip("ABCDEFGHIJKL", (9, 15, 38, 18, 10, 14, 12, 12, 13, 12, 13, 11)):
        sheet.column_dimensions[col_letter].width = width
    if changed:
        sheet.auto_filter.ref = f"A3:L{sheet.max_row}"
    sheet.freeze_panes = "A4"
    workbook.save(dest_path)
    return dest_path
