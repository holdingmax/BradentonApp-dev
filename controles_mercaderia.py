"""
Controles: Mercadería en C-Store -- cuarto módulo de la sección Controles.

Cruza el total de facturas cargadas en el mes en el libro de Proveedores
("Bradenton. Cta Cte Proveedores.xlsx", una hoja por proveedor) contra el
movimiento neto (Débito - Crédito) del Mayor de la cuenta 1.04.02.00 -
MERCADERIA EN C-STORE para ese mismo período.

Pautas de negocio confirmadas con el usuario (2026-09-08):
- Todas las hojas de proveedor cuentan para esta cuenta EXCEPTO AIRGAS,
  KOOLER ICE y SIGNARAMA (no van a Mercadería), más "J.H W" -- que a pesar
  del nombre es "J.H. Williams OIL" (cuenta 2.01.01.00, combustible, sin
  relación con Mercadería) -- encontrado revisando el libro real, no
  nombrado explícitamente por el usuario pero confirmado con él antes de
  excluirlo. No hace falta mantener una lista de "quién sí cuenta": una
  hoja de proveedor nuevo que se agregue al libro entra sola (se
  reconoce cualquier hoja con al menos una fecha real en columna A,
  mismo criterio que ya usa `proveedores._find_last_real_row` -- esto
  también es lo que excluye solas a las hojas que no son de proveedores,
  como "RESUMEN COMPRAS" o "Resumen").
- El período a revisar se toma de la línea "Desde: ... Hasta: ..." que ya
  trae el Mayor exportado (mismo mecanismo que controles_cupones.py) --
  no hay selector de mes manual, para no arriesgar que quede
  desincronizado con el archivo subido.
- Las filas "invoice" (columna DEBE) suman al total y cuentan para
  proveedores_invoice_count. Las filas "OP" (pagos) nunca cuentan --
  son plata que sale del banco, no mercadería. Cualquier OTRA fila con
  fecha en el período (ej. "CREDIT MM" en HT Hackney, o las variantes
  "AJUSTE"/"ANULACION"/"Void"/"Credit" que aparecen en otras hojas)
  se toma como nota de crédito/ajuste real contra Mercadería y se RESTA
  por su columna HABER -- confirmado con el usuario (2026-09-08): "H.T.
  tiene notas de crédito que se cargan en contra de Mercadería en
  C-Store, y que disminuyen su valor". No se depende del texto exacto
  de la etiqueta (cada proveedor usa una distinta) -- alcanza con que
  la fila no sea "invoice" ni "OP" y tenga algo en HABER, exactamente
  el mismo criterio que ya usa la fórmula de BALANCE de cada hoja
  (`=+anterior+DEBE-HABER`). Estas filas de crédito NO suman a
  proveedores_invoice_count (no son una factura nueva).

Es de solo lectura: nunca escribe nada en ningún Excel ni genera un
archivo para descargar -- el resultado se muestra en pantalla, según lo
ya decidido para toda la sección Controles.
"""

import os
import re
from datetime import datetime

import pandas as pd
from openpyxl import load_workbook

from proveedores import COL_COMPROB, COL_DATE, COL_DEBE, COL_HABER, _find_last_real_row

# Confirmado con el usuario (2026-09-08): a diferencia de Cupones, acá los
# dos lados representan el mismo conjunto de facturas físicas del mes --
# tienen que coincidir casi exacto. Margen de ±$0.05 (no $0 exacto): al
# sumar el neto DEBE-HABER de ~20 hojas por separado, cada una redondeada
# a 2 decimales antes de sumarse, se acumulan un par de centavos de
# diferencia contra el total del Mayor aun cuando todo está cargado bien
# -- confirmado con un caso real (julio-2026, notas de crédito de HT
# Hackney incluidas): diferencia final de $0.03, no un error de carga.
# Una diferencia real (factura faltante/de más) da un número mucho mayor
# a un par de centavos, así que este margen no tapa un error real.
TOLERANCE = 0.05

EXCLUDED_SHEETS = {"AIRGAS", "KOOLER ICE", "SIGNARAMA", "J.H W"}

_PERIOD_LINE_RE = re.compile(r"desde:\s*([\d/-]+)\s*hasta:\s*([\d/-]+)", re.IGNORECASE)
_MAYOR_DATE_FORMATS = ("%Y-%m-%d", "%d/%m/%Y", "%m/%d/%Y")


def _parse_mayor_date(text):
    text = text.strip()
    for fmt in _MAYOR_DATE_FORMATS:
        try:
            return datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _format_date(value):
    """Formatea una fecha para el template -- tolera texto en vez de fecha real."""
    if value is None:
        return None
    if hasattr(value, "strftime"):
        return value.strftime("%d/%m/%Y")
    text = str(value).strip()
    return text or None


def _read_mercaderia_mayor(path):
    """
    Lee el Mayor de la cuenta Mercadería en C-Store (.xls o .xlsx).

    Busca la línea "Desde: ... Hasta: ..." y la fila de encabezado real
    por su propio contenido ("Fecha"/"Debito"/"Credito", nunca una fila
    fija), igual que controles_cupones.py. Suma Débito y Crédito de cada
    fila con una fecha real -- la fila "SALDO INICIAL" y la de totales al
    pie no tienen fecha, así que quedan afuera solas; no depende de que
    exista una fila de totales con una forma exacta, se recalcula desde
    las filas reales.

    Devuelve (period_from, period_to, total_debito, total_credito, invoice_count)
    -- invoice_count cuenta solo las filas con Débito (una factura real),
    nunca las de Crédito (notas de crédito, ej. "CREDIT MEMO ..."), para
    poder comparar cantidad de facturas contra el lado de Proveedores.
    """
    df = pd.read_excel(path, header=None)

    period_from = period_to = None
    header_row_index = None
    col_index = {}

    for idx in range(len(df)):
        cells = ["" if pd.isna(v) else str(v).strip() for v in df.iloc[idx]]
        if period_from is None:
            joined = " ".join(c for c in cells if c)
            match = _PERIOD_LINE_RE.search(joined)
            if match:
                period_from = _parse_mayor_date(match.group(1))
                period_to = _parse_mayor_date(match.group(2))
        if header_row_index is None:
            lowered = [c.lower() for c in cells]
            if "fecha" in lowered and "debito" in lowered and "credito" in lowered:
                header_row_index = idx
                col_index = {
                    "fecha": lowered.index("fecha"),
                    "debito": lowered.index("debito"),
                    "credito": lowered.index("credito"),
                }

    if header_row_index is None:
        raise ValueError('No se encontró la fila de encabezado ("Fecha"/"Debito"/"Credito") en el Mayor.')
    if period_from is None or period_to is None:
        raise ValueError('No se encontró la línea "Desde: ... Hasta: ..." en el Mayor.')

    total_debito = 0.0
    total_credito = 0.0
    invoice_count = 0
    for idx in range(header_row_index + 1, len(df)):
        fecha = df.iat[idx, col_index["fecha"]]
        if pd.isna(fecha):
            continue
        debito = df.iat[idx, col_index["debito"]]
        credito = df.iat[idx, col_index["credito"]]
        if pd.notna(debito) and float(debito) > 0:
            total_debito += float(debito)
            invoice_count += 1
        total_credito += float(credito) if pd.notna(credito) else 0.0

    return period_from, period_to, round(total_debito, 2), round(total_credito, 2), invoice_count


def _sum_sheet_invoices_for_period(sheet, period_from, period_to):
    """
    Neto de mercadería de una hoja de proveedor dentro del período: suma
    DEBE de las filas "invoice" (cuentan para el conteo de facturas) y
    resta HABER de cualquier otra fila que no sea "OP" (notas de
    crédito/ajustes -- cada proveedor usa su propia etiqueta para esto,
    ver el docstring del módulo) -- nunca cuentan como factura nueva.
    """
    total = 0.0
    count = 0
    last_row = _find_last_real_row(sheet)
    if last_row is None:
        return 0.0, 0
    for row in range(1, last_row + 1):
        date_value = sheet.cell(row=row, column=COL_DATE).value
        if not isinstance(date_value, datetime):
            continue
        invoice_date = date_value.date()
        if not (period_from <= invoice_date <= period_to):
            continue
        comprob = sheet.cell(row=row, column=COL_COMPROB).value
        comprob_text = comprob.strip().lower() if isinstance(comprob, str) else ""
        if comprob_text == "op":
            continue
        if comprob_text == "invoice":
            amount = sheet.cell(row=row, column=COL_DEBE).value
            if isinstance(amount, (int, float)):
                total += float(amount)
                count += 1
        else:
            credit = sheet.cell(row=row, column=COL_HABER).value
            if isinstance(credit, (int, float)):
                total -= float(credit)
    return round(total, 2), count


def check_mercaderia_invoices(proveedores_path, mayor_path):
    """
    Cruza el total de facturas del mes en el libro de Proveedores (todas
    las hojas, salvo EXCLUDED_SHEETS) contra el neto Débito-Crédito del
    Mayor de Mercadería en C-Store para el mismo período.

    Devuelve un dict con el período del Mayor, los totales de cada lado,
    el desglose por proveedor, la diferencia y el chequeo (ok). Nunca
    escribe ni descarga nada.
    """
    proveedores_path = os.path.abspath(str(proveedores_path).strip())
    if not os.path.isfile(proveedores_path):
        raise FileNotFoundError(f"Excel de Proveedores no encontrado: {proveedores_path}")
    extension = os.path.splitext(proveedores_path)[1].lower()
    if extension not in {".xlsx", ".xlsm"}:
        raise ValueError("El Excel de Proveedores debe ser .xlsx o .xlsm.")

    mayor_path = os.path.abspath(str(mayor_path).strip())
    if not os.path.isfile(mayor_path):
        raise FileNotFoundError(f"Mayor de Mercadería en C-Store no encontrado: {mayor_path}")

    period_from, period_to, mayor_debito, mayor_credito, mayor_invoice_count = _read_mercaderia_mayor(mayor_path)
    mayor_total = round(mayor_debito - mayor_credito, 2)

    workbook = load_workbook(proveedores_path, data_only=False)
    try:
        breakdown = []
        proveedores_total = 0.0
        proveedores_invoice_count = 0
        for sheet_name in workbook.sheetnames:
            if sheet_name.strip() in EXCLUDED_SHEETS:
                continue
            sheet = workbook[sheet_name]
            amount, count = _sum_sheet_invoices_for_period(sheet, period_from, period_to)
            if count > 0:
                breakdown.append({"supplier": sheet_name.strip(), "count": count, "amount": amount})
                proveedores_total += amount
                proveedores_invoice_count += count
    finally:
        workbook.close()

    breakdown.sort(key=lambda item: item["amount"], reverse=True)
    proveedores_total = round(proveedores_total, 2)
    diff = round(mayor_total - proveedores_total, 2)

    return {
        "period_from": period_from,
        "period_to": period_to,
        "period_from_display": _format_date(period_from),
        "period_to_display": _format_date(period_to),
        "mayor_debito": mayor_debito,
        "mayor_credito": mayor_credito,
        "mayor_total": mayor_total,
        "mayor_invoice_count": mayor_invoice_count,
        "proveedores_total": proveedores_total,
        "proveedores_invoice_count": proveedores_invoice_count,
        "invoice_count_diff": proveedores_invoice_count - mayor_invoice_count,
        "breakdown": breakdown,
        "diff": diff,
        "ok": abs(diff) <= TOLERANCE,
    }
